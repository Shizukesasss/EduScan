def generate_gradebook_report_xlsx(db: Session, gradebook_id: int, region: str = "", division: str = "", school_id: str = "") -> Path:
    import openpyxl
    from .grading import get_gradebook_full
    from ..models import Gradebook

    full = get_gradebook_full(db, gradebook_id)
    if not full:
        raise HTTPException(status_code=404, detail="Gradebook not found")

    gb = full["gradebook"]
    quarter: int = gb["quarter"]
    subject_name: str = gb["subject_name"] or ""
    is_mapeh = "MAPEH" in subject_name.upper()

    # 1. Fetch siblings if MAPEH
    gradebooks_to_process = []
    if is_mapeh:
        siblings = db.query(Gradebook).filter(
            Gradebook.school_year_id == gb["school_year_id"],
            Gradebook.grading_period_id == gb["grading_period_id"],
            Gradebook.section_id == gb["section_id"],
            Gradebook.subject_name.like("MAPEH%")
        ).all()
        for sib in siblings:
            sib_full = get_gradebook_full(db, sib.id)
            if sib_full:
                gradebooks_to_process.append(sib_full)
    else:
        gradebooks_to_process.append(full)

    # 2. Select the correct template
    TEMPLATES_DIR = Path(__file__).parent.parent.parent / "templates"
    subject_upper = subject_name.upper().strip()
    template_file = None

    if is_mapeh:
        q_strs = ["1ST", "2ND", "3RD", "4TH"]
        q_str = q_strs[max(0, min(3, quarter - 1))]
        template_file = TEMPLATES_DIR / f"GRADE 7-10_MAPEH {q_str} QUARTER.xlsx"
    else:
        SUBJECT_MAP = {
            "ARALING PANLIPUNAN": "GRADE 7-10_ARALING PANLIPUNAN.xlsx",
            "EDUKASYON SA PAGPAPAKATAO": "GRADE 7-10_EDUKASYON SA PAGPAPAKATAO.xlsx",
            "ESP": "GRADE 7-10_EDUKASYON SA PAGPAPAKATAO.xlsx",
            "ENGLISH": "GRADE 7-10_ENGLISH.xlsx",
            "FILIPINO": "GRADE 7-10_FILIPINO.xlsx",
            "HOME ECONOMICS": "GRADE 7-10_HOME ECONOMICS.xlsx",
            "MATHEMATICS": "GRADE 7-10_MATHEMATICS.xlsx",
            "MATH": "GRADE 7-10_MATHEMATICS.xlsx",
            "SCIENCE": "GRADE 7-10_SCIENCE.xlsx",
            "TLE": "GRADE 7-10_TLE.xlsx",
            "TECHNOLOGY AND LIVELIHOOD EDUCATION": "GRADE 7-10_TLE.xlsx",
            "TECHNICAL VOCATIONAL EDUCATION": "GRADE 7-10_TLE.xlsx",
            "TECHNICAL": "GRADE 7-10_TLE.xlsx",
            "VOCATIONAL": "GRADE 7-10_TLE.xlsx",
            "ECHNICAL": "GRADE 7-10_TLE.xlsx",
        }
        for key, fname in SUBJECT_MAP.items():
            if key in subject_upper or subject_upper.startswith(key):
                template_file = TEMPLATES_DIR / fname
                break

    if not template_file or not template_file.exists():
        available = list(TEMPLATES_DIR.glob("GRADE 7-10_*.xlsx"))
        if not available:
            raise HTTPException(status_code=500, detail="No DepEd Excel templates found in backend/templates/")
        template_file = available[0]

    wb = openpyxl.load_workbook(str(template_file))

    # Safe write helper
    def write(sheet, coord, value):
        import openpyxl.cell
        cell = sheet[coord]
        if isinstance(cell, openpyxl.cell.MergedCell):
            for rng in sheet.merged_cells.ranges:
                if coord in rng:
                    top_left = rng.coord.split(':')[0]
                    sheet[top_left] = value
                    return
        sheet[coord] = value

    # 3. Fill INPUT DATA sheet (use the first/current gradebook for header data)
    inp = wb["INPUT DATA"]
    write(inp, "G4", region)
    write(inp, "O4", division)
    write(inp, "G5", gb.get("school_name", "San Jose National High School"))
    write(inp, "X5", school_id)
    write(inp, "AG5", gb.get("school_year_name", ""))
    write(inp, "K7", f"Grade {gb['grade_name']} - {gb['section_name']}")
    write(inp, "S7", gb.get("teacher_name", ""))
    
    display_subject = "MAPEH" if is_mapeh else ("TLE" if template_file and "TLE.xlsx" in template_file.name else subject_name)
    write(inp, "AG7", display_subject)

    students = full["students"]
    males = [s for s in students if (s.get("sex") or "").upper() in ("M", "MALE")]
    females = [s for s in students if (s.get("sex") or "").upper() in ("F", "FEMALE")]
    if not males and not females:
        males = students

    MALE_NAME_ROW = 12
    FEMALE_NAME_ROW = 63

    for i, student in enumerate(males):
        inp.cell(row=MALE_NAME_ROW + i, column=2, value=student["full_name"])
    for i, student in enumerate(females):
        inp.cell(row=FEMALE_NAME_ROW + i, column=2, value=student["full_name"])

    # 4. Fill Quarter sheet(s)
    WW_START_COL = 6
    WW_MAX_ITEMS = 10
    PT_START_COL = 19
    PT_MAX_ITEMS = 10
    QA_COL = 32
    HPS_ROW = 10
    MALE_DATA_ROW = 12
    FEMALE_DATA_ROW = 63

    def _match_comp(comps, keyword: str) -> dict | None:
        for c in comps:
            name_up = (c.get("name") or "").upper()
            type_up = (c.get("component_type") or "").upper()
            if keyword in name_up or keyword in type_up:
                return c
        return None

    for proc_full in gradebooks_to_process:
        pgb = proc_full["gradebook"]
        pcomps = proc_full["components"]
        pstudents = proc_full["students"]
        psubj = (pgb["subject_name"] or "").upper()
        
        # Select target sheet
        if is_mapeh:
            target_sheet_name = None
            if "MUSIC" in psubj:
                target_sheet_name = next((s for s in wb.sheetnames if "MUSIC" in s.upper()), None)
            elif "ARTS" in psubj:
                target_sheet_name = next((s for s in wb.sheetnames if "ARTS" in s.upper()), None)
            elif "PE" in psubj:
                target_sheet_name = next((s for s in wb.sheetnames if "PE" in s.upper()), None)
            elif "HEALTH" in psubj:
                target_sheet_name = next((s for s in wb.sheetnames if "HEALTH" in s.upper()), None)
            
            if not target_sheet_name:
                continue
            q_sheet = wb[target_sheet_name]
        else:
            quarter_idx = max(1, min(4, quarter))
            q_sheet = wb.worksheets[quarter_idx]

        # Write header to sheet
        write(q_sheet, "G4", region)
        write(q_sheet, "O4", division)
        write(q_sheet, "G5", pgb.get("school_name", "San Jose National High School"))
        write(q_sheet, "X5", school_id)
        write(q_sheet, "AG5", pgb.get("school_year_name", ""))
        write(q_sheet, "K7", f"Grade {pgb['grade_name']} - {pgb['section_name']}")
        write(q_sheet, "S7", pgb.get("teacher_name", ""))

        ww_comp = _match_comp(pcomps, "WRITTEN")
        pt_comp = _match_comp(pcomps, "PERFORMANCE")
        qa_comp = next((c for c in pcomps if "QUARTERLY" in (c.get("name") or "").upper() or "QUARTERLY" in (c.get("component_type") or "").upper()), 
                       next((c for c in pcomps if "ASSESSMENT" in (c.get("name") or "").upper() or "ASSESSMENT" in (c.get("component_type") or "").upper()), None))

        if ww_comp:
            for i, item in enumerate(ww_comp.get("items") or []):
                if i >= WW_MAX_ITEMS:
                    break
                q_sheet.cell(row=HPS_ROW, column=WW_START_COL + i, value=item.get("max_score", 0))

        if pt_comp:
            for i, item in enumerate(pt_comp.get("items") or []):
                if i >= PT_MAX_ITEMS:
                    break
                q_sheet.cell(row=HPS_ROW, column=PT_START_COL + i, value=item.get("max_score", 0))

        if qa_comp:
            qa_items = qa_comp.get("items") or []
            if qa_items:
                qa_hps_total = sum(item.get("max_score", 0) for item in qa_items)
                q_sheet.cell(row=HPS_ROW, column=QA_COL, value=qa_hps_total)

        pmales = [s for s in pstudents if (s.get("sex") or "").upper() in ("M", "MALE")]
        pfemales = [s for s in pstudents if (s.get("sex") or "").upper() in ("F", "FEMALE")]
        if not pmales and not pfemales:
            pmales = pstudents
            
        def fill_student_row(student: dict, data_row: int) -> None:
            scores = student.get("scores") or {}
            q_sheet.cell(row=data_row, column=2, value=student["full_name"])
            
            if ww_comp:
                for i, item in enumerate(ww_comp.get("items") or []):
                    if i >= WW_MAX_ITEMS: break
                    score = scores.get(str(item["id"]))
                    if score is not None:
                        q_sheet.cell(row=data_row, column=WW_START_COL + i, value=float(score))
                        
            if pt_comp:
                for i, item in enumerate(pt_comp.get("items") or []):
                    if i >= PT_MAX_ITEMS: break
                    score = scores.get(str(item["id"]))
                    if score is not None:
                        q_sheet.cell(row=data_row, column=PT_START_COL + i, value=float(score))
                        
            if qa_comp:
                qa_items = qa_comp.get("items") or []
                if qa_items:
                    qa_score_total: float | None = None
                    for item in qa_items:
                        raw = scores.get(str(item["id"]))
                        if raw is not None:
                            qa_score_total = (qa_score_total or 0.0) + float(raw)
                    if qa_score_total is not None:
                        q_sheet.cell(row=data_row, column=QA_COL, value=qa_score_total)

        for i, student in enumerate(pmales):
            fill_student_row(student, MALE_DATA_ROW + i)
        for i, student in enumerate(pfemales):
            fill_student_row(student, FEMALE_DATA_ROW + i)

    # 5. Fill the Summary sheet
    if "SUMMARY OF QUARTERLY GRADES" in wb.sheetnames:
        sum_sheet = wb["SUMMARY OF QUARTERLY GRADES"]
        SUMMARY_MALE_ROW = 13
        SUMMARY_FEMALE_ROW = 64
        for i, student in enumerate(males):
            sum_sheet.cell(row=SUMMARY_MALE_ROW + i, column=2, value=student["full_name"])
        for i, student in enumerate(females):
            sum_sheet.cell(row=SUMMARY_FEMALE_ROW + i, column=2, value=student["full_name"])

    # 6. Save
    safe_subj_name = "MAPEH" if is_mapeh else subject_name
    target = EXPORT_DIR / (
        f"ClassRecord-{_safe_name(gb['school_year_name'])}-Q{quarter}"
        f"-Grade{_safe_name(gb['grade_name'])}-{_safe_name(gb['section_name'])}"
        f"-{_safe_name(safe_subj_name)}.xlsx"
    )
    wb.save(str(target))
    return target
