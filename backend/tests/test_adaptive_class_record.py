"""Custom workbook calculations; no school database or real student records."""
import copy
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.services.adaptive_class_record import _term_sheet
from app.services.grading_engine import calculate_student_grade, build_default_transmutation_table


def sample_record(method='half_up', exams=False):
    components = [
        dict(id=1, name='Written Work', weight=25, items=[dict(id=1, label='Essay', max_score=30, active=True), dict(id=2, label='Quiz', max_score=20, active=True), dict(id=3, label='Unused', max_score=100, active=False)]),
        dict(id=2, name='Performance Tasks', weight=35, items=[dict(id=4, label='Project', max_score=40, active=True)]),
        dict(id=3, name='Quarterly Assessment', weight=30, component_type='Term Examinations' if exams else 'Quarterly Assessment', items=[dict(id=5, label='ST1', max_score=30, active=True), dict(id=6, label='ST2', max_score=20, active=True), dict(id=7, label='TE', max_score=50, active=True)]),
        dict(id=4, name='Attendance', weight=10, items=[dict(id=8, label='Days present', max_score=20, active=True)]),
    ]
    scores = {str(key): value for key, value in ((1, 24), (2, 18), (4, 33), (5, 25), (6, 14), (7, 41), (8, 19))}
    table = build_default_transmutation_table()
    result = calculate_student_grade([dict(item_id=int(key), score=value, status='Scored') for key, value in scores.items()], components,
                                     {component['id']: [item for item in component['items'] if item['active']] for component in components},
                                     table, rounding_method=method)
    student = dict(person_id=1, full_name='Synthetic, Learner', sex='Male', external_id='TEST-001', lrn='TEST-001', scores=scores,
                   score_statuses={key: 'Scored' for key in scores}, **result)
    return dict(gradebook=dict(quarter=1, subject_name='Custom Subject', school_year_name='2026-2027', grade_name='7', section_name='Jose',
                               teacher_name='Synthetic Teacher', school_head='Synthetic Head', passing_grade=75),
                policy=dict(transmutation_table=table, rounding_method=method, rounding_decimal_places=2, rounding_final_decimal_places=0),
                components=components, students=[student])


class AdaptiveClassRecordTests(unittest.TestCase):
    def test_all_terms_export_requires_three_complete_records_and_preserves_grades(self):
        from fastapi import HTTPException
        from app.services.adaptive_class_record import export_adaptive

        terms = {}
        books = []
        for period in (1, 2, 3):
            full = sample_record(exams=True)
            full['gradebook'].update(id=period, quarter=period, status='Locked' if period == 3 else 'Finalized')
            full['statistics'] = {'incomplete': 0}
            terms[period] = full
            books.append(SimpleNamespace(id=period, quarter=period, teacher_user_id=7, school_year_id=1,
                                         grade_level_id=1, section_id=1, subject_id=1))
        before = copy.deepcopy(terms)
        db = Mock()
        db.scalars.return_value.all.return_value = books
        with tempfile.TemporaryDirectory() as directory, \
                patch('app.services.adaptive_class_record.get_gradebook_full', side_effect=lambda db, key: terms[key]), \
                patch('app.services.reports.EXPORT_DIR', Path(directory)):
            path = export_adaptive(db, books[0], terms[1], 'V', 'V', '301874', all_terms=True)
            self.assertIn('AllTerms', path.name)
            formulas = load_workbook(path)
            values = load_workbook(path, data_only=True)
            try:
                self.assertEqual(values.sheetnames, ['INPUT DATA', 'Term 1', 'Term 2', 'Term 3'])
                for period in (1, 2, 3):
                    sheet = values[f'Term {period}']
                    headings = {cell.value: cell.column for cell in sheet[10]}
                    for heading, key in (('INITIAL GRADE', 'initial_grade'), ('TERM GRADE', 'reported_grade')):
                        self.assertEqual(sheet.cell(12, headings[heading]).value, terms[period]['students'][0][key])
                        self.assertEqual(formulas[sheet.title].cell(12, headings[heading]).data_type, 'f')
            finally:
                formulas.close()
                values.close()
            self.assertEqual(terms, before)
            query = str(db.scalars.call_args.args[0].compile(compile_kwargs={'literal_binds': True}))
            for constraint in ('teacher_user_id = 7', 'school_year_id = 1', 'section_id = 1', 'subject_id = 1', 'quarter IN (1, 2, 3)'):
                self.assertIn(constraint, query)

            for invalid in ('missing', 'draft', 'incomplete', 'empty', 'duplicate'):
                with self.subTest(invalid=invalid):
                    terms[3] = copy.deepcopy(before[3])
                    db.scalars.return_value.all.return_value = books
                    if invalid == 'missing':
                        db.scalars.return_value.all.return_value = books[:2]
                    elif invalid == 'draft':
                        terms[3]['gradebook']['status'] = 'Draft'
                    elif invalid == 'incomplete':
                        terms[3]['statistics']['incomplete'] = 1
                    elif invalid == 'empty':
                        terms[3]['students'] = []
                    else:
                        db.scalars.return_value.all.return_value = books + [books[2]]
                    with self.assertRaises(HTTPException) as error:
                        export_adaptive(db, books[0], terms[1], 'V', 'V', '301874', all_terms=True)
                    self.assertEqual(error.exception.status_code, 409)
                    self.assertEqual(len(list(Path(directory).rglob('*.xlsx'))), 1)

    def test_variable_maxima_components_and_formula_caches_match_engine(self):
        from app.services.reports import _write_class_record_workbook
        for exams in (False, True):
            for method in ('half_up', 'floor', 'ceiling'):
                with self.subTest(exams=exams, method=method), tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1] / 'data') as directory:
                    full, cached, workbook = sample_record(method, exams), {}, Workbook()
                    before = copy.deepcopy(full)
                    sheet = _term_sheet(workbook, full, cached)
                    self.assertEqual(full, before)
                    self.assertEqual([sheet.cell(11, col).value for col in (5, 6)], [30, 20])
                    self.assertNotIn('Unused', [cell.value for cell in sheet[10]])
                    self.assertIn('Attendance', [cell.value for cell in sheet[9]])
                    headings = {cell.value: cell.column for cell in sheet[10]}
                    for name in ('INITIAL GRADE', 'TERM GRADE'):
                        self.assertEqual(sheet.cell(12, headings[name]).data_type, 'f')
                    self.assertTrue(any('SUM(IF(ISNUMBER' in str(cell.value) for cell in sheet[12]))
                    source, target = Path(directory) / 'source.xlsx', Path(directory) / 'export.xlsx'
                    workbook.save(source)
                    _write_class_record_workbook(source, target, {'Term 1': {}}, formula_values=cached)
                    data = load_workbook(target, data_only=True)
                    try:
                        for heading, key in (('INITIAL GRADE', 'initial_grade'), ('TERM GRADE', 'reported_grade')):
                            self.assertEqual(data['Term 1'].cell(12, headings[heading]).value, full['students'][0][key])
                    finally:
                        data.close()

    def test_names_and_activity_labels_are_text_not_formulas(self):
        full = sample_record()
        full['students'][0]['full_name'] = '=1+1'
        full['components'][0]['items'][0]['label'] = '=1+1'
        full['gradebook']['teacher_name'] = '=1+1'
        sheet = _term_sheet(Workbook(), full, {})
        for cell in ('B12', 'E10', 'D5'):
            self.assertEqual(sheet[cell].data_type, 's')
            self.assertEqual(sheet[cell].value, '=1+1')

    def test_exam_conversion_uses_unrounded_percentages_at_grade_boundaries(self):
        full = sample_record(exams=True)
        sheet = _term_sheet(Workbook(), full, {})
        column = next(cell.column for cell in sheet[10] if cell.value == 'TERM GRADE')
        formula = sheet.cell(12, column).value
        self.assertIn('*100000000', formula)
        self.assertNotIn('ROUND(', formula)
        self.assertNotIn('LOOKUP(', formula)

    def test_zero_scores_are_real_scores_and_mismatched_grades_stop_export(self):
        from fastapi import HTTPException
        full = sample_record()
        full['students'][0]['scores']['1'] = 0
        with self.assertRaises(HTTPException):
            _term_sheet(Workbook(), full, {})


if __name__ == '__main__':
    unittest.main()
