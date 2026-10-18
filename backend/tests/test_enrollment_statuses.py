"""Lifecycle checks use temporary SQLite and disposable exports only."""
import io
import unittest
from datetime import date, datetime, time, timedelta
from pathlib import Path
from unittest.mock import patch

import test_attendance_sync as fixtures
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from app.models import AttendanceEvent, Person, SmsOutbox
from app.schemas import PersonCreate
from app.services.attendance import MANILA, raw_row, list_rows, record_gate_match, close_day
from app.services.enrollment import STUDENT_STATUSES, STAFF_STATUSES, DEPARTURE_STATUSES
from app.services.reports import attendance_range_rows
from app.services.roster import parse_roster
from app.services import sf2


class EnrollmentStatusesTest(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.AttendanceSyncTest()
        self.fixture.setUp()
        self.db, self.client = self.fixture.db, self.fixture.client
        self.student, self.staff = self.fixture.student, self.fixture.personnel
        self.day = self.fixture.day

    def tearDown(self):
        self.fixture.tearDown()

    def payload(self, role, status):
        return dict(external_id='TEST-ID', first_name='Test', sex='Male', role=role,
                    grade='10', section='Rizal', enrollment_status=status,
                    enrollment_start_date='2026-10-01', enrollment_end_date=str(self.day) if status in DEPARTURE_STATUSES else None,
                    enrollment_note='Confirmed departure' if status == 'Other' else None)

    def test_role_specific_statuses_and_required_dates(self):
        for role, statuses in [('Student', STUDENT_STATUSES), ('Faculty', STAFF_STATUSES), ('Non-teaching Personnel', STAFF_STATUSES)]:
            for status in statuses:
                with self.subTest(role=role, status=status):
                    PersonCreate(**self.payload(role, status))
                    if status in DEPARTURE_STATUSES:
                        with self.assertRaises(ValidationError):
                            PersonCreate(**{**self.payload(role, status), 'enrollment_end_date': None})
                    if status in {'Transferred In', 'Appointed'}:
                        with self.assertRaises(ValidationError):
                            PersonCreate(**{**self.payload(role, status), 'enrollment_start_date': None})
        for role, status in [('Student', 'Resigned'), ('Faculty', 'Graduated'), ('Non-teaching Personnel', 'Dropped Out')]:
            with self.assertRaises(ValidationError):
                PersonCreate(**self.payload(role, status))
        for changes in [{'enrollment_note': ''}, {'enrollment_start_date': '2026-11-01'}]:
            with self.assertRaises(ValidationError):
                PersonCreate(**{**self.payload('Student', 'Other'), **changes})

    def test_every_departure_blocks_scans_assistance_and_absence_after_last_day(self):
        for role, person, statuses in [('Student', self.student, STUDENT_STATUSES), ('Faculty', self.staff, STAFF_STATUSES), ('Non-teaching Personnel', self.staff, STAFF_STATUSES)]:
            for status in statuses & DEPARTURE_STATUSES:
                with self.subTest(role=role, status=status):
                    person.role, person.enrollment_status = role, status
                    person.enrollment_start_date, person.enrollment_end_date = self.day - timedelta(days=5), self.day
                    self.db.commit()
                    self.assertEqual(raw_row(self.db, person, self.day)['status'], 'No scan')
                    following = self.day + timedelta(days=1)
                    self.assertEqual(raw_row(self.db, person, following)['status'], status)
                    with patch('app.services.attendance.queue_notice') as sms:
                        result = record_gate_match(self.db, person, 20, datetime.combine(following, time(8), MANILA))
                        self.assertFalse(result['recorded'])
                        self.assertEqual(result['status'], status)
                        sms.assert_not_called()
                    assisted = self.client.post('/api/attendance/corrections', headers=self.fixture.admin_headers, json={
                        'person_id': person.id, 'attendance_date': str(following), 'status': 'Present', 'time_in': '08:00',
                        'assisted': True, 'identity_verified': True, 'reason': 'Synthetic assistance check'})
                    self.assertEqual(assisted.status_code, 422, assisted.text)
        close_day(self.db, self.day + timedelta(days=1))
        self.assertEqual(self.db.scalars(select(AttendanceEvent)).all(), [])
        self.assertEqual(self.db.scalars(select(SmsOutbox)).all(), [])

    def test_start_date_and_last_day_boundary_preserve_gate_rules(self):
        self.staff.enrollment_start_date = self.day + timedelta(days=1)
        self.assertEqual(raw_row(self.db, self.staff, self.day)['status'], 'Not Started')
        self.student.enrollment_status = 'Withdrawn'
        self.student.enrollment_end_date = self.day
        self.db.commit()
        result = record_gate_match(self.db, self.student, 20, datetime.combine(self.day, time(8), MANILA))
        self.assertTrue(result['recorded'])
        self.assertEqual(raw_row(self.db, self.student, self.day)['status'], 'Present')

    def test_departed_inactive_people_remain_in_historical_reports(self):
        self.student.enrollment_status = 'Dropped Out'
        self.student.enrollment_end_date = self.day
        self.student.active = False
        self.db.commit()
        rows = list_rows(self.db, self.day + timedelta(days=1))
        self.assertEqual(next(row for row in rows if row['person_id'] == self.student.id)['status'], 'Dropped Out')
        rows = attendance_range_rows(self.db, self.day - timedelta(days=1), self.day + timedelta(days=1), person_id=self.student.id)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[-1]['status'], 'Dropped Out')

    def test_manual_correction_cannot_create_absence_after_departure(self):
        self.student.enrollment_status = 'Withdrawn'
        self.student.enrollment_end_date = self.day
        self.db.commit()
        for status in ('Absent', 'Present', 'Late'):
            response = self.client.post('/api/attendance/corrections', headers=self.fixture.admin_headers, json={
                'person_id': self.student.id, 'attendance_date': str(self.day + timedelta(days=1)),
                'status': status, 'reason': 'Synthetic manual correction check'})
            self.assertEqual(response.status_code, 422, response.text)

    def test_sf2_departure_remarks_and_totals_for_all_student_departures(self):
        self.student.active = False
        self.student.enrollment_start_date = date(2026, 10, 1)
        self.student.enrollment_end_date = self.day
        self.student.enrollment_note = 'Confirmed by school staff'
        self.db.add(AttendanceEvent(id='fixture-absence', person_id=self.student.id, event_date=date(2026, 10, 12),
                                    event_time=time(9), direction='Daily Close', status='Absent', source='Test'))
        self.db.commit()
        template = Path(__file__).resolve().parents[1] / 'data/templates/School Form 2 (SF2) Daily Attendance Report of Learners.xlsx'
        for status in STUDENT_STATUSES & DEPARTURE_STATUSES:
            with self.subTest(status=status), patch.object(sf2, 'template_path', return_value=template):
                self.student.enrollment_status = status
                self.db.commit()
                path = sf2.generate_sf2(self.db, 2026, 10, '10', 'Rizal', '301874', '2026-2027', 'Test school')
                book = load_workbook(path, data_only=True)
                sheet = book[sf2.SHEET_NAME]
                self.assertEqual(sheet['B14'].value, self.student.full_name.upper())
                self.assertIn(status.upper(), sheet['AE14'].value)
                self.assertIn('2026-10-13', sheet['AE14'].value)
                if status != 'Transferred Out':
                    self.assertNotIn('TRANSFERRED OUT', sheet['AE14'].value)
                self.assertEqual(sheet['AC14'].value, 1)
                book.close()

    def test_api_update_preserves_identity_and_existing_attendance(self):
        old_id, external_id = self.student.id, self.student.external_id
        record_gate_match(self.db, self.student, 20, datetime.combine(self.day, time(8), MANILA))
        payload = {**self.payload('Student', 'Graduated'), 'external_id': external_id}
        response = self.client.patch(f'/api/persons/{old_id}', headers=self.fixture.admin_headers, json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['enrollment_status'], 'Graduated')
        self.db.expire_all()
        self.assertEqual(self.db.get(Person, old_id).external_id, external_id)
        self.assertEqual(len(self.db.scalars(select(AttendanceEvent)).all()), 1)

    def test_roster_import_uses_same_role_and_departure_validation(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['Employee ID', 'Name', 'Sex', 'Role', 'Employment Status', 'Enrollment end date'])
        sheet.append(['STAFF-2', 'Resigned Staff', 'Male', 'Faculty', 'Resigned', '2026-10-13'])
        output = io.BytesIO(); book.save(output)
        rows, errors = parse_roster(output.getvalue())
        self.assertFalse(errors)
        self.assertEqual(rows[0]['enrollment_status'], 'Resigned')
        sheet['F2'] = None
        output = io.BytesIO(); book.save(output)
        rows, errors = parse_roster(output.getvalue())
        self.assertTrue(errors)
