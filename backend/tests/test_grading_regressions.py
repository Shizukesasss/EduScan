"""Class-record regressions, isolated from the school's live database."""
import json
import copy
import os
import sys
import tempfile
import unittest
from datetime import date, datetime, time
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
TEST_DIR = tempfile.TemporaryDirectory(
    prefix="eduscan-record-tests-",
    dir=Path(__file__).resolve().parents[1] / "data",
)
os.environ["DATABASE_URL"] = "sqlite://"
os.environ["EDUSCAN_DATA_DIR"] = TEST_DIR.name
os.environ["EDUSCAN_SECRET_KEY"] = "isolated-class-record-regression-tests"
os.environ["ATTENDANCE_SCHEDULER_ENABLED"] = "false"

from app.database import Base  # noqa: E402
from app.models import ClassSchedule, GradeLevel, GradingPeriod, Person, SchoolSection, SchoolYear, Subject, SubjectTeachingAssignment, User  # noqa: E402
from app.models import GradebookState, GradeComponent  # noqa: E402
from app.models_grading import AssessmentItem, Gradebook, GradebookAuditEntry, GradeAdjustmentRequest, GradeCriteriaAdjustmentRequest, GradingPolicy, StudentScore  # noqa: E402
from app.services import grading  # noqa: E402
from app.services.grading_engine import calculate_weighted_grade  # noqa: E402
from app.schemas import SchedulePayload, SectionPayload, SubjectPayload, SubjectTeachingAssignmentPayload  # noqa: E402
from app.main import academic_structure, ensure_report_access, report_parameters_allowed, report_sections, schedules, save_schedule, save_section, save_subject, save_subject_teaching_assignments  # noqa: E402
from app.main import finalize_gradebook_new_endpoint, validate_gradebook_endpoint  # noqa: E402
from app.main import finalize_gradebook as finalize_legacy_gradebook, gradebook_report_xlsx_new  # noqa: E402
from app.schemas import GradebookActionPayload  # noqa: E402
from app.schemas_grading import GradebookWorkflowPayload  # noqa: E402
from app.services.academic import change_deleted_reference  # noqa: E402
from fastapi import HTTPException  # noqa: E402


def definitions(weights=(0.3, 0.5, 0.2), counts=(1, 1, 1)):
    return json.dumps([{'name': name, 'weight': weight, 'max_items': count}
                       for name, weight, count in zip(('Written Works', 'Performance Tasks', 'Quarterly Assessment'), weights, counts)])


class ClassRecordRegressionTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine('sqlite://')
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine, expire_on_commit=False)
        self.admin = User(username='admin', password_hash='unused', full_name='Administrator', role='admin')
        self.teacher = User(username='teacher', password_hash='unused', full_name='Class Adviser', role='teacher')
        self.year = SchoolYear(name='2026-2027', starts_on=date(2026, 6, 1), ends_on=date(2027, 3, 31))
        self.grade = GradeLevel(name='Grade 11')
        self.subject = Subject(code='SHS-CORE', name='Core Subject', category='SHS-Core')
        self.db.add_all([self.admin, self.teacher, self.year, self.grade, self.subject])
        self.db.flush()
        self.period = GradingPeriod(school_year_id=self.year.id, quarter=1, name='Q1', starts_on=date(2026, 6, 1), ends_on=date(2026, 8, 31))
        self.section = SchoolSection(grade_level_id=self.grade.id, name='Berlin', adviser_user_id=self.teacher.id)
        self.general = GradingPolicy(name='General policy', school_year=self.year.name, grade_levels='7,8,9,10,11,12', subject_category='General', component_definitions_json=definitions())
        self.specific = GradingPolicy(name='SHS policy', school_year=self.year.name, grade_levels='11,12', subject_category='SHS-Core', component_definitions_json=definitions((0.25, 0.5, 0.25), (10, 8, 1)), transmutation_table_json=json.dumps([{'min': 0, 'max': 100, 'value': 88}]), rounding_decimal_places=3, rounding_method='floor')
        self.first = Person(external_id='FIRST', full_name='First Student', sex='Male', role='Student', grade='11', section='Berlin')
        self.second = Person(external_id='SECOND', full_name='Second Student', sex='Female', role='Student', grade='grade11', section=' berlin ')
        self.outside = Person(external_id='OUTSIDE', full_name='Other Class', sex='Male', role='Student', grade='12', section='Berlin')
        self.db.add_all([self.period, self.section, self.general, self.specific, self.first, self.second, self.outside])
        self.db.commit()
        from app.services.settings_store import set_json
        set_json(self.db, 'school.heads', {str(self.year.id): {'full_name': 'Buen Owogowog', 'title': 'Principal'}})
        from app.models import SectionAdviserAssignment
        self.db.add(SectionAdviserAssignment(school_year_id=self.year.id, section_id=self.section.id, teacher_user_id=self.teacher.id))
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def create(self, policy=None, **overrides):
        context = dict(school_year_id=self.year.id, grading_period_id=self.period.id, grade_level_id=self.grade.id, section_id=self.section.id, subject_id=self.subject.id, user=self.admin, grading_policy_id=policy.id if policy else None)
        context.update(overrides)
        return grading.create_or_get_gradebook(self.db, **context)

    def test_prefixed_grade_and_trimmed_section_find_correct_students(self):
        full = grading.get_gradebook_full(self.db, self.create().id)
        self.assertEqual(full['gradebook']['grade_name'], '11')
        self.assertEqual({row['person_id'] for row in full['students']}, {self.first.id, self.second.id})

    def test_school_head_access_is_read_only(self):
        from app.main import app, list_gradebooks, get_gradebook_endpoint, report_history
        from app.models import GeneratedReport
        head = User(username='principal', full_name='Buen Owogowog', password_hash='unused', role='school_head', must_change_password=False)
        other_teacher = User(username='otherteacher', full_name='Other Teacher', password_hash='unused', role='teacher')
        other_section = SchoolSection(grade_level_id=self.grade.id, name='Other Section')
        self.db.add_all([head, other_teacher, other_section])
        self.db.commit()
        book = self.create(self.general)
        other_book = self.create(self.general, section_id=other_section.id)
        book.teacher_user_id = self.teacher.id
        other_book.teacher_user_id = other_teacher.id
        for teacher in (self.teacher, other_teacher):
            self.db.add(GeneratedReport(id=f'report-{teacher.id}', report_type='Class record',
                filename=f'class-{teacher.id}.xlsx', file_path='unused-test-path', file_sha256='0' * 64,
                parameters_json='{}', generated_by=teacher.id, generated_by_name=teacher.full_name))
        self.db.commit()
        readable = [('GET', '/api/gradebooks'), ('GET', '/api/gradebooks/{gradebook_id}'),
                    ('GET', '/api/reports/history'), ('GET', '/api/reports/history/{report_id}/download'),
                    ('GET', '/api/gradebooks/{gradebook_id}/report.xlsx'),
                    ('GET', '/api/gradebooks/{gradebook_id}/audit')]
        blocked = [('PUT', '/api/gradebooks/{gradebook_id}/scores'),
                   ('PUT', '/api/gradebooks/{gradebook_id}/components'),
                   ('POST', '/api/gradebooks'),
                   ('POST', '/api/gradebooks/{gradebook_id}/finalize'),
                   ('POST', '/api/gradebooks/{gradebook_id}/reopen'),
                   ('POST', '/api/gradebooks/{gradebook_id}/lock'),
                   ('PUT', '/api/adjustment-requests/{request_id}'),
                   ('DELETE', '/api/adjustment-requests/{request_id}'),
                   ('POST', '/api/gradebooks/{gradebook_id}/adjustments'),
                   ('POST', '/api/gradebooks/{gradebook_id}/criteria-adjustments'),
                   ('POST', '/api/reports/history/{report_id}/review'),
                   ('PUT', '/api/admin/school-heads')]
        for method, path in readable + blocked:
            route = next(route for route in app.routes if getattr(route, 'path', None) == path and method in getattr(route, 'methods', set()))
            authorization = next(dependency.call for dependency in route.dependant.dependencies if dependency.name in {'user', 'actor', '_'})
            if (method, path) in readable:
                self.assertIs(authorization(user=head), head)
            else:
                with self.assertRaises(HTTPException) as denied:
                    authorization(user=head)
                self.assertEqual(denied.exception.status_code, 403)
        self.assertEqual({book.id, other_book.id}, {row['id'] for row in list_gradebooks(db=self.db, user=head)})
        for record in (book, other_book):
            self.assertEqual(get_gradebook_endpoint(record.id, db=self.db, user=head)['gradebook']['id'], record.id)
        self.assertEqual({row['id'] for row in report_history(db=self.db, user=head)},
                         {f'report-{self.teacher.id}', f'report-{other_teacher.id}'})

    def test_school_head_settings_are_saved_per_school_year(self):
        from app.main import save_school_head, school_heads
        from app.schemas import SchoolHeadPayload
        from app.services.settings_store import set_json
        set_json(self.db, 'school.heads', {'9999': {'full_name': 'Previous Principal', 'title': 'Principal'}})
        save_school_head(SchoolHeadPayload(school_year_id=self.year.id, full_name='Buen Owogowog', title='Principal'), db=self.db, user=self.admin)
        saved = school_heads(db=self.db, user=self.admin)
        self.assertEqual(saved[str(self.year.id)]['full_name'], 'Buen Owogowog')
        self.assertEqual(saved['9999']['full_name'], 'Previous Principal')

    def test_admin_reads_each_teachers_saved_record_without_separate_grade_calculation(self):
        from app.main import get_gradebook_endpoint, ensure_gradebook_access
        other_teacher = User(username='parity-teacher', full_name='Other Teacher', password_hash='unused', role='teacher')
        other_section = SchoolSection(grade_level_id=self.grade.id, name='Parity Section')
        other_student = Person(external_id='PARITY', full_name='Other Learner', sex='Female', role='Student', grade='11', section='Parity Section')
        self.db.add_all([other_teacher, other_section, other_student])
        self.db.commit()
        books = [self.create(self.general), self.create(self.general, section_id=other_section.id)]
        for book, teacher, learner, percentage in zip(books, (self.teacher, other_teacher), (self.first, other_student), (80, 65)):
            book.teacher_user_id, book.teacher_name = teacher.id, teacher.full_name
            self.db.add(SubjectTeachingAssignment(school_year_id=self.year.id, section_id=book.section_id,
                subject_id=self.subject.id, teacher_user_id=teacher.id, active=True))
            self.db.commit()
            entries = [{'assessment_item_id': item.id, 'score': item.max_score * percentage / 100, 'status': 'Scored'}
                       for component in book.components for item in component.items]
            grading.save_gradebook_scores(self.db, book.id, teacher, {str(learner.id): entries}, 'Isolated saved-grade parity check')
            teacher_view = get_gradebook_endpoint(book.id, db=self.db, user=teacher)
            admin_view = get_gradebook_endpoint(book.id, db=self.db, user=self.admin)
            self.assertEqual(admin_view, teacher_view)
            entries[0]['score'] = 0
            grading.save_gradebook_scores(self.db, book.id, teacher, {str(learner.id): entries}, 'Isolated saved-grade update check')
            refreshed = get_gradebook_endpoint(book.id, db=self.db, user=self.admin)
            self.assertEqual(refreshed, get_gradebook_endpoint(book.id, db=self.db, user=teacher))
            self.assertNotEqual(refreshed['students'], admin_view['students'])
        with self.assertRaises(HTTPException):
            ensure_gradebook_access(self.db, self.teacher, books[1])

    def test_specific_subject_policy_wins_over_general(self):
        self.assertEqual(self.create().grading_policy_id, self.specific.id)
        self.assertIsNone(grading.get_active_policy_for(self.db, self.year.name, '6', 'SHS-Core'))

    def test_fractional_policy_weights_and_item_limits(self):
        book = self.create()
        self.assertEqual([component.weight for component in book.components], [25, 50, 25])
        self.assertEqual([len(component.items) for component in book.components], [10, 8, 1])
        self.assertEqual([component.component_type for component in book.components], ['Written Work', 'Performance Task', 'Quarterly Assessment'])

    def test_existing_fractional_weights_are_read_without_changing_storage(self):
        book = self.create(self.general)
        for component in book.components:
            component.weight /= 100
        self.db.commit()
        full = grading.get_gradebook_full(self.db, book.id)
        self.assertEqual([component['weight'] for component in full['components']], [30, 50, 20])
        self.assertEqual([component.weight for component in book.components], [0.3, 0.5, 0.2])

    def test_browser_receives_saved_transmutation_and_rounding_rules(self):
        full = grading.get_gradebook_full(self.db, self.create().id)
        self.assertEqual(full['policy']['transmutation_table'][0]['value'], 88)
        self.assertEqual(full['policy']['rounding_decimal_places'], 3)
        self.assertEqual(full['policy']['rounding_method'], 'floor')

    def test_gradebook_mutations_use_a_postgresql_row_lock(self):
        from sqlalchemy.dialects import postgresql

        class CapturingSession:
            statement = None

            def scalar(self, statement):
                self.statement = statement
                return None

        session = CapturingSession()
        grading._gradebook_for_update(session, 1)
        self.assertIn('FOR UPDATE', str(session.statement.compile(dialect=postgresql.dialect())))

    def test_incomplete_grade_stays_blank_and_weights_do_not_expand(self):
        grade = calculate_weighted_grade([{'name': 'WW', 'weight': 30, 'percentage': 90, 'complete': True}, {'name': 'PT', 'weight': 50, 'percentage': None, 'complete': False}, {'name': 'QA', 'weight': 20, 'percentage': None, 'complete': False}])
        self.assertIsNone(grade['initial_grade'])
        self.assertEqual(grade['component_grades'][0]['weighted_score'], 27)

    def test_complete_fractional_weights_compute_full_grade(self):
        grade = calculate_weighted_grade([{'name': name, 'weight': weight, 'percentage': 90, 'complete': True} for name, weight in zip(('WW', 'PT', 'QA'), (0.3, 0.5, 0.2))])
        self.assertEqual(grade['initial_grade'], 90)

    def test_only_class_used_activities_are_required_and_each_learner_must_complete_them(self):
        book = self.create(self.general)
        self.db.add(SubjectTeachingAssignment(school_year_id=self.year.id, subject_id=self.subject.id,
                                            section_id=self.section.id, teacher_user_id=self.teacher.id))
        ww, pt, qa = book.components
        ww_second = AssessmentItem(label='Written Work 2', max_score=20, sequence=1)
        pt_second = AssessmentItem(label='Performance Task 2', max_score=20, sequence=1)
        unused_ww = AssessmentItem(label='Unused Written Work', max_score=20, sequence=2)
        ww.items.extend((ww_second, unused_ww))
        pt.items.append(pt_second)
        self.db.commit()
        active_items = [ww.items[0], ww_second, pt.items[0], pt_second, qa.items[0]]
        first_entries = [{'assessment_item_id': item.id, 'score': item.max_score, 'status': 'Scored'}
                         for item in active_items]
        second_entries = [{'assessment_item_id': item.id, 'score': item.max_score, 'status': 'Scored'}
                          for item in active_items if item is not ww_second]
        grading.save_gradebook_scores(self.db, book.id, self.teacher,
                                      {str(self.first.id): first_entries, str(self.second.id): second_entries},
                                      'Save partial class scores')
        full = grading.get_gradebook_full(self.db, book.id)
        self.assertEqual(full['statistics']['incomplete'], 1)
        self.assertTrue(full['students'][0]['complete'])
        self.assertIsNotNone(full['students'][0]['initial_grade'])
        self.assertFalse(full['students'][1]['complete'])
        self.assertIsNone(full['students'][1]['initial_grade'])
        self.assertFalse(next(item for item in full['components'][0]['items'] if item['id'] == unused_ww.id)['active'])
        validation = validate_gradebook_endpoint(book.id, self.db, self.teacher)
        self.assertFalse(validation['valid'])
        incomplete_issue = next(issue for issue in validation['issues'] if issue['code'] == 'INCOMPLETE_LEARNERS')
        self.assertIn('Second Student', incomplete_issue['message'])
        self.assertIn('Written Work 2', incomplete_issue['message'])
        before = len(self.db.scalars(select(GradebookAuditEntry)).all())
        with self.assertRaises(HTTPException) as blocked:
            finalize_gradebook_new_endpoint(book.id, GradebookWorkflowPayload(reason='Finalize term record'), self.db, self.teacher)
        self.assertEqual(blocked.exception.status_code, 409)
        self.assertEqual(blocked.exception.detail['code'], 'INCOMPLETE_GRADEBOOK')
        self.assertIn('Second Student', blocked.exception.detail['message'])
        self.assertIn('Written Work 2', blocked.exception.detail['message'])
        self.assertEqual(book.status, 'Draft')
        self.assertIsNone(book.finalized_at)
        self.assertEqual(len(self.db.scalars(select(GradebookAuditEntry)).all()), before)
        grading.save_gradebook_scores(self.db, book.id, self.teacher,
                                      {str(self.second.id): [{'assessment_item_id': ww_second.id, 'score': 0, 'status': 'Scored'}]},
                                      'Complete the class-used activities')
        self.assertTrue(validate_gradebook_endpoint(book.id, self.db, self.teacher)['valid'])
        result = finalize_gradebook_new_endpoint(book.id, GradebookWorkflowPayload(reason='Finalize complete record'), self.db, self.teacher)
        self.assertTrue(result['finalized'])
        with self.assertRaisesRegex(ValueError, 'Cannot edit scores'):
            grading.save_gradebook_scores(self.db, book.id, self.teacher, {}, 'Attempt a locked score edit')

    def test_finalize_error_lists_missing_exam_scores_for_each_learner(self):
        book = self.create(self.general)
        self.db.add(SubjectTeachingAssignment(school_year_id=self.year.id, subject_id=self.subject.id,
                                            section_id=self.section.id, teacher_user_id=self.teacher.id))
        exam_component = book.components[-1]
        exam_component.items[0].label = 'ST1'
        exam_component.items.extend((AssessmentItem(label='ST2', max_score=50, sequence=1),
                                     AssessmentItem(label='TE', max_score=50, sequence=2)))
        exam_component.component_type = 'Term Examinations'
        self.db.commit()
        validation = validate_gradebook_endpoint(book.id, self.db, self.teacher)
        message = next(issue['message'] for issue in validation['issues'] if issue['code'] == 'INCOMPLETE_LEARNERS')
        for learner_name in ('First Student', 'Second Student'):
            self.assertIn(learner_name, message)
        for label in ('ST1', 'ST2', 'TE'):
            self.assertEqual(message.count(f'{label} (score missing)'), 2, message)
        before = self.db.scalars(select(StudentScore)).all()
        with self.assertRaises(HTTPException) as blocked:
            finalize_gradebook_new_endpoint(book.id, GradebookWorkflowPayload(reason='Check missing exam scores'), self.db, self.teacher)
        self.assertEqual(blocked.exception.detail['code'], 'INCOMPLETE_GRADEBOOK')
        for label in ('ST1', 'ST2', 'TE'):
            self.assertEqual(blocked.exception.detail['message'].count(f'{label} (score missing)'), 2)
        self.assertEqual(self.db.scalars(select(StudentScore)).all(), before)
        self.assertEqual(book.status, 'Draft')

    def test_every_unresolved_score_status_blocks_finalization(self):
        book = self.create(self.general)
        entries = [{'assessment_item_id': item.id, 'score': 0, 'status': 'Scored'}
                   for component in book.components for item in component.items]
        grading.save_gradebook_scores(self.db, book.id, self.teacher,
                                      {str(student.id): entries for student in (self.first, self.second)},
                                      'Record complete zero scores')
        score = self.db.scalar(select(StudentScore).where(StudentScore.person_id == self.first.id))
        for status in ('Missing', 'Absent', 'Excused', 'Incomplete', 'Not Applicable', 'Scored'):
            with self.subTest(status=status):
                score.status, score.score = status, None
                self.db.commit()
                full = grading.get_gradebook_full(self.db, book.id)
                self.assertEqual(full['statistics']['incomplete'], 1)
                with self.assertRaisesRegex(ValueError, 'missing or incomplete'):
                    grading.finalize_gradebook(self.db, book.id, self.teacher, 'Attempt incomplete finalization')
                self.assertEqual(book.status, 'Draft')

    def test_official_download_cannot_bypass_finalization_or_score_completion(self):
        book = self.create(self.general)
        with patch('app.main.generate_gradebook_report_xlsx') as export:
            with self.assertRaises(HTTPException) as draft:
                gradebook_report_xlsx_new(book.id, db=self.db, user=self.admin)
            self.assertEqual(draft.exception.status_code, 409)
            for status in ('Finalized', 'Locked'):
                book.status = status
                self.db.commit()
                with self.assertRaises(HTTPException) as incomplete:
                    gradebook_report_xlsx_new(book.id, db=self.db, user=self.admin)
                self.assertEqual(incomplete.exception.status_code, 409)
                self.assertIn('incomplete', incomplete.exception.detail)
            export.assert_not_called()

    def test_legacy_blank_manual_grade_cannot_turn_into_a_passing_finalized_record(self):
        self.grade.name = '11'
        state = GradebookState(class_key='legacy-incomplete', school_year=self.year.name, quarter=1,
                              subject='Core Subject', grade='11', section='Berlin')
        manual = GradeComponent(class_key=state.class_key, category='Manual Overall', label='Overall Grade',
                                weight=100, max_score=100, sequence=0)
        self.db.add_all([state, manual])
        self.db.commit()
        summary = grading.gradebook_summary(self.db, state.class_key)
        self.assertEqual(summary['statistics']['incomplete'], 1)
        self.assertIsNone(summary['rows'][0]['transmuted_grade'])
        with self.assertRaises(HTTPException) as blocked:
            finalize_legacy_gradebook(state.class_key, GradebookActionPayload(reason='Finalize missing manual grade'), self.db, self.teacher)
        self.assertEqual(blocked.exception.status_code, 409)
        self.assertEqual(state.status, 'Draft')

    def test_legacy_empty_class_cannot_be_finalized(self):
        self.grade.name = '11'
        state = GradebookState(class_key='legacy-empty', school_year=self.year.name, quarter=1,
                              subject='Core Subject', grade='11', section='Berlin')
        self.db.add(state)
        self.db.commit()
        with self.assertRaises(HTTPException) as blocked:
            finalize_legacy_gradebook(state.class_key, GradebookActionPayload(reason='Finalize empty legacy record'), self.db, self.teacher)
        self.assertEqual(blocked.exception.status_code, 409)
        self.assertEqual(state.status, 'Draft')

    def test_invalid_weights_do_not_generate_final_grade(self):
        result = calculate_weighted_grade([{'name': 'WW', 'weight': 99, 'percentage': 90, 'complete': True}])
        self.assertIsNone(result['initial_grade'])

    def test_wrong_class_student_cannot_receive_scores(self):
        book = self.create(self.general)
        item = book.components[0].items[0]
        with self.assertRaisesRegex(ValueError, 'active students'):
            grading.save_gradebook_scores(self.db, book.id, self.admin, {str(self.outside.id): [{'assessment_item_id': item.id, 'score': 10, 'status': 'Scored'}]}, 'Invalid class')
        self.assertEqual(self.db.scalars(select(StudentScore)).all(), [])

    def test_wrong_grade_section_context_cannot_create_record(self):
        grade = GradeLevel(name='12')
        self.db.add(grade)
        self.db.commit()
        with self.assertRaisesRegex(ValueError, 'section does not belong'):
            self.create(grade_level_id=grade.id)
        self.assertEqual(self.db.scalars(select(Gradebook)).all(), [])

    def test_jhs_subject_cannot_create_shs_record(self):
        subject = Subject(code='JHS-MATH', name='Mathematics', category='JHS')
        self.db.add(subject)
        self.db.commit()
        with self.assertRaisesRegex(ValueError, 'subject does not belong'):
            self.create(subject_id=subject.id)

    def test_new_record_shows_assigned_teacher_instead_of_admin(self):
        book = self.create()
        self.assertEqual(book.teacher_user_id, self.teacher.id)
        self.assertEqual(book.teacher_name, self.teacher.full_name)

    def test_one_teacher_can_handle_a_subject_in_multiple_sections(self):
        first = self.create(self.general)
        second_section = SchoolSection(grade_level_id=self.grade.id, name="Paris")
        self.db.add(second_section)
        self.db.commit()
        self.db.add(SubjectTeachingAssignment(
            school_year_id=self.year.id, subject_id=self.subject.id,
            section_id=second_section.id, teacher_user_id=self.teacher.id,
        ))
        self.db.commit()
        second = self.create(self.general, section_id=second_section.id, user=self.teacher)
        self.assertEqual(first.teacher_user_id, self.teacher.id)
        self.assertEqual(second.teacher_user_id, self.teacher.id)
        self.assertEqual(
            grading.teacher_grading_section_ids(self.db, self.teacher),
            {self.section.id, second_section.id},
        )

    def test_subject_teacher_can_teach_outside_advisory_section_but_only_assigned_subject(self):
        other_section = SchoolSection(grade_level_id=self.grade.id, name="Oslo")
        self.db.add(other_section)
        self.db.flush()
        self.db.add(SubjectTeachingAssignment(
            school_year_id=self.year.id, subject_id=self.subject.id,
            section_id=other_section.id, teacher_user_id=self.teacher.id,
        ))
        self.db.commit()

        book = self.create(self.general, section_id=other_section.id, user=self.teacher)
        self.assertEqual(book.teacher_user_id, self.teacher.id)
        self.assertTrue(grading.teacher_can_access_gradebook_context(
            self.db, self.teacher, self.year.id, other_section.id, self.subject.id,
        ))
        self.assertFalse(grading.teacher_can_access_gradebook_context(
            self.db, self.teacher, self.year.id, other_section.id, 99999,
        ))
        self.assertFalse(grading.teacher_can_access_gradebook_context(
            self.db, self.teacher, self.year.id, self.section.id, self.subject.id,
        ))
        another_subject = Subject(code="SHS-OTHER", name="Other Subject", category="SHS")
        self.db.add(another_subject)
        self.db.commit()
        self.assertFalse(grading.teacher_can_access_gradebook_context(
            self.db, self.teacher, self.year.id, self.section.id, another_subject.id,
        ))
        self.assertIn(other_section.id, grading.teacher_grading_section_ids(self.db, self.teacher))

    def test_teacher_with_conflicting_subject_assignments_cannot_open_either_gradebook(self):
        second_subject = Subject(code="SHS-SECOND", name="Second Subject", category="SHS")
        self.db.add(second_subject)
        self.db.flush()
        self.db.add_all([
            SubjectTeachingAssignment(
                school_year_id=self.year.id, subject_id=self.subject.id,
                section_id=self.section.id, teacher_user_id=self.teacher.id,
            ),
            SubjectTeachingAssignment(
                school_year_id=self.year.id, subject_id=second_subject.id,
                section_id=self.section.id, teacher_user_id=self.teacher.id,
            ),
        ])
        self.db.commit()

        self.assertFalse(grading.teacher_can_access_gradebook_context(
            self.db, self.teacher, self.year.id, self.section.id, self.subject.id,
        ))
        self.assertFalse(grading.teacher_can_access_gradebook_context(
            self.db, self.teacher, self.year.id, self.section.id, second_subject.id,
        ))

    def test_teacher_report_sections_follow_advisory_and_selected_school_year_dates(self):
        teaching_section = SchoolSection(grade_level_id=self.grade.id, name="Teaching only")
        self.db.add(teaching_section)
        self.db.flush()
        self.db.add(SubjectTeachingAssignment(
            school_year_id=self.year.id, subject_id=self.subject.id,
            section_id=teaching_section.id, teacher_user_id=self.teacher.id,
        ))
        self.db.commit()

        current = report_sections(self.db, self.teacher, date(2026, 10, 1), date(2026, 10, 31))
        outside = report_sections(self.db, self.teacher, date(2025, 10, 1), date(2025, 10, 31))
        self.assertEqual({item.id for item in current}, {self.section.id, teaching_section.id})
        self.assertEqual({item.id for item in outside}, set())
        ensure_report_access(
            self.db, self.teacher, "Grade 11", "Teaching only",
            date(2026, 10, 1), date(2026, 10, 31),
        )
        with self.assertRaises(HTTPException):
            ensure_report_access(
                self.db, self.teacher, "Grade 11", "Teaching only",
                date(2025, 10, 1), date(2025, 10, 31),
            )
        with self.assertRaises(HTTPException):
            ensure_report_access(
                self.db, self.teacher, "Grade 11", self.section.name,
                date(2025, 10, 1), date(2025, 10, 31),
            )
        self.assertTrue(report_parameters_allowed(self.db, self.teacher, {
            "grade": "Grade 11", "section": "Teaching only",
            "starts_on": "2026-10-01", "ends_on": "2026-10-31",
        }))

    def test_teacher_schedule_list_contains_only_their_own_schedules(self):
        self.db.add_all([
            ClassSchedule(
                grade="11", section="Berlin", subject="Core Subject",
                teacher_name=self.teacher.full_name, weekdays="0,1,2,3,4",
                start_time=time(7, 30), end_time=time(8, 30),
            ),
            ClassSchedule(
                grade="11", section="Berlin", subject="Other Subject",
                teacher_name="Another Teacher", weekdays="0,1,2,3,4",
                start_time=time(8, 30), end_time=time(9, 30),
            ),
        ])
        self.db.commit()

        rows = schedules(db=self.db, user=self.teacher)
        self.assertEqual([item["teacher_name"] for item in rows], [self.teacher.full_name])

    def test_teacher_schedule_requires_exact_active_subject_assignment(self):
        payload = SchedulePayload(
            grade=self.grade.name, section=self.section.name, subject=self.subject.name,
            teacher_name=self.teacher.full_name, start_time=time(7, 30), end_time=time(8, 30),
        )
        with self.assertRaises(HTTPException) as error:
            save_schedule(payload, self.db, self.teacher)
        self.assertEqual(error.exception.status_code, 403)

        self.db.add(SubjectTeachingAssignment(
            school_year_id=self.year.id, subject_id=self.subject.id,
            section_id=self.section.id, teacher_user_id=self.teacher.id,
        ))
        self.db.commit()
        result = save_schedule(payload, self.db, self.admin)
        self.assertIsNotNone(self.db.get(ClassSchedule, result["id"]))

        other_subject = Subject(code="SHS-OTHER", name="Other Subject", category="SHS-Core")
        self.db.add(other_subject)
        self.db.commit()
        payload.subject = other_subject.name
        with self.assertRaises(HTTPException) as error:
            save_schedule(payload, self.db, self.teacher)
        self.assertEqual(error.exception.status_code, 403)

    def test_admin_can_schedule_teacher_in_non_advisory_section(self):
        teaching_section = SchoolSection(grade_level_id=self.grade.id, name="Jose")
        self.db.add(teaching_section)
        self.db.flush()
        self.db.add(SubjectTeachingAssignment(
            school_year_id=self.year.id, subject_id=self.subject.id,
            section_id=teaching_section.id, teacher_user_id=self.teacher.id,
        ))
        self.db.commit()
        payload = SchedulePayload(
            grade=self.grade.name, section=teaching_section.name, subject=self.subject.name,
            teacher_name=self.teacher.full_name, start_time=time(7, 30), end_time=time(8, 30),
        )

        result = save_schedule(payload, self.db, self.admin)

        self.assertEqual(self.db.get(ClassSchedule, result["id"]).section, "Jose")

    def test_shs_elective_is_restricted_to_its_section_track(self):
        elective = Subject(code="STEM-ELECTIVE", name="STEM Elective", category="SHS",
                           is_elective=True, elective_track="STEM")
        self.section.track = "ABM"
        self.db.add(elective)
        self.db.commit()
        with self.assertRaisesRegex(ValueError, "only to the STEM track"):
            self.create(self.general, subject_id=elective.id)

        self.section.track = "STEM"
        self.db.commit()
        self.assertEqual(self.create(self.general, subject_id=elective.id).subject_id, elective.id)

    def test_admin_can_mark_jhs_elective_and_shs_electives_require_a_track(self):
        jhs_elective = save_subject(SubjectPayload(
            code="JHS-CHOICE", name="JHS Choice", category="JHS", is_elective=True,
        ), self.db, self.admin)
        stored_jhs_elective = self.db.get(Subject, jhs_elective["id"])
        self.assertTrue(stored_jhs_elective.is_elective)
        self.assertEqual(stored_jhs_elective.elective_track, "STE")
        with self.assertRaises(HTTPException) as error:
            save_subject(SubjectPayload(
                code="SHS-CHOICE", name="SHS Choice", category="SHS", is_elective=True,
            ), self.db, self.admin)
        self.assertEqual(error.exception.status_code, 422)

    def test_jhs_electives_are_assignable_to_ste_but_not_bep_sections(self):
        grade = GradeLevel(name="Grade 9")
        self.db.add(grade)
        self.db.flush()
        bep = save_section(SectionPayload(
            grade_level_id=grade.id, name="BEP Section", adviser_user_id="", track="bep",
        ), self.db, self.admin)
        ste = save_section(SectionPayload(
            grade_level_id=grade.id, name="STE Section", adviser_user_id="", track="ste",
        ), self.db, self.admin)
        elective = Subject(code="JHS-STE-CHOICE", name="JHS STE Choice", category="JHS", is_elective=True)
        self.db.add(elective)
        self.db.commit()
        with self.assertRaisesRegex(HTTPException, "only available to the STE program"):
            save_subject_teaching_assignments(SubjectTeachingAssignmentPayload(
                school_year_id=self.year.id, subject_id=elective.id,
                teacher_user_id=self.teacher.id, section_ids=[bep["id"]],
            ), self.db, self.admin)
        assigned = save_subject_teaching_assignments(SubjectTeachingAssignmentPayload(
            school_year_id=self.year.id, subject_id=elective.id,
            teacher_user_id=self.teacher.id, section_ids=[ste["id"]],
        ), self.db, self.admin)
        self.assertEqual(assigned["section_count"], 1)
        self.assertEqual(self.db.get(SchoolSection, bep["id"]).track, "BEP")
        self.assertEqual(self.db.get(SchoolSection, ste["id"]).track, "STE")
        with self.assertRaisesRegex(ValueError, "STE program"):
            self.create(grade_level_id=grade.id, section_id=bep["id"], subject_id=elective.id)
        self.assertEqual(self.create(grade_level_id=grade.id, section_id=ste["id"], subject_id=elective.id).subject_id, elective.id)
        structure = academic_structure(db=self.db, user=self.admin)
        visible_elective = next(item for item in structure["subjects"] if item["id"] == elective.id)
        self.assertEqual(visible_elective["elective_track"], "STE")

    def test_blank_adviser_selection_is_saved_as_unassigned_and_linked_name_is_visible(self):
        payload = SectionPayload(grade_level_id=self.grade.id, name="No Adviser", adviser_user_id="")
        result = save_section(payload, self.db, self.admin)
        self.assertIsNone(self.db.get(SchoolSection, result["id"]).adviser_user_id)

        self.section.adviser_name = None
        self.db.commit()
        structure = academic_structure(db=self.db, user=self.admin)
        visible_section = next(item for item in structure["sections"] if item["id"] == self.section.id)
        self.assertEqual(visible_section["adviser_name"], self.teacher.full_name)

    def test_teacher_cannot_be_adviser_of_two_active_sections(self):
        other_section = SchoolSection(grade_level_id=self.grade.id, name="Oslo")
        self.db.add(other_section)
        self.db.commit()
        with self.assertRaises(HTTPException) as error:
            save_section(SectionPayload(
                id=other_section.id, grade_level_id=self.grade.id, name=other_section.name,
                adviser_user_id=self.teacher.id,
            ), self.db, self.admin)
        self.assertEqual(error.exception.status_code, 409)

    def test_restoring_section_preserves_annual_adviser_history(self):
        from app.models import SectionAdviserAssignment
        self.section.deleted_at = datetime.utcnow()
        self.db.commit()
        self.assertTrue(change_deleted_reference(self.db, "sections", self.section.id, self.admin, restore=True)["restored"])
        assignment = self.db.scalar(select(SectionAdviserAssignment).where(
            SectionAdviserAssignment.school_year_id == self.year.id,
            SectionAdviserAssignment.section_id == self.section.id))
        self.assertEqual(assignment.teacher_user_id, self.teacher.id)

    def test_admin_can_assign_same_subject_teacher_to_multiple_sections(self):
        other_section = SchoolSection(grade_level_id=self.grade.id, name="Oslo")
        self.db.add(other_section)
        self.db.commit()
        result = save_subject_teaching_assignments(SubjectTeachingAssignmentPayload(
            school_year_id=self.year.id, subject_id=self.subject.id,
            teacher_user_id=self.teacher.id, section_ids=[self.section.id, other_section.id],
        ), self.db, self.admin)
        self.assertEqual(result["section_count"], 2)
        assignments = self.db.scalars(select(SubjectTeachingAssignment).where(
            SubjectTeachingAssignment.school_year_id == self.year.id,
            SubjectTeachingAssignment.subject_id == self.subject.id,
            SubjectTeachingAssignment.active.is_(True),
        )).all()
        self.assertEqual({item.section_id for item in assignments}, {self.section.id, other_section.id})

    def test_teacher_can_only_have_one_subject_per_year_and_sections_must_share_a_grade(self):
        other_subject = Subject(code="SHS-SECOND", name="Second Subject", category="SHS")
        grade_twelve = GradeLevel(name="Grade 12")
        grade_twelve_section = SchoolSection(grade_level_id=grade_twelve.id, name="Other Grade")
        same_grade_section = SchoolSection(grade_level_id=self.grade.id, name="Second Section")
        self.db.add_all([other_subject, grade_twelve, same_grade_section])
        self.db.flush()
        grade_twelve_section.grade_level_id = grade_twelve.id
        self.db.add(grade_twelve_section)
        self.db.commit()

        with self.assertRaisesRegex(HTTPException, "one grade level"):
            save_subject_teaching_assignments(SubjectTeachingAssignmentPayload(
                school_year_id=self.year.id, subject_id=self.subject.id,
                teacher_user_id=self.teacher.id, section_ids=[self.section.id, grade_twelve_section.id],
            ), self.db, self.admin)

        save_subject_teaching_assignments(SubjectTeachingAssignmentPayload(
            school_year_id=self.year.id, subject_id=self.subject.id,
            teacher_user_id=self.teacher.id, section_ids=[self.section.id, same_grade_section.id],
        ), self.db, self.admin)
        with self.assertRaisesRegex(HTTPException, "already assigned to Core Subject"):
            save_subject_teaching_assignments(SubjectTeachingAssignmentPayload(
                school_year_id=self.year.id, subject_id=other_subject.id,
                teacher_user_id=self.teacher.id, section_ids=[self.section.id],
            ), self.db, self.admin)

        # Existing assignments can be explicitly removed to resolve a conflict.
        save_subject_teaching_assignments(SubjectTeachingAssignmentPayload(
            school_year_id=self.year.id, subject_id=self.subject.id,
            teacher_user_id=self.teacher.id, section_ids=[],
        ), self.db, self.admin)
        result = save_subject_teaching_assignments(SubjectTeachingAssignmentPayload(
            school_year_id=self.year.id, subject_id=other_subject.id,
            teacher_user_id=self.teacher.id, section_ids=[self.section.id],
        ), self.db, self.admin)
        self.assertEqual(result["section_count"], 1)

    def test_subject_can_have_different_teachers_per_section_without_overwriting_assignments(self):
        another_teacher = User(username="another-subject-teacher", password_hash="unused",
                               full_name="Another Subject Teacher", role="teacher")
        other_section = SchoolSection(grade_level_id=self.grade.id, name="Oslo")
        self.db.add_all([another_teacher, other_section])
        self.db.commit()

        save_subject_teaching_assignments(SubjectTeachingAssignmentPayload(
            school_year_id=self.year.id, subject_id=self.subject.id,
            teacher_user_id=self.teacher.id, section_ids=[self.section.id],
        ), self.db, self.admin)
        save_subject_teaching_assignments(SubjectTeachingAssignmentPayload(
            school_year_id=self.year.id, subject_id=self.subject.id,
            teacher_user_id=another_teacher.id, section_ids=[other_section.id],
        ), self.db, self.admin)

        def active_teacher_by_section():
            return {item.section_id: item.teacher_user_id for item in self.db.scalars(select(SubjectTeachingAssignment).where(
                SubjectTeachingAssignment.school_year_id == self.year.id,
                SubjectTeachingAssignment.subject_id == self.subject.id,
                SubjectTeachingAssignment.active.is_(True),
            )).all()}

        self.assertEqual(active_teacher_by_section(), {
            self.section.id: self.teacher.id,
            other_section.id: another_teacher.id,
        })
        self.assertTrue(grading.teacher_can_access_gradebook_context(
            self.db, self.teacher, self.year.id, self.section.id, self.subject.id,
        ))
        self.assertTrue(grading.teacher_can_access_gradebook_context(
            self.db, another_teacher, self.year.id, other_section.id, self.subject.id,
        ))
        self.assertFalse(grading.teacher_can_access_gradebook_context(
            self.db, self.teacher, self.year.id, other_section.id, self.subject.id,
        ))

        save_subject_teaching_assignments(SubjectTeachingAssignmentPayload(
            school_year_id=self.year.id, subject_id=self.subject.id,
            teacher_user_id=self.teacher.id, section_ids=[],
        ), self.db, self.admin)
        self.assertEqual(active_teacher_by_section(), {other_section.id: another_teacher.id})

    def test_teacher_sees_only_owned_sections_but_knows_subject_is_assigned_elsewhere(self):
        other_teacher = User(username="other-teacher", password_hash="unused",
                              full_name="Other Teacher", role="teacher")
        private_section = SchoolSection(grade_level_id=self.grade.id, name="Private")
        self.db.add_all([other_teacher, private_section])
        self.db.flush()
        self.db.add(SubjectTeachingAssignment(
            school_year_id=self.year.id, subject_id=self.subject.id,
            section_id=private_section.id, teacher_user_id=other_teacher.id,
        ))
        self.db.commit()

        structure = academic_structure("grading", self.db, self.teacher)
        self.assertEqual({item["id"] for item in structure["sections"]}, {self.section.id})
        self.assertTrue(all(item["section_id"] != private_section.id for item in structure["subject_assignments"]))
        self.assertIn((self.year.id, self.subject.id), structure["subject_assignment_subjects"])
        self.assertFalse(grading.teacher_can_access_gradebook_context(
            self.db, self.teacher, self.year.id, self.section.id, self.subject.id,
        ))

    def test_unassigned_legacy_gradebook_is_not_accessible_to_teacher(self):
        self.create(self.general)
        self.assertFalse(grading.teacher_can_access_gradebook_context(
            self.db, self.teacher, self.year.id, self.section.id, self.subject.id,
        ))
        another_teacher = User(username="another-teacher", password_hash="unused",
                               full_name="Different Teacher", role="teacher")
        self.db.add(another_teacher)
        self.db.flush()
        other_section = SchoolSection(grade_level_id=self.grade.id, name="Rome",
                                      adviser_user_id=another_teacher.id)
        self.db.add(other_section)
        self.db.commit()
        with self.assertRaisesRegex(ValueError, "not assigned to your teacher account"):
            self.create(self.general, section_id=other_section.id, user=another_teacher)
        self.assertEqual(len(self.db.scalars(select(Gradebook)).all()), 1)

    def test_opening_same_context_returns_existing_record(self):
        first, second = self.create(), self.create()
        self.assertEqual(first.id, second.id)
        self.assertEqual(len(self.db.scalars(select(Gradebook)).all()), 1)

    def test_empty_roster_cannot_be_finalized(self):
        book = self.create()
        self.first.active = False
        self.second.active = False
        self.db.commit()
        with self.assertRaisesRegex(ValueError, 'without active students'):
            grading.finalize_gradebook(self.db, book.id, self.admin, 'Trying empty class')
        self.assertEqual(book.status, 'Draft')

    def test_assessment_maximum_cannot_be_lowered_below_saved_score(self):
        book = self.create(self.general)
        self.score_student(book, 90)
        payload = grading.get_gradebook_full(self.db, book.id)['components']
        payload[0]['items'][0]['max_score'] = 1
        with self.assertRaisesRegex(ValueError, 'below a saved score'):
            grading.update_gradebook_components(self.db, book.id, self.admin, payload, 'Adjust assessment maximum')
        self.assertEqual(book.components[0].items[0].max_score, 20)

    def test_recorded_assessments_cannot_be_silently_removed(self):
        book = self.create(self.general)
        self.score_student(book, 90)
        payload = grading.get_gradebook_full(self.db, book.id)['components'][1:]
        payload[0]['weight'] = 80
        with self.assertRaisesRegex(ValueError, 'recorded scores'):
            grading.update_gradebook_components(self.db, book.id, self.admin, payload, 'Remove written assessment')
        self.assertEqual(len(book.components), 3)

    def test_invalid_score_batch_does_not_partially_change_saved_scores(self):
        book = self.create(self.general)
        self.score_student(book, 90)
        items = [component.items[0] for component in book.components]
        entries = [{'assessment_item_id': items[0].id, 'score': 0, 'status': 'Scored'},
                   {'assessment_item_id': items[1].id, 'score': 9999, 'status': 'Scored'}]
        with self.assertRaisesRegex(ValueError, 'out of bounds'):
            grading.save_gradebook_scores(self.db, book.id, self.admin, {str(self.first.id): entries}, 'Bulk invalid score update')
        saved = self.db.scalar(select(StudentScore).where(StudentScore.person_id == self.first.id, StudentScore.assessment_item_id == items[0].id))
        self.assertEqual(saved.score, 18)

    def test_calculation_modal_receives_component_names_totals_and_passing_status(self):
        book = self.create(self.general)
        self.score_student(book, 90)
        breakdown = grading.get_student_breakdown(self.db, book.id, self.first.id)
        self.assertEqual(breakdown['components'][0]['name'], 'Written Works')
        self.assertEqual(breakdown['components'][0]['total_score'], 18)
        self.assertEqual(breakdown['components'][0]['total_max'], 20)
        self.assertEqual(breakdown['status'], 'Passing')

    def test_adjustments_validate_class_bounds_and_stale_approvals(self):
        book = self.create(self.general)
        self.score_student(book, 90)
        book.status = 'Finalized'
        self.db.commit()
        item = book.components[0].items[0]
        for kwargs, message in [({'person_id': self.outside.id, 'new_score': 19}, 'learner'),
                                ({'person_id': self.first.id, 'new_score': 999}, 'out of bounds')]:
            with self.assertRaisesRegex(ValueError, message):
                grading.create_adjustment_request(self.db, book.id, self.teacher, assessment_item_id=item.id,
                                                  new_status='Scored', reason='Correction based on rubric', **kwargs)
        request = grading.create_adjustment_request(self.db, book.id, self.teacher, self.first.id, item.id, 19, 'Scored', 'Correction based on rubric')
        saved = self.db.scalar(select(StudentScore).where(StudentScore.person_id == self.first.id, StudentScore.assessment_item_id == item.id))
        saved.score = 17
        self.db.commit()
        with self.assertRaisesRegex(ValueError, 'score changed'):
            grading.review_adjustment_request(self.db, request.id, self.admin, 'Approved', 'Review revised rubric')
        self.assertEqual(request.status, 'Pending')
        self.assertEqual(saved.score, 17)

    def test_approved_criteria_adjustment_updates_weights_without_changing_scores(self):
        book = self.create(self.general)
        self.score_student(book, 90)
        book.status = 'Finalized'
        self.db.commit()
        components = [{'component_id': component.id, 'weight': weight}
                      for component, weight in zip(book.components, (20, 50, 30))]
        request = grading.create_criteria_adjustment_request(
            self.db, book.id, self.teacher, components, 'Align criteria with approved workbook')
        self.assertEqual(request.status, 'Pending')
        saved_before = {(score.person_id, score.assessment_item_id): (score.score, score.status)
                        for score in self.db.scalars(select(StudentScore)).all()}
        with self.assertRaisesRegex(ValueError, 'total exactly 100'):
            grading.create_criteria_adjustment_request(
                self.db, book.id, self.teacher,
                [{'component_id': component.id, 'weight': weight}
                 for component, weight in zip(book.components, (20, 50, 25))],
                'Invalid criteria total for test')
        grading.review_criteria_adjustment_request(
            self.db, request.id, self.admin, 'Approved', 'Approved to match workbook')
        self.assertEqual([component.weight for component in book.components], [20, 50, 30])
        saved_after = {(score.person_id, score.assessment_item_id): (score.score, score.status)
                       for score in self.db.scalars(select(StudentScore)).all()}
        self.assertEqual(saved_after, saved_before)
        audit = self.db.scalars(select(GradebookAuditEntry).where(
            GradebookAuditEntry.action == 'CriteriaAdjustmentApproved')).one()
        self.assertIn('20', audit.new_value)

    def test_approved_adjustment_cannot_make_finalized_grade_incomplete(self):
        book = self.create(self.general)
        self.score_student(book, 90)
        self.score_second_student(book, 90)
        book.status = 'Finalized'
        self.db.commit()
        item = book.components[0].items[0]
        with self.assertRaisesRegex(ValueError, 'must include a numeric scored grade'):
            grading.create_adjustment_request(self.db, book.id, self.teacher, self.first.id, item.id, None,
                                              'Missing', 'Assessment was not completed')
        saved = self.db.scalar(select(StudentScore).where(
            StudentScore.person_id == self.first.id, StudentScore.assessment_item_id == item.id))
        request = GradeAdjustmentRequest(
            gradebook_id=book.id, person_id=self.first.id, assessment_item_id=item.id,
            old_score=saved.score, old_status=saved.status, new_score=None, new_status='Missing',
            reason='Legacy incomplete adjustment', requested_by=self.teacher.id,
            requested_by_name=self.teacher.full_name, status='Pending',
        )
        self.db.add(request)
        self.db.commit()
        with self.assertRaisesRegex(ValueError, 'must leave a numeric scored grade'):
            grading.review_adjustment_request(self.db, request.id, self.admin, 'Approved', 'Reviewed missing entry')
        self.assertEqual(request.status, 'Pending')
        self.assertEqual(grading.get_gradebook_full(self.db, book.id)['statistics']['incomplete'], 0)

    def test_renamed_academic_context_refreshes_class_record_labels(self):
        book = self.create(self.general)
        self.section.name = 'Updated Section'
        self.subject.name = 'Updated Core Subject'
        self.db.commit()
        reopened = self.create()
        self.assertEqual(reopened.id, book.id)
        self.assertEqual(reopened.section_name, 'Updated Section')
        self.assertEqual(reopened.subject_name, 'Updated Core Subject')

    def score_student(self, book, percentage):
        entries = [{'assessment_item_id': item.id, 'score': item.max_score * percentage / 100, 'status': 'Scored'} for component in book.components for item in component.items]
        grading.save_gradebook_scores(self.db, book.id, self.admin, {str(self.first.id): entries}, 'Test complete assessment scores')

    def score_second_student(self, book, percentage):
        entries = [{'assessment_item_id': item.id, 'score': item.max_score * percentage / 100, 'status': 'Scored'} for component in book.components for item in component.items]
        grading.save_gradebook_scores(self.db, book.id, self.admin, {str(self.second.id): entries}, 'Complete second learner scores for export test')

    def export(self, book):
        from app.services import reports
        import openpyxl
        with patch.object(reports, 'EXPORT_DIR', Path(TEST_DIR.name)):
            path = reports.generate_gradebook_report_xlsx(self.db, book.id, 'V', 'Test Division', '123456')
        self.export_path = path
        workbook = openpyxl.load_workbook(path)
        full = grading.get_gradebook_full(self.db, book.id)
        cell = 'E16'
        self.assertEqual(workbook['INPUT DATA'][cell].value,
                         full['gradebook']['school_head'])
        return workbook

    def configure_subject_export(self, kind='core'):
        from app.services.class_record_export import template_table, subject_template
        from app.services.reports import TEMPLATES_DIR
        self.grade.name = self.first.grade = self.second.grade = '7'
        self.subject.name, self.subject.category = 'English', 'JHS'
        if kind == 'gmrc':
            self.subject.name = 'GMRC'
            components = [(name, weight, 3 if 'Assessment' in name else 1) for name, weight in (
                ('Written Work Cognitive', 10), ('Written Work Affective', 10),
                ('Performance Task Cognitive', 10), ('Performance Task Affective', 10),
                ('Performance Task Behavioral', 30), ('Quarterly Assessment', 30))]
            self.general.component_definitions_json = json.dumps([
                {'name': name, 'weight': weight, 'max_items': count} for name, weight, count in components])
        else:
            self.general.component_definitions_json = definitions((20, 50 if kind == 'core' else 60, 30 if kind == 'core' else 20), (1, 1, 3))
        component_definitions = json.loads(self.general.component_definitions_json)
        for component in component_definitions:
            if component['name'] == 'Quarterly Assessment':
                component['component_type'] = 'Term Examinations'
        self.general.component_definitions_json = json.dumps(component_definitions)
        _, filename = subject_template(self.subject.name)
        self.general.transmutation_table_json = json.dumps(template_table(TEMPLATES_DIR / 'SUBJECT CLASS RECORDS' / filename))
        self.db.commit()

    def label_export_exams(self, book):
        for item, label in zip(book.components[-1].items, ('ST1', 'ST2', 'TE')):
            item.label = label
        self.db.commit()

    def test_school_head_in_every_core_subject_across_teachers_and_sections(self):
        from app.services import reports
        from app.services.class_record_export import subject_template
        from xml.etree import ElementTree as ET
        self.configure_subject_export()
        self.subject.name = 'Policy Source'
        other_teacher = User(username='subject-export-teacher', full_name='Other Subject Teacher', password_hash='unused', role='teacher')
        other_section = SchoolSection(grade_level_id=self.grade.id, name='Other Section')
        other_student = Person(external_id='EXPORT-OTHER', full_name='Other Learner', sex='Male', role='Student', grade='7', section='Other Section')
        self.db.add_all([other_teacher, other_section, other_student])
        self.db.commit()
        for index, name in enumerate(('English', 'Science', 'Mathematics', 'Filipino', 'Araling Panlipunan', 'TLE', 'Robotics')):
            with self.subTest(subject=name):
                subject = Subject(code=f'HEAD-{index}', name=name, category='JHS')
                self.db.add(subject)
                self.db.commit()
                section, teacher, learners = ((self.section, self.teacher, (self.first, self.second)) if index % 2 == 0
                                               else (other_section, other_teacher, (other_student,)))
                book = self.create(self.general, subject_id=subject.id, section_id=section.id)
                book.teacher_user_id, book.teacher_name = teacher.id, teacher.full_name
                self.label_export_exams(book)
                entries = [{'assessment_item_id': item.id, 'score': item.max_score * 0.8, 'status': 'Scored'}
                           for component in book.components for item in component.items]
                grading.save_gradebook_scores(self.db, book.id, teacher,
                    {str(learner.id): entries for learner in learners}, 'Isolated subject export verification')
                book.status = 'Finalized'
                self.db.commit()
                before = [(score.id, score.score, score.status) for score in self.db.scalars(select(StudentScore)).all()]
                workbook = self.export(book)
                try:
                    self.assertEqual(workbook['INPUT DATA']['E16'].value, 'Buen Owogowog')
                    self.assertEqual(workbook['INPUT DATA']['E23'].value, teacher.full_name)
                    self.assertEqual(workbook['INPUT DATA']['E26'].value, section.name)
                finally:
                    workbook.close()
                _, filename = subject_template(name)
                with ZipFile(reports.TEMPLATES_DIR / 'SUBJECT CLASS RECORDS' / filename) as source, ZipFile(self.export_path) as exported:
                    term_one_path = reports._xlsx_sheet_paths(source)['TERM 1']
                    for part in source.namelist():
                        if part.startswith('xl/worksheets/') and part.endswith('.xml'):
                            def formulas(raw):
                                root = ET.fromstring(raw)
                                return {cell.get('r'): formula.text or '' for cell in root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c')
                                        if (formula := cell.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f')) is not None}
                            before_formulas, after_formulas = formulas(source.read(part)), formulas(exported.read(part))
                            changed = {cell for cell in before_formulas.keys() | after_formulas.keys()
                                       if before_formulas.get(cell) != after_formulas.get(cell)}
                            self.assertLessEqual(changed, {'AC18', 'AC69'} if part == term_one_path else set(), part)
                            for cell in changed:
                                self.assertIn(f'ROUND(AB{cell[2:]},12)', after_formulas[cell], (part, cell))
                self.assertEqual(before, [(score.id, score.score, score.status) for score in self.db.scalars(select(StudentScore)).all()])

    def test_jhs_elective_exports_core_template_with_own_subject_name(self):
        import openpyxl
        self.configure_subject_export()
        self.subject.name = 'Research Elective'
        self.subject.is_elective = True
        self.section.track = 'STE'
        self.db.commit()
        book = self.create(self.general)
        self.label_export_exams(book)
        self.score_student(book, 80)
        self.score_second_student(book, 100)
        grading.finalize_gradebook(self.db, book.id, self.teacher, 'Isolated elective export test')
        before = grading.get_gradebook_full(self.db, book.id)
        workbook = self.export(book)
        self.assertEqual(workbook['INPUT DATA']['E24'].value, 'Research Elective')
        self.assertTrue(workbook['TERM 1']['AB18'].value.startswith('='))
        workbook.close()
        values = openpyxl.load_workbook(self.export_path, data_only=True)
        self.assertEqual(values['TERM 1']['AB18'].value, 80)
        self.assertEqual(values['TERM 1']['AC18'].value, 83)
        values.close()
        self.assertEqual(before, grading.get_gradebook_full(self.db, book.id))

    def test_subject_export_preserves_formulas_and_matches_saved_grades(self):
        import openpyxl
        from app.services import reports
        from app.services.settings_store import set_json
        set_json(self.db, 'school.heads', {str(self.year.id): {'full_name': 'Buen Owogowog', 'title': 'Principal'}, '9999': {'full_name': 'Previous Head'}})
        self.configure_subject_export()
        book = self.create(self.general)
        self.label_export_exams(book)
        self.score_student(book, 80)
        self.score_second_student(book, 100)
        book.status = 'Finalized'
        self.db.commit()
        before = [(score.id, score.score, score.status) for score in self.db.scalars(select(StudentScore)).all()]
        workbook = self.export(book)
        try:
            self.assertEqual(workbook['INPUT DATA']['E11'].value, 'Test Division')
            self.assertEqual(workbook['INPUT DATA']['E13'].value, '123456')
            self.assertEqual(workbook['INPUT DATA']['E15'].value, '2026-2027')
            self.assertEqual(workbook['INPUT DATA']['E16'].value, 'Buen Owogowog')
            self.assertEqual(workbook['INPUT DATA']['E16'].value, grading.get_gradebook_full(self.db, book.id)['gradebook']['school_head'])
            self.assertEqual(workbook['INPUT DATA']['E25'].value, 7)
            self.assertEqual(workbook['INPUT DATA']['E26'].value, 'Berlin')
            self.assertTrue(workbook['TERM 1']['AB18'].value.startswith('='))
            self.assertEqual(workbook['TERM 1']['F15'].value, 20)
            self.assertEqual(workbook['TERM 1']['T18'].value, 40)
        finally:
            workbook.close()
        data = openpyxl.load_workbook(self.export_path, data_only=True)
        self.assertEqual(data['TERM 1']['AB18'].value, 80)
        self.assertEqual(data['TERM 1']['AC18'].value, 83)
        self.assertEqual(data['TERM 1']['AC69'].value, 100)
        data.close()
        self.assertEqual(before, [(score.id, score.score, score.status) for score in self.db.scalars(select(StudentScore)).all()])

    def test_export_rejects_policy_and_exam_mismatches_without_changing_scores(self):
        from app.services.class_record_export import fill_term, template_table
        from app.services.reports import TEMPLATES_DIR
        self.configure_subject_export()
        book = self.create(self.general)
        self.label_export_exams(book)
        self.score_student(book, 80)
        self.score_second_student(book, 80)
        book.status = 'Finalized'
        self.db.commit()
        full = grading.get_gradebook_full(self.db, book.id)
        table = template_table(TEMPLATES_DIR / 'SUBJECT CLASS RECORDS' / 'ENGLISH - Grade 7-10 Class Record.xlsx')
        for mismatch in ('weight', 'table', 'exam'):
            data = copy.deepcopy(full)
            if mismatch == 'weight':
                data['components'][0]['weight'] = 30
            elif mismatch == 'table':
                data['policy']['transmutation_table'] = []
            else:
                data['components'][-1]['items'][0]['label'] = 'Unknown examination'
            with self.subTest(mismatch=mismatch), self.assertRaises(HTTPException):
                fill_term(data, 'core', 'TERM 1', {}, {}, data['students'][:1], data['students'][1:], table)

    def test_all_subject_files_and_aliases_resolve_to_installed_workbooks(self):
        from app.services.class_record_export import subject_template
        from app.services.reports import TEMPLATES_DIR
        for subject in ('English', 'Science', 'Mathematics', 'Math', 'Filipino', 'Araling Panlipunan', 'AP', 'TLE',
                        'Technology and Livelihood Education', 'GMRC', 'Values Education', 'EsP', 'SPA', 'SPS',
                        'Music and Arts', 'Physical Education and Health', 'PE & Health'):
            with self.subTest(subject=subject):
                _, filename = subject_template(subject)
                self.assertTrue((TEMPLATES_DIR / 'SUBJECT CLASS RECORDS' / filename).is_file())

    def test_unknown_jhs_subject_exports_its_saved_gradebook_layout(self):
        from app.services import reports
        self.grade.name, self.subject.name, self.subject.category = '7', 'Unsupported Subject', 'JHS'
        self.first.grade, self.second.grade = '7', '7'
        self.db.commit()
        book = self.create(self.general)
        self.score_student(book, 80)
        self.score_second_student(book, 80)
        book.status = 'Finalized'
        self.db.commit()
        workbook = self.export(book)
        try:
            self.assertEqual(workbook['INPUT DATA']['E24'].value, 'Unsupported Subject')
            self.assertIn('Term 1', workbook.sheetnames)
        finally:
            workbook.close()

    def test_missing_selected_template_uses_the_gradebook_instead_of_an_unrelated_workbook(self):
        from fastapi import HTTPException
        from app.services import reports
        self.grade.name, self.subject.name, self.subject.category = '7', 'English', 'JHS'
        self.first.grade, self.second.grade = '7', '7'
        self.db.commit()
        book = self.create(self.general)
        self.score_student(book, 80)
        self.score_second_student(book, 80)
        book.status = 'Finalized'
        self.db.commit()
        with tempfile.TemporaryDirectory(dir=TEST_DIR.name) as directory:
            templates = Path(directory)
            (templates / 'UPDATED CLASS RECORD').mkdir()
            (templates / 'UPDATED CLASS RECORD' / 'unrelated.xlsx').write_bytes(b'unrelated template must not be opened')
            with patch.object(reports, 'TEMPLATES_DIR', templates):
                workbook = self.export(book)
                try:
                    self.assertEqual(workbook['INPUT DATA']['E24'].value, 'English')
                    self.assertIn('Term 1', workbook.sheetnames)
                finally:
                    workbook.close()

    def test_shs_export_uses_all_three_terms_and_authoritative_results(self):
        first_semester = self.create(self.general)
        self.score_student(first_semester, 90)
        second = GradingPeriod(school_year_id=self.year.id, quarter=2, name='Term 2', starts_on=date(2026, 9, 1), ends_on=date(2026, 10, 31))
        third = GradingPeriod(school_year_id=self.year.id, quarter=3, name='Term 3', starts_on=date(2026, 11, 1), ends_on=date(2027, 1, 31))
        self.db.add(second)
        self.db.commit()
        second_book = self.create(self.general, grading_period_id=second.id)
        self.score_student(second_book, 70)
        self.score_second_student(second_book, 70)
        second_book.status = 'Finalized'
        self.db.add(third)
        self.db.commit()
        third_book = self.create(self.general, grading_period_id=third.id)
        self.score_student(third_book, 40)
        self.score_second_student(third_book, 40)
        self.score_second_student(first_semester, 90)
        third_book.status = 'Finalized'
        first_semester.status = 'Finalized'
        self.db.commit()
        from app.services.settings_store import set_json
        set_json(self.db, 'school.heads', {str(self.year.id): {'full_name': 'Buen Owogowog', 'title': 'Principal'}})
        workbook = self.export(third_book)
        try:
            self.assertEqual(workbook['INPUT DATA']['E16'].value, 'Buen Owogowog')
            self.assertEqual(workbook['INPUT DATA']['E25'].value, '11')
            self.assertEqual(workbook['Term 1']['B12'].value, 'First Student')
            self.assertTrue(workbook['Term 1']['Q12'].value.startswith('='))
            self.assertTrue(workbook['Term 1']['R12'].value.startswith('='))
        finally:
            workbook.close()
        import openpyxl
        workbook = openpyxl.load_workbook(self.export_path, data_only=True)
        try:
            for term, initial, grade in ((1, 90, 93), (2, 70, 81), (3, 40, 70)):
                self.assertEqual(workbook[f'Term {term}']['Q12'].value, initial)
                self.assertEqual(workbook[f'Term {term}']['R12'].value, grade)
        finally:
            workbook.close()

    def test_custom_weights_maxima_and_attendance_export_without_changing_saved_scores(self):
        self.configure_subject_export()
        definitions = json.loads(self.general.component_definitions_json)
        definitions[0]['weight'], definitions[1]['weight'] = 25, 35
        definitions.append({'name': 'Attendance', 'weight': 10, 'max_items': 1})
        self.general.component_definitions_json = json.dumps(definitions)
        self.db.commit()
        book = self.create(self.general)
        for component in book.components:
            if component.name == 'Quarterly Assessment':
                for item, label, maximum in zip(component.items, ('ST1', 'ST2', 'TE'), (30, 20, 50)):
                    item.label, item.max_score = label, maximum
        book.components[0].items[0].max_score = 30
        self.db.commit()
        self.score_student(book, 80)
        self.score_second_student(book, 70)
        book.status = 'Finalized'
        self.db.commit()
        before = grading.get_gradebook_full(self.db, book.id)
        workbook = self.export(book)
        try:
            sheet = workbook['Term 1']
            self.assertEqual(sheet['E11'].value, 30)
            self.assertEqual(sheet['H11'].value, 0.25)
            self.assertIn('Attendance', [cell.value for cell in sheet[9]])
            self.assertTrue(any(cell.value == 'TERM GRADE' for cell in sheet[10]))
        finally:
            workbook.close()
        after = grading.get_gradebook_full(self.db, book.id)
        self.assertEqual(before, after)
        with self.assertRaises(HTTPException) as error:
            from app.services.reports import generate_gradebook_report_xlsx
            generate_gradebook_report_xlsx(self.db, book.id, 'V', 'V', '301874', allowed_teacher_id=book.teacher_user_id + 99999)
        self.assertEqual(error.exception.status_code, 403)

    def test_custom_music_and_health_setups_keep_the_combined_workbook(self):
        self.configure_subject_export('mapeh')
        self.general.component_definitions_json = definitions((30, 50, 20), (1, 1, 3))
        self.subject.name = 'Music and Arts'
        self.db.commit()
        arts = self.create(self.general)
        health = Subject(code='PEH-CUSTOM', name='Physical Education and Health', category='JHS')
        self.db.add(health)
        self.db.commit()
        pe = self.create(self.general, subject_id=health.id)
        for book in (arts, pe):
            self.label_export_exams(book)
            self.score_student(book, 80)
            self.score_second_student(book, 80)
            book.status = 'Finalized'
        self.db.commit()
        workbook = self.export(arts)
        try:
            self.assertIn('Term 1 M and A', workbook.sheetnames)
            self.assertIn('Term 1 PE and H', workbook.sheetnames)
            self.assertTrue(workbook['FINAL GRADES']['E5'].value.startswith('='))
        finally:
            workbook.close()
        import openpyxl
        workbook = openpyxl.load_workbook(self.export_path, data_only=True)
        try:
            self.assertEqual(workbook['FINAL GRADES']['E5'].value, 83)
            self.assertIsNone(workbook['FINAL GRADES']['L5'].value)
        finally:
            workbook.close()

    def test_export_uses_the_saved_calculation_even_when_template_weights_match(self):
        for method, examination_type in (('floor', 'Term Examinations'), ('half_up', 'Quarterly Assessment')):
            with self.subTest(method=method, examination_type=examination_type):
                self.configure_subject_export()
                self.general.rounding_method = method
                self.db.commit()
                book = self.create(self.general)
                book.status = 'Draft'
                book.components[-1].component_type = examination_type
                self.db.commit()
                self.label_export_exams(book)
                self.score_student(book, 83.123)
                self.score_second_student(book, 71.2)
                book.status = 'Finalized'
                self.db.commit()
                workbook = self.export(book)
                try:
                    self.assertIn('Term 1', workbook.sheetnames)
                finally:
                    workbook.close()
    def test_mapeh_export_uses_updated_three_term_workbook(self):
        self.configure_subject_export('mapeh')
        self.subject.name = 'Music and Arts'
        other = Subject(code='PEH', name='Physical Education and Health', category='JHS')
        self.db.add(other)
        self.db.commit()
        books = []
        for subject in (self.subject, other):
            book = self.create(self.general, subject_id=subject.id)
            self.label_export_exams(book)
            self.score_student(book, 80)
            self.score_second_student(book, 100)
            book.status = 'Finalized'
            books.append(book)
        self.db.commit()
        workbook = self.export(books[0])
        try:
            self.assertEqual(workbook['TERM 1 M and A']['F18'].value, 16)
            self.assertEqual(workbook['TERM 1 PE and H']['F18'].value, 16)
            self.assertEqual(workbook['FINAL GRADES']['B65'].value, 'FEMALE')
            self.assertIn('COUNT(F66:G66)=2', workbook['FINAL GRADES']['H66'].value)
            self.assertIn('AC69', workbook['FINAL GRADES']['F66'].value)
        finally:
            workbook.close()
        from app.services import reports
        import openpyxl
        books[0].teacher_user_id = self.teacher.id
        other_teacher = User(username='other-export-teacher', full_name='Other Teacher', password_hash='unused', role='teacher')
        self.db.add(other_teacher)
        self.db.flush()
        books[1].teacher_user_id = other_teacher.id
        self.db.commit()
        with patch.object(reports, 'EXPORT_DIR', Path(TEST_DIR.name)):
            path = reports.generate_gradebook_report_xlsx(self.db, books[0].id, 'V', 'Test Division', '123456', allowed_teacher_id=self.teacher.id)
        scoped = openpyxl.load_workbook(path)
        try:
            self.assertEqual(scoped['TERM 1 M and A']['F18'].value, 16)
            self.assertIsNone(scoped['TERM 1 PE and H']['F18'].value)
        finally:
            scoped.close()

    def check_records_officer_three_term_export(self, grade):
        import openpyxl
        from app.services import reports
        self.grade.name = self.first.grade = self.second.grade = str(grade)
        self.subject.category = 'JHS' if grade < 11 else 'SHS-Core'
        self.outside.grade = '10'
        officer = User(username='export-officer', full_name='Test Records Officer', password_hash='unused', role='records_officer')
        self.db.add(officer)
        self.db.commit()
        books = []
        for quarter in (1, 2, 3):
            period = self.period if quarter == 1 else GradingPeriod(
                school_year_id=self.year.id, quarter=quarter, name=f'Term {quarter}',
                starts_on=date(2026, 6, 1), ends_on=date(2027, 3, 31))
            self.db.add(period)
            self.db.commit()
            book = self.create(self.general, grading_period_id=period.id)
            book.teacher_user_id, book.teacher_name = self.teacher.id, self.teacher.full_name
            self.score_student(book, 60 + quarter * 10)
            self.score_second_student(book, 65 + quarter * 10)
            book.status = 'Locked' if quarter == 3 else 'Finalized'
            self.db.commit()
            books.append(book)
        snapshots = {book.id: grading.get_gradebook_full(self.db, book.id) for book in books}
        with patch.object(reports, 'EXPORT_DIR', Path(TEST_DIR.name)):
            response = gradebook_report_xlsx_new(books[0].id, region='V', division='V', school_id='301874',
                                                all_terms=True, db=self.db, user=officer)
            self.assertIn('AllTerms', response.filename)
            formulas = openpyxl.load_workbook(response.path)
            values = openpyxl.load_workbook(response.path, data_only=True)
            try:
                self.assertEqual(values.sheetnames, ['INPUT DATA', 'Term 1', 'Term 2', 'Term 3'])
                for book in books:
                    full = snapshots[book.id]
                    sheet = values[f'Term {book.quarter}']
                    headings = {cell.value: cell.column for cell in sheet[10]}
                    learners = {student['external_id']: student for student in full['students']}
                    self.assertEqual(len(learners), 2)
                    for row in (12, 13):
                        student = learners[sheet.cell(row, 3).value]
                        for heading, key in (('INITIAL GRADE', 'initial_grade'), ('TERM GRADE', 'reported_grade')):
                            self.assertEqual(sheet.cell(row, headings[heading]).value, student[key])
                            self.assertEqual(formulas[sheet.title].cell(row, headings[heading]).data_type, 'f')
                        column = 5
                        for component, result in zip(full['components'], student['component_grades']):
                            for item in (item for item in component['items'] if item['active']):
                                self.assertEqual(sheet.cell(row, column).value, student['scores'][str(item['id'])])
                                self.assertEqual(sheet.cell(11, column).value, item['max_score'])
                                column += 1
                            self.assertEqual(sheet.cell(row, column + 1).value, result['percentage'])
                            self.assertEqual(sheet.cell(row, column + 2).value, result['weighted_score'])
                            column += 3
                self.assertEqual(values['INPUT DATA']['E25'].value, str(grade))
                self.assertEqual(values['INPUT DATA']['E16'].value, 'Buen Owogowog')
            finally:
                formulas.close()
                values.close()
            single = gradebook_report_xlsx_new(books[0].id, region='V', division='V', school_id='301874',
                                              all_terms=False, db=self.db, user=officer)
            self.assertIn('Term1', single.filename)
            self.assertEqual(snapshots, {book.id: grading.get_gradebook_full(self.db, book.id) for book in books})
            books[1].status = 'Draft'
            self.db.commit()
            with self.assertRaises(HTTPException) as blocked:
                gradebook_report_xlsx_new(books[0].id, region='V', division='V', school_id='301874',
                                         all_terms=True, db=self.db, user=officer)
            self.assertEqual(blocked.exception.status_code, 409)

    def test_records_officer_all_terms_grade_7(self):
        self.check_records_officer_three_term_export(7)

    def test_records_officer_all_terms_grade_11(self):
        self.check_records_officer_three_term_export(11)

    def test_records_officer_all_terms_grade_12(self):
        self.check_records_officer_three_term_export(12)

    def test_jhs_export_writes_students_and_all_three_term_grades(self):
        self.configure_subject_export()
        self.subject.name, self.subject.category = 'English', 'JHS'
        self.db.commit()
        books = []
        for quarter in (1, 2, 3):
            period = self.period if quarter == 1 else GradingPeriod(
                school_year_id=self.year.id, quarter=quarter, name=f'Term {quarter}',
                starts_on=date(2026, 6, 1), ends_on=date(2027, 3, 31))
            if quarter != 1:
                self.db.add(period)
                self.db.commit()
            book = self.create(self.general, grading_period_id=period.id)
            self.label_export_exams(book)
            self.score_student(book, 70 if quarter == 1 else 71.18)
            self.score_second_student(book, 70 if quarter == 1 else 71.18)
            book.status = 'Finalized'
            books.append(book)
        self.db.commit()
        workbook = self.export(books[-1])
        workbook.close()
        import openpyxl
        workbook = openpyxl.load_workbook(self.export_path, data_only=True)
        try:
            self.assertEqual(workbook['Term 1']['B12'].value, 'First Student')
            self.assertEqual(workbook['Term 1']['B13'].value, 'Second Student')
            for book in books:
                sheet = workbook[f'Term {book.quarter}']
                column = next(cell.column for cell in sheet[10] if cell.value == 'TERM GRADE')
                expected = grading.get_gradebook_full(self.db, book.id)['students'][0]['reported_grade']
                self.assertEqual(sheet.cell(12, column).value, expected)
        finally:
            workbook.close()
        from app.services import reports
        print_report = reports.gradebook_report_html(self.db, books[-1].id)
        self.assertIn('Term 3', print_report)
        self.assertNotIn('Quarter 3', print_report)

    def test_jhs_export_leaves_draft_sibling_term_blank(self):
        self.configure_subject_export()
        self.subject.name, self.subject.category = 'English', 'JHS'
        self.db.commit()
        books = []
        for quarter in (1, 2, 3):
            period = self.period if quarter == 1 else GradingPeriod(
                school_year_id=self.year.id, quarter=quarter, name=f'Term {quarter}',
                starts_on=date(2026, 6, 1), ends_on=date(2027, 3, 31))
            if quarter != 1:
                self.db.add(period)
                self.db.commit()
            book = self.create(self.general, grading_period_id=period.id)
            self.label_export_exams(book)
            self.score_student(book, 70 + quarter)
            self.score_second_student(book, 70 + quarter)
            book.status = 'Draft' if quarter == 2 else 'Finalized'
            books.append(book)
        self.db.commit()
        workbook = self.export(books[-1])
        workbook.close()
        import openpyxl
        workbook = openpyxl.load_workbook(self.export_path, data_only=True)
        try:
            self.assertEqual(workbook['TERM 1']['AC18'].value, 75)
            self.assertIsNone(workbook['TERM 2']['F18'].value)
            self.assertEqual(workbook['TERM 3']['AC18'].value, 77)
        finally:
            workbook.close()

    def test_gmrc_spa_and_sps_exports_use_their_updated_workbooks(self):
        for name in ('GMRC', 'Values Education', 'SPA', 'SPS'):
            if name != 'Values Education':
                self.configure_subject_export('gmrc' if name == 'GMRC' else 'spa')
            self.subject.name = 'Policy Source'
            self.db.commit()
            subject = Subject(code=f'EXPORT-{name}', name=name, category='JHS')
            self.db.add(subject)
            self.db.commit()
            book = self.create(self.general, subject_id=subject.id)
            self.label_export_exams(book)
            self.score_student(book, 80)
            self.score_second_student(book, 100)
            book.status = 'Finalized'
            self.db.commit()
            workbook = self.export(book)
            try:
                self.assertEqual(workbook['INPUT DATA']['K11'].value, 'First Student')
                hps_row = 16 if name in {'GMRC', 'Values Education'} else 15
                self.assertEqual(workbook['TERM 1'][f'F{hps_row}'].value, 20)
                if name in {'GMRC', 'Values Education'}:
                    self.assertEqual(workbook['TERM 1']['F15'].value, 1)
                    self.assertEqual(workbook['TERM 1']['N16'].value, 20)
                    self.assertEqual(workbook['TERM 1']['AH16'].value, 50)
                    self.assertTrue(workbook['TERM 1']['AV19'].value.startswith('='))
            finally:
                workbook.close()

    def test_export_does_not_merge_overflow_or_misreport_examination_weights(self):
        from app.services.class_record_export import fill_term, template_table
        from app.services.reports import TEMPLATES_DIR
        self.configure_subject_export()
        book = self.create(self.general)
        self.label_export_exams(book)
        self.score_student(book, 80)
        self.score_second_student(book, 80)
        full = grading.get_gradebook_full(self.db, book.id)
        table = template_table(TEMPLATES_DIR / 'SUBJECT CLASS RECORDS' / 'ENGLISH - Grade 7-10 Class Record.xlsx')
        excessive = copy.deepcopy(full)
        excessive['components'][0]['items'] *= 6
        with self.assertRaisesRegex(HTTPException, 'silently merged'):
            fill_term(excessive, 'core', 'TERM 1', {}, {}, excessive['students'][:1], excessive['students'][1:], table)
        unequal = copy.deepcopy(full)
        exam_ids = [str(item['id']) for item in unequal['components'][-1]['items']]
        unequal['students'][0]['scores'][exam_ids[0]] = 0
        with self.assertRaisesRegex(HTTPException, 'calculation differs'):
            fill_term(unequal, 'core', 'TERM 1', {}, {}, unequal['students'][:1], unequal['students'][1:], table)

    def test_export_requires_school_details_and_preserves_every_core_formula(self):
        from app.services import reports
        from xml.etree import ElementTree as ET
        self.configure_subject_export()
        book = self.create(self.general)
        self.label_export_exams(book)
        self.score_student(book, 80)
        self.score_second_student(book, 80)
        book.status = 'Finalized'
        self.db.commit()
        with self.assertRaises(HTTPException) as error:
            reports.generate_gradebook_report_xlsx(self.db, book.id, 'V', '', '123456')
        self.assertEqual(error.exception.status_code, 422)
        self.export(book).close()
        source = reports.TEMPLATES_DIR / 'SUBJECT CLASS RECORDS' / 'ENGLISH - Grade 7-10 Class Record.xlsx'
        with ZipFile(source) as original, ZipFile(self.export_path) as exported:
            term_one_path = reports._xlsx_sheet_paths(original)['TERM 1']
            for name in original.namelist():
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    def formulas(raw):
                        root = ET.fromstring(raw)
                        return {cell.get('r'): formula.text or '' for cell in root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c')
                                if (formula := cell.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f')) is not None}
                    before_formulas, after_formulas = formulas(original.read(name)), formulas(exported.read(name))
                    changed = {cell for cell in before_formulas.keys() | after_formulas.keys()
                               if before_formulas.get(cell) != after_formulas.get(cell)}
                    self.assertLessEqual(changed, {'AC18', 'AC69'} if name == term_one_path else set(), name)
                    for cell in changed:
                        self.assertIn(f'ROUND(AB{cell[2:]},12)', after_formulas[cell], (name, cell))


    def test_template_policy_alignment_keeps_every_score_and_existing_maximum(self):
        from app.services.template_policies import ensure_template_policies, align_open_gradebook
        self.grade.name = self.first.grade = self.second.grade = '7'
        self.subject.name, self.subject.category = 'Mathematics', 'JHS'
        self.general.component_definitions_json = definitions((40, 40, 20))
        self.db.commit()
        book = self.create(self.general)
        self.score_student(book, 80)
        self.score_second_student(book, 100)
        original = [(row.id, row.assessment_item_id, row.score, row.status, row.reason, row.updated_at)
                    for row in self.db.scalars(select(StudentScore).order_by(StudentScore.id)).all()]
        maxima = {item.id: item.max_score for component in book.components for item in component.items}
        old_policy = self.general.component_definitions_json
        previous_exam = book.components[-1].items[0].id
        policies = ensure_template_policies(self.db, self.year.name, self.admin)
        align_open_gradebook(self.db, book, self.admin, policies)
        align_open_gradebook(self.db, book, self.admin, policies)
        self.assertEqual([component.weight for component in book.components], [20, 50, 30])
        self.assertEqual({item.label for item in book.components[-1].items}, {'ST1', 'ST2', 'TE'})
        self.assertEqual(next(item.label for item in book.components[-1].items if item.id == previous_exam), 'TE')
        self.assertEqual(original, [(row.id, row.assessment_item_id, row.score, row.status, row.reason, row.updated_at)
                                   for row in self.db.scalars(select(StudentScore).order_by(StudentScore.id)).all()])
        self.assertTrue(all(item.max_score == maxima[item.id] for component in book.components for item in component.items if item.id in maxima))
        self.assertEqual(self.general.component_definitions_json, old_policy)
        self.assertEqual(grading.get_gradebook_full(self.db, book.id)['statistics']['incomplete'], 2)
        self.assertEqual(len(self.db.scalars(select(GradebookAuditEntry).where(GradebookAuditEntry.action == 'PolicyAlignment')).all()), 1)
        with self.assertRaises(ValueError) as blocked:
            grading.finalize_gradebook(self.db, book.id, self.teacher, 'Missing ST1/ST2')
        for learner_name in ('First Student', 'Second Student'):
            self.assertIn(f'{learner_name}:', str(blocked.exception))
        for label in ('ST1', 'ST2'):
            self.assertEqual(str(blocked.exception).count(f'{label} (score missing)'), 2)
        self.score_student(book, 80)
        self.score_second_student(book, 100)
        book.status = 'Finalized'
        self.db.commit()
        self.export(book).close()
        import openpyxl
        workbook = openpyxl.load_workbook(self.export_path, data_only=True)
        self.assertEqual(workbook['TERM 1']['AC18'].value, 83)
        workbook.close()
        with self.assertRaisesRegex(ValueError, 'Finalized record'):
            align_open_gradebook(self.db, book, self.admin, policies)

    def test_updated_exam_weighting_and_conversion_match_excel_boundaries(self):
        from app.services.grading_engine import calculate_component_percentage, calculate_student_grade
        items = [dict(id=index, label=label, max_score=50) for index, label in enumerate(('ST1', 'ST2', 'TE'), 1)]
        scores = [dict(item_id=index, score=score, status='Scored') for index, score in enumerate((50, 0, 25), 1)]
        self.assertEqual(calculate_component_percentage(scores, items, 'Term Examinations')['percentage'], 50)
        self.assertFalse(calculate_component_percentage(scores[2:], items, 'Term Examinations')['complete'])
        from app.services.class_record_export import template_table
        from app.services.reports import TEMPLATES_DIR
        table = template_table(TEMPLATES_DIR / 'SUBJECT CLASS RECORDS' / 'ENGLISH - Grade 7-10 Class Record.xlsx')
        for score, reported in ((71.175, 75), (71.18, 76), (0, 60), (100, 100)):
            exam_items = [dict(item, max_score=100) for item in items]
            result = calculate_student_grade([dict(item_id=item['id'], score=score, status='Scored') for item in exam_items],
                                             [dict(id=1, name='Examinations', weight=100, component_type='Term Examinations')],
                                             {1: exam_items}, table)
            self.assertEqual(result['reported_grade'], reported)


if __name__ == '__main__':
    try:
        unittest.main()
    finally:
        TEST_DIR.cleanup()

