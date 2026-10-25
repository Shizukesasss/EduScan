"""Integration tests for EduScan Grading Management API."""
from __future__ import annotations

import os
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))
TEST_DIR = tempfile.TemporaryDirectory(prefix="eduscan-grading-test-")
database_path = Path(TEST_DIR.name) / "grading_test.db"
os.environ["DATABASE_URL"] = f"sqlite:///{database_path.as_posix()}"
os.environ["EDUSCAN_DATA_DIR"] = TEST_DIR.name
os.environ["EDUSCAN_SECRET_KEY"] = "grading-test-secret-key-2026"
os.environ["SMS_GATEWAY_ENABLED"] = "false"
os.environ["ATTENDANCE_SCHEDULER_ENABLED"] = "false"

from fastapi.testclient import TestClient  # noqa: E402
from app.main import app  # noqa: E402


class GradingApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.client_context = TestClient(app)
        cls.client = cls.client_context.__enter__()
        # Login admin
        resp = cls.client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
        assert resp.status_code == 200, resp.text
        cls.admin_headers = {"Authorization": f"Bearer {resp.json()['access_token']}"}
        # Change password to satisfy security requirement
        cls.client.post("/api/auth/change-password", headers=cls.admin_headers, json={
            "current_password": "admin123", "new_password": "Grading-Admin-Pass-2026!",
        })
        teacher_login = cls.client.post('/api/auth/login', json={'username': 'teacher', 'password': 'teacher123'})
        assert teacher_login.status_code == 200, teacher_login.text
        cls.teacher_headers = {'Authorization': f"Bearer {teacher_login.json()['access_token']}"}
        changed = cls.client.post('/api/auth/change-password', headers=cls.teacher_headers,
                                  json={'current_password': 'teacher123', 'new_password': 'Teacher-Grading-Pass-2026!'})
        assert changed.status_code == 200, changed.text
        cls.teacher = next(row for row in cls.client.get('/api/admin/users', headers=cls.admin_headers).json() if row['username'] == 'teacher')
        structure = cls.client.get('/api/admin/academic-structure', headers=cls.admin_headers).json()
        first_section = structure['sections'][0]
        result = cls.client.post('/api/admin/sections', headers=cls.admin_headers,
                                 json={**first_section, 'adviser_user_id': cls.teacher['id'], 'adviser_name': cls.teacher['full_name']})
        assert result.status_code == 200, result.text
        for year in structure['school_years']:
            for subject in structure['subjects']:
                is_shs_subject = str(subject.get('category', 'JHS')).upper().startswith('SHS')
                section_ids = [section['id'] for section in structure['sections']
                               if (str(section['grade']).strip().lower().replace('grade', '').strip() in {'11', '12'}) == is_shs_subject
                               and (not subject.get('is_elective') or not subject.get('elective_track')
                                    or section.get('track') == subject['elective_track'])]
                if section_ids:
                    assigned = cls.client.post('/api/admin/subject-teaching-assignments', headers=cls.admin_headers,
                                               json={'school_year_id': year['id'], 'subject_id': subject['id'],
                                                     'teacher_user_id': cls.teacher['id'], 'section_ids': section_ids})
                    assert assigned.status_code == 200, assigned.text
        for username, role, target in [('grade-records', 'records_officer', 'records_headers'), ('other-teacher', 'teacher', 'other_teacher_headers')]:
            created = cls.client.post('/api/admin/users', headers=cls.admin_headers,
                                      json={'username': username, 'password': 'Initial-Grading-Pass-2026!', 'role': role, 'full_name': username.title(), 'active': True})
            assert created.status_code == 200, created.text
            login = cls.client.post('/api/auth/login', json={'username': username, 'password': 'Initial-Grading-Pass-2026!'})
            assert login.status_code == 200, login.text
            headers = {'Authorization': f"Bearer {login.json()['access_token']}"}
            changed = cls.client.post('/api/auth/change-password', headers=headers,
                                      json={'current_password': 'Initial-Grading-Pass-2026!', 'new_password': 'Updated-Grading-Pass-2026!'})
            assert changed.status_code == 200, changed.text
            setattr(cls, target, headers)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.client_context.__exit__(None, None, None)
        from app.database import engine
        engine.dispose()
        TEST_DIR.cleanup()

    def assert_exports_match_gradebook(self, gb_id):
        snapshot = self.client.get(f'/api/gradebooks/{gb_id}', headers=self.teacher_headers).json()
        for headers in (self.other_teacher_headers, self.admin_headers):
            self.assertEqual(self.client.get(f'/api/gradebooks/{gb_id}/report.xlsx?all_terms=true', headers=headers).status_code, 403)
        self.assertEqual(self.client.get(f'/api/gradebooks/{gb_id}', headers=self.admin_headers).json(), snapshot)
        for role, headers in (('teacher', self.teacher_headers), ('records officer', self.records_headers)):
            with self.subTest(export_role=role):
                exported = self.client.get(f'/api/gradebooks/{gb_id}/report.xlsx', headers=headers,
                                           params={'region': 'V', 'division': 'V', 'school_id': '301874'})
                self.assertEqual(exported.status_code, 200, exported.text if exported.status_code != 200 else '')
                formulas = load_workbook(BytesIO(exported.content))
                values = load_workbook(BytesIO(exported.content), data_only=True)
                try:
                    term_name = f"Term {snapshot['gradebook']['term_order']}"
                    sheet = values[term_name]
                    headings = {cell.value: cell.column for cell in sheet[10]}
                    for heading in ('INITIAL GRADE', 'TERM GRADE'):
                        self.assertEqual(formulas[term_name].cell(12, headings[heading]).data_type, 'f')
                    metadata = snapshot['gradebook']
                    for cell, key in (('E15', 'school_year_name'), ('E23', 'teacher_name'), ('E24', 'subject_name'), ('E25', 'grade_name'), ('E26', 'section_name')):
                        self.assertEqual(values['INPUT DATA'][cell].value, metadata[key])
                    self.assertEqual(values['INPUT DATA']['E13'].value, '301874')
                    self.assertEqual(values['INPUT DATA']['E10'].value, 'V')
                    self.assertEqual(values['INPUT DATA']['E11'].value, 'V')
                    self.assertEqual(values['INPUT DATA']['E16'].value or '', metadata['school_head'] or '')
                    by_identifier = {student['lrn'] or student['external_id']: student for student in snapshot['students']}
                    for row in range(12, 12 + len(by_identifier)):
                        student = by_identifier[sheet.cell(row, 3).value]
                        self.assertEqual(sheet.cell(row, headings['INITIAL GRADE']).value, student['initial_grade'])
                        self.assertEqual(sheet.cell(row, headings['TERM GRADE']).value, student['reported_grade'])
                        column = 5
                        for component, detail, result in zip(snapshot['components'], student['component_details'], student['component_grades']):
                            self.assertEqual(formulas[term_name].cell(9, column).value, component['name'])
                            items = [item for item in component['items'] if item['active']]
                            for item in items:
                                self.assertEqual(sheet.cell(11, column).value, item['max_score'])
                                self.assertEqual(sheet.cell(row, column).value, student['scores'][str(item['id'])])
                                column += 1
                            self.assertEqual(sheet.cell(11, column + 2).value, component['weight'] / 100)
                            for offset, expected in enumerate((detail['total_earned'], result['percentage'], result['weighted_score'])):
                                self.assertAlmostEqual(sheet.cell(row, column + offset).value, expected, places=10)
                            column += 3
                finally:
                    formulas.close()
                    values.close()
        self.assertEqual(self.client.get(f'/api/gradebooks/{gb_id}/report.xlsx', headers=self.other_teacher_headers).status_code, 403)
        self.assertEqual(self.client.get(f'/api/gradebooks/{gb_id}/report.xlsx', headers=self.admin_headers).status_code, 403)
        self.assertEqual(self.client.get(f'/api/gradebooks/{gb_id}', headers=self.teacher_headers).json(), snapshot)


    def test_01_grading_policies_list_seeded(self) -> None:
        resp = self.client.get("/api/grading-policies", headers=self.admin_headers)
        self.assertEqual(resp.status_code, 200, resp.text)
        policies = resp.json()
        self.assertGreaterEqual(len(policies), 4)
        names = [p["name"] for p in policies]
        self.assertTrue(any("Languages/AP/EsP" in n for n in names))
        self.assertTrue(any("Math & Science" in n for n in names))
        self.assertTrue(any("MAPEH & TLE" in n for n in names))
        self.assertTrue(any("General Secondary" in n for n in names))

    def test_02_grading_policy_crud_and_validation(self) -> None:
        # Invalid weights (not 100%)
        bad_policy = {
            "name": "Invalid Policy Test",
            "school_year": "2026-2027",
            "component_definitions": [
                {"name": "Written Work", "component_type": "Written Work", "weight": 40.0},
                {"name": "Performance Tasks", "component_type": "Performance Task", "weight": 40.0},
            ],
        }
        resp = self.client.post("/api/grading-policies", headers=self.teacher_headers, json=bad_policy)
        self.assertEqual(resp.status_code, 422)

        # Valid custom policy
        good_policy = {
            "name": "SHS TVL Track Policy",
            "description": "Custom policy for Senior High School TVL track",
            "school_year": "2026-2027",
            "grade_levels": "11,12",
            "subject_category": "TVL",
            "component_definitions": [
                {"name": "Written Work", "component_type": "Written Work", "weight": 20.0},
                {"name": "Performance Tasks", "component_type": "Performance Task", "weight": 60.0},
                {"name": "Quarterly Assessment", "component_type": "Quarterly Assessment", "weight": 20.0},
            ],
            "passing_grade": 75,
            "rounding_decimal_places": 2,
            "rounding_final_decimal_places": 0,
            "rounding_method": "half_up",
        }
        resp = self.client.post("/api/grading-policies", headers=self.teacher_headers, json=good_policy)
        self.assertEqual(resp.status_code, 200, resp.text)
        policy_id = resp.json()["id"]

        # Read back
        get_resp = self.client.get(f"/api/grading-policies/{policy_id}", headers=self.admin_headers)
        self.assertEqual(get_resp.status_code, 200)
        self.assertEqual(get_resp.json()["name"], "SHS TVL Track Policy")
        self.assertEqual(get_resp.json()["reporting_method"], "adjusted_2026")
        self.assertTrue(get_resp.json()["transmutation_table"])

        # Update
        good_policy["description"] = "Updated TVL policy description"
        good_policy["status"] = "Active"
        update_resp = self.client.put(f"/api/grading-policies/{policy_id}", headers=self.teacher_headers, json=good_policy)
        self.assertEqual(update_resp.status_code, 200)

        future_policy = {**good_policy, "name": "Direct Reporting Test", "school_year": "2027-2028",
                         "rounding_method": "floor", "rounding_final_decimal_places": 2,
                         "transmutation_table": [{"min": 0, "max": 100, "value": 88}]}
        future_resp = self.client.post("/api/grading-policies", headers=self.teacher_headers, json=future_policy)
        self.assertEqual(future_resp.status_code, 200, future_resp.text)
        future_details = self.client.get(f"/api/grading-policies/{future_resp.json()['id']}", headers=self.admin_headers).json()
        self.assertEqual(future_details["reporting_method"], "direct")
        self.assertIsNone(future_details["transmutation_table"])
        self.assertEqual(future_details["rounding_method"], "half_up")
        self.assertEqual(future_details["rounding_final_decimal_places"], 0)

    def test_03_gradebook_lifecycle_scores_and_calculations(self) -> None:
        # 1. Fetch academic structure
        struct_resp = self.client.get("/api/admin/academic-structure", headers=self.admin_headers)
        self.assertEqual(struct_resp.status_code, 200)
        struct = struct_resp.json()
        sy = struct["school_years"][0]
        gp = struct["grading_periods"][0]
        gl = struct["grade_levels"][0]
        sec = [s for s in struct["sections"] if s["grade_level_id"] == gl["id"]][0]
        subj = struct["subjects"][0]

        # Add 2 test students in this grade/section
        s1 = self.client.post("/api/persons", headers=self.admin_headers, json={
            "external_id": "TEST-STU-001", "lrn": "111111111111", "surname": "DELA CRUZ", "first_name": "JUAN",
            "sex": "Male", "date_of_birth": "2012-01-01", "role": "Student", "grade": gl["name"], "section": sec["name"],
            "guardian_phone": "+639170000001", "biometric_consent": False,
        })
        self.assertEqual(s1.status_code, 200, s1.text)
        s1_id = s1.json()["id"]

        s2 = self.client.post("/api/persons", headers=self.admin_headers, json={
            "external_id": "TEST-STU-002", "lrn": "111111111112", "surname": "SANTOS", "first_name": "MARIA",
            "sex": "Female", "date_of_birth": "2012-01-01", "role": "Student", "grade": gl["name"], "section": sec["name"],
            "guardian_phone": "+639170000002", "biometric_consent": False,
        })
        self.assertEqual(s2.status_code, 200, s2.text)
        s2_id = s2.json()["id"]

        # 2. Create Gradebook
        gb_create = self.client.post("/api/gradebooks", headers=self.teacher_headers, json={
            "school_year_id": sy["id"],
            "grading_period_id": gp["id"],
            "grade_level_id": gl["id"],
            "section_id": sec["id"],
            "subject_id": subj["id"],
        })
        self.assertEqual(gb_create.status_code, 200, gb_create.text)
        gb_id = gb_create.json()["id"]

        # 3. Load full gradebook
        gb_resp = self.client.get(f"/api/gradebooks/{gb_id}", headers=self.admin_headers)
        self.assertEqual(gb_resp.status_code, 200)
        gb_data = gb_resp.json()
        self.assertEqual(gb_data["gradebook"]["status"], "Draft")
        self.assertGreaterEqual(len(gb_data["components"]), 1)
        self.assertGreaterEqual(len(gb_data["students"]), 2)

        # 4. Update components to standard 3-component structure
        components_payload = [
            {
                "name": "Written Work",
                "component_type": "Written Work",
                "weight": 30.0,
                "items": [
                    {"label": "Quiz 1", "max_score": 20.0},
                    {"label": "Quiz 2", "max_score": 30.0},
                ],
            },
            {
                "name": "Performance Tasks",
                "component_type": "Performance Task",
                "weight": 50.0,
                "items": [
                    {"label": "Project 1", "max_score": 50.0},
                ],
            },
            {
                "name": "Quarterly Assessment",
                "component_type": "Quarterly Assessment",
                "weight": 20.0,
                "items": [
                    {"label": "Periodical Exam", "max_score": 50.0},
                ],
            },
        ]
        comp_update = self.client.put(f"/api/gradebooks/{gb_id}/components", headers=self.teacher_headers, json={
            "components": components_payload,
            "change_reason": "Configure Q1 assessment components",
        })
        self.assertEqual(comp_update.status_code, 200, comp_update.text)

        # Reload to get generated item IDs
        gb_data = self.client.get(f"/api/gradebooks/{gb_id}", headers=self.admin_headers).json()
        items = []
        for c in gb_data["components"]:
            for item in c["items"]:
                items.append(item)
        self.assertEqual(len(items), 4)

        q1_id, q2_id, proj_id, exam_id = [item["id"] for item in items]

        # 5. Save scores for student 1 (perfect scores: 20/20, 30/30, 50/50, 50/50 -> 100 transmuted)
        scores_payload = {
            "scores": {
                str(s1_id): [
                    {"assessment_item_id": q1_id, "score": 20.0, "status": "Scored"},
                    {"assessment_item_id": q2_id, "score": 30.0, "status": "Scored"},
                    {"assessment_item_id": proj_id, "score": 50.0, "status": "Scored"},
                    {"assessment_item_id": exam_id, "score": 50.0, "status": "Scored"},
                ],
                str(s2_id): [
                    # Student 2 is partially entered (missing exam)
                    {"assessment_item_id": q1_id, "score": 15.0, "status": "Scored"},
                    {"assessment_item_id": q2_id, "score": 20.0, "status": "Scored"},
                    {"assessment_item_id": proj_id, "score": 40.0, "status": "Scored"},
                    {"assessment_item_id": exam_id, "score": None, "status": "Missing"},
                ],
            },
            "change_reason": "Enter first batch of scores",
        }
        score_resp = self.client.put(f"/api/gradebooks/{gb_id}/scores", headers=self.teacher_headers, json=scores_payload)
        self.assertEqual(score_resp.status_code, 200, score_resp.text)

        cleared_column = self.client.put(f'/api/gradebooks/{gb_id}/scores', headers=self.teacher_headers,
                                        json={'scores': {str(person_id): [{'assessment_item_id': q2_id, 'score': None, 'status': 'Missing'}]
                                                         for person_id in (s1_id, s2_id)},
                                              'change_reason': 'Leave one assessment column unfinished'})
        self.assertEqual(cleared_column.status_code, 200, cleared_column.text)
        unused_column = self.client.get(f'/api/gradebooks/{gb_id}', headers=self.teacher_headers).json()
        self.assertEqual(unused_column['statistics']['incomplete'], 1)
        self.assertTrue(next(student for student in unused_column['students'] if student['person_id'] == s1_id)['complete'])
        self.assertFalse(next(student for student in unused_column['students'] if student['person_id'] == s2_id)['complete'])
        # A column unused by the whole class is excluded; once used, every learner needs it.
        used_column = self.client.put(f'/api/gradebooks/{gb_id}/scores', headers=self.teacher_headers,
                                      json={'scores': {str(s2_id): [{'assessment_item_id': q2_id, 'score': 20.0, 'status': 'Scored'}]},
                                            'change_reason': 'Use Quiz 2 for one learner while another is missing it'})
        self.assertEqual(used_column.status_code, 200, used_column.text)
        unfinished = self.client.get(f'/api/gradebooks/{gb_id}', headers=self.teacher_headers).json()
        self.assertEqual(unfinished['statistics']['incomplete'], 2)
        self.assertTrue(all(student['reported_grade'] is None for student in unfinished['students']))
        validation = self.client.post(f'/api/gradebooks/{gb_id}/validate', headers=self.teacher_headers)
        self.assertFalse(validation.json()['valid'])
        incomplete_issue = next(issue for issue in validation.json()['issues'] if issue['code'] == 'INCOMPLETE_LEARNERS')
        self.assertIn('DELA CRUZ, JUAN', incomplete_issue['message'])
        self.assertIn('Quiz 2', incomplete_issue['message'])
        self.assertIn('SANTOS, MARIA', incomplete_issue['message'])
        self.assertIn('Periodical Exam', incomplete_issue['message'])
        blocked = self.client.post(f'/api/gradebooks/{gb_id}/finalize', headers=self.teacher_headers,
                                   json={'reason': 'Try to finalize the unfinished record'})
        self.assertEqual(blocked.status_code, 409, blocked.text)
        self.assertEqual(blocked.json()['detail']['code'], 'INCOMPLETE_GRADEBOOK')
        self.assertIn('DELA CRUZ, JUAN', blocked.json()['detail']['message'])
        self.assertIn('Quiz 2', blocked.json()['detail']['message'])
        self.assertIn('SANTOS, MARIA', blocked.json()['detail']['message'])
        self.assertIn('Periodical Exam', blocked.json()['detail']['message'])
        draft_export = self.client.get(f'/api/gradebooks/{gb_id}/report.xlsx', headers=self.teacher_headers)
        self.assertEqual(draft_export.status_code, 409, draft_export.text)
        restored = self.client.put(f'/api/gradebooks/{gb_id}/scores', headers=self.teacher_headers, json=scores_payload)
        self.assertEqual(restored.status_code, 200, restored.text)

        # 6. Verify calculations
        gb_data = self.client.get(f"/api/gradebooks/{gb_id}", headers=self.admin_headers).json()
        s1_row = next(s for s in gb_data["students"] if s["person_id"] == s1_id)
        self.assertEqual(s1_row["initial_grade"], 100.0)
        self.assertEqual(s1_row["reported_grade"], 100)
        self.assertEqual(s1_row["status"], "Passing")
        self.assertTrue(s1_row["complete"])

        s2_row = next(s for s in gb_data["students"] if s["person_id"] == s2_id)
        self.assertIsNone(s2_row["reported_grade"])
        self.assertEqual(s2_row["status"], "Incomplete")
        self.assertFalse(s2_row["complete"])

        # 7. Check calculation breakdown modal data
        calc_resp = self.client.get(f"/api/gradebooks/{gb_id}/calculation/{s1_id}", headers=self.admin_headers)
        self.assertEqual(calc_resp.status_code, 200)
        calc_data = calc_resp.json()
        self.assertEqual(calc_data["student_name"], "DELA CRUZ, JUAN")
        self.assertEqual(len(calc_data["components"]), 3)
        self.assertEqual(calc_data["reported_grade"], 100)

        # 8. Workflow: Submit
        sub_resp = self.client.post(f"/api/gradebooks/{gb_id}/submit", headers=self.teacher_headers, json={
            "reason": "Ready for quarterly review",
        })
        self.assertEqual(sub_resp.status_code, 200)
        self.assertEqual(sub_resp.json()["status"], "Submitted")

        # 9. Finalize should fail because student 2 is incomplete
        fin_fail = self.client.post(f"/api/gradebooks/{gb_id}/finalize", headers=self.teacher_headers, json={
            "reason": "Finalizing Q1 gradebook",
        })
        self.assertEqual(fin_fail.status_code, 409)
        self.assertIn("incomplete", fin_fail.text)

        # Complete student 2's score
        self.client.put(f"/api/gradebooks/{gb_id}/scores", headers=self.teacher_headers, json={
            "scores": {
                str(s2_id): [
                    {"assessment_item_id": exam_id, "score": 35.0, "status": "Scored"},
                ],
            },
            "change_reason": "Enter missing exam for Student 2",
        })

        # 10. Finalize should now succeed
        fin_ok = self.client.post(f"/api/gradebooks/{gb_id}/finalize", headers=self.teacher_headers, json={
            "reason": "All scores entered and verified",
        })
        self.assertEqual(fin_ok.status_code, 200)
        self.assertEqual(fin_ok.json()["status"], "Finalized")

        self.assert_exports_match_gradebook(gb_id)

        # Editing scores while finalized must fail
        edit_blocked = self.client.put(f"/api/gradebooks/{gb_id}/scores", headers=self.teacher_headers, json={
            "scores": {str(s1_id): [{"assessment_item_id": q1_id, "score": 10.0, "status": "Scored"}]},
            "change_reason": "Unauthorized modification",
        })
        self.assertEqual(edit_blocked.status_code, 422)

        # 11. Lock gradebook
        lock_resp = self.client.post(f"/api/gradebooks/{gb_id}/lock", headers=self.records_headers, json={
            "reason": "Quarter grades archived",
        })
        self.assertEqual(lock_resp.status_code, 200)
        self.assertEqual(lock_resp.json()["status"], "Locked")

        # 12. Reopen with reason
        denied_reopen = self.client.post(f"/api/gradebooks/{gb_id}/reopen", headers=self.admin_headers, json={
            "reason": "Correcting student score following parent conference",
        })
        self.assertEqual(denied_reopen.status_code, 403, denied_reopen.text)
        reopen_resp = self.client.post(f"/api/gradebooks/{gb_id}/reopen", headers=self.records_headers, json={
            "reason": "Correcting student score following parent conference",
        })
        self.assertEqual(reopen_resp.status_code, 200)
        self.assertEqual(reopen_resp.json()["status"], "Draft")
        reopened = self.client.get(f"/api/gradebooks/{gb_id}", headers=self.teacher_headers).json()
        self.assertEqual(reopened['gradebook']['status'], 'Draft')
        reopened_student = next(student for student in reopened['students'] if student['person_id'] == s1_id)
        self.assertEqual(reopened_student['scores'][str(q1_id)], 20.0)
        edit_after_reopen = self.client.put(f"/api/gradebooks/{gb_id}/scores", headers=self.teacher_headers, json={
            "scores": {str(s1_id): [{"assessment_item_id": q1_id, "score": 19.0, "status": "Scored"}]},
            "change_reason": "Corrected after administrator reopened the record",
        })
        self.assertEqual(edit_after_reopen.status_code, 200, edit_after_reopen.text)

        # 13. Check audit history
        audit_resp = self.client.get(f"/api/gradebooks/{gb_id}/audit", headers=self.admin_headers)
        self.assertEqual(audit_resp.status_code, 200)
        audit_list = audit_resp.json()
        self.assertGreaterEqual(len(audit_list), 5)
        actions = [a["action"] for a in audit_list]
        self.assertIn("Create", actions)
        self.assertIn("Submit", actions)
        self.assertIn("Finalize", actions)
        self.assertIn("Lock", actions)
        self.assertIn("Reopen", actions)

    def test_04_grade_adjustment_request_flow(self) -> None:
        struct = self.client.get("/api/admin/academic-structure", headers=self.admin_headers).json()
        sy = struct["school_years"][0]
        gp = struct["grading_periods"][1] if len(struct["grading_periods"]) > 1 else struct["grading_periods"][0]
        gl = struct["grade_levels"][0]
        sec = [s for s in struct["sections"] if s["grade_level_id"] == gl["id"]][0]
        subj = struct["subjects"][1] if len(struct["subjects"]) > 1 else struct["subjects"][0]

        # Create gradebook
        gb_create = self.client.post("/api/gradebooks", headers=self.teacher_headers, json={
            "school_year_id": sy["id"], "grading_period_id": gp["id"],
            "grade_level_id": gl["id"], "section_id": sec["id"], "subject_id": subj["id"],
        })
        gb_id = gb_create.json()["id"]
        gb_data = self.client.get(f"/api/gradebooks/{gb_id}", headers=self.admin_headers).json()
        item_id = gb_data["components"][0]["items"][0]["id"]
        student_id = gb_data["students"][0]["person_id"]

        # Enter score
        self.client.put(f"/api/gradebooks/{gb_id}/scores", headers=self.teacher_headers, json={
            "scores": {str(student_id): [{"assessment_item_id": item_id, "score": 15.0, "status": "Scored"}]},
            "change_reason": "Initial score entry",
        })

        # The reference's correction workflow applies after finalization.
        complete_scores = {str(student['person_id']): [
            {'assessment_item_id': item['id'], 'score': 15.0 if item['id'] == item_id else item['max_score'], 'status': 'Scored'}
            for component in gb_data['components'] for item in component['items']
        ] for student in gb_data['students']}
        complete = self.client.put(f'/api/gradebooks/{gb_id}/scores', headers=self.teacher_headers,
                                  json={'scores': complete_scores, 'change_reason': 'Complete scores before finalization'})
        self.assertEqual(complete.status_code, 200, complete.text)
        finalized = self.client.post(f'/api/gradebooks/{gb_id}/finalize', headers=self.teacher_headers,
                                     json={'reason': 'All scores verified before requesting correction'})
        self.assertEqual(finalized.status_code, 200, finalized.text)

        # Request adjustment
        adj_resp = self.client.post(f"/api/gradebooks/{gb_id}/adjustments", headers=self.teacher_headers, json={
            "person_id": student_id,
            "assessment_item_id": item_id,
            "new_score": 18.0,
            "new_status": "Scored",
            "reason": "Teacher re-checked essay rubric and awarded +3 points",
        })
        self.assertEqual(adj_resp.status_code, 200)
        adj_id = adj_resp.json()["id"]

        # List adjustments
        list_adj = self.client.get(f"/api/gradebooks/{gb_id}/adjustments", headers=self.admin_headers)
        self.assertEqual(list_adj.status_code, 200)
        self.assertEqual(len(list_adj.json()), 1)
        self.assertEqual(list_adj.json()[0]["status"], "Pending")

        # Approve adjustment
        app_resp = self.client.put(f"/api/adjustment-requests/{adj_id}", headers=self.records_headers, json={
            "status": "Approved",
            "note": "Approved per submitted revised rubric",
        })
        self.assertEqual(app_resp.status_code, 200)
        self.assertEqual(app_resp.json()["status"], "Approved")

        # Verify score is now 18.0
        gb_after = self.client.get(f"/api/gradebooks/{gb_id}", headers=self.admin_headers).json()
        student_after = next(s for s in gb_after["students"] if s["person_id"] == student_id)
        self.assertEqual(student_after["scores"][str(item_id)], 18.0)


    def test_05_records_officer_only_reviews_teacher_adjustments(self):
        structure = self.client.get('/api/admin/academic-structure', headers=self.admin_headers).json()
        teacher_structure = self.client.get('/api/admin/academic-structure?context=grading', headers=self.teacher_headers).json()
        other_teacher_structure = self.client.get('/api/admin/academic-structure?context=grading', headers=self.other_teacher_headers).json()
        self.assertTrue(teacher_structure['sections'])
        self.assertEqual(other_teacher_structure['sections'], [])
        self.assertEqual(other_teacher_structure['grade_levels'], [])
        section = structure['sections'][0]
        period = next(row for row in structure['grading_periods'] if row['term_order'] == 3)
        context = {'school_year_id': period['school_year_id'], 'grading_period_id': period['id'],
                   'grade_level_id': section['grade_level_id'], 'section_id': section['id'], 'subject_id': structure['subjects'][0]['id']}
        books_before = self.client.get('/api/gradebooks', headers=self.admin_headers).json()
        empty = self.client.get('/api/gradebooks/lookup', headers=self.admin_headers, params=context)
        self.assertEqual(empty.status_code, 200, empty.text)
        self.assertIsNone(empty.json()['id'])
        self.assertEqual(len(self.client.get('/api/gradebooks', headers=self.admin_headers).json()), len(books_before))
        denied = self.client.post('/api/gradebooks', headers=self.admin_headers, json=context)
        self.assertEqual(denied.status_code, 403, denied.text)
        created = self.client.post('/api/gradebooks', headers=self.teacher_headers, json=context)
        self.assertEqual(created.status_code, 200, created.text)
        book_id = created.json()['id']
        admin_records = self.client.get('/api/gradebooks', headers=self.admin_headers)
        self.assertEqual(admin_records.status_code, 200, admin_records.text)
        self.assertIn(book_id, {book['id'] for book in admin_records.json()})
        lookup = self.client.get('/api/gradebooks/lookup', headers=self.admin_headers, params=context)
        self.assertEqual(lookup.status_code, 200, lookup.text)
        self.assertEqual(lookup.json()['id'], book_id)
        data = self.client.get(f'/api/gradebooks/{book_id}', headers=self.teacher_headers).json()
        self.assertEqual(data['gradebook']['grading_period_name'], period['name'])
        admin_view = self.client.get(f'/api/gradebooks/{book_id}', headers=self.admin_headers)
        self.assertEqual(admin_view.status_code, 200, admin_view.text)
        self.assertEqual(admin_view.json()['students'], data['students'])
        self.assertEqual(admin_view.json()['statistics'], data['statistics'])
        self.assertEqual(data['gradebook']['teacher_name'], self.teacher['full_name'])
        components = [{**component, 'items': component['items'] if component['component_type'] == 'Term Examinations' else component['items'][:1]}
                      for component in data['components']]
        scores = {str(student['person_id']): [{'assessment_item_id': activity['id'], 'score': activity['max_score'], 'status': 'Scored'}
                  for component in components for activity in component['items']] for student in data['students']}
        for path, payload in [('components', {'components': components, 'change_reason': 'Unauthorized admin assessment setup'}),
                              ('scores', {'scores': scores, 'change_reason': 'Unauthorized admin score modification'})]:
            denied = self.client.put(f'/api/gradebooks/{book_id}/{path}', headers=self.admin_headers, json=payload)
            self.assertEqual(denied.status_code, 403, denied.text)
        for action in ('submit', 'finalize', 'lock', 'reopen'):
            denied = self.client.post(f'/api/gradebooks/{book_id}/{action}', headers=self.admin_headers,
                                      json={'reason': 'Unauthorized admin class record action'})
            self.assertEqual(denied.status_code, 403, denied.text)
        for action in ('finalize', 'reopen'):
            denied = self.client.post(f'/api/gradebook/legacy-record/{action}', headers=self.admin_headers,
                                      json={'reason': 'Unauthorized legacy class record action'})
            self.assertEqual(denied.status_code, 403, denied.text)
        setup = self.client.put(f'/api/gradebooks/{book_id}/components', headers=self.teacher_headers,
                               json={'components': components, 'change_reason': 'Teacher configured quarterly activities'})
        self.assertEqual(setup.status_code, 200, setup.text)
        saved = self.client.put(f'/api/gradebooks/{book_id}/scores', headers=self.teacher_headers,
                               json={'scores': scores, 'change_reason': 'Teacher completed quarterly scores'})
        self.assertEqual(saved.status_code, 200, saved.text)
        for action in ('submit', 'finalize'):
            result = self.client.post(f'/api/gradebooks/{book_id}/{action}', headers=self.teacher_headers,
                                     json={'reason': 'Teacher verified complete term results'})
            self.assertEqual(result.status_code, 200, result.text)
        item = components[0]['items'][0]
        learner_id = data['students'][0]['person_id']
        proposal = {'person_id': learner_id, 'assessment_item_id': item['id'], 'new_score': item['max_score'] - 1,
                    'new_status': 'Scored', 'reason': 'Teacher reviewed the quiz rubric'}
        self.assertEqual(self.client.post(f'/api/gradebooks/{book_id}/adjustments', headers=self.admin_headers, json=proposal).status_code, 403)
        request = self.client.post(f'/api/gradebooks/{book_id}/adjustments', headers=self.teacher_headers, json=proposal)
        self.assertEqual(request.status_code, 200, request.text)
        request_id = request.json()['id']
        before = self.client.get(f'/api/gradebooks/{book_id}', headers=self.teacher_headers).json()
        self.assertEqual(before['students'][0]['scores'][str(item['id'])], item['max_score'])
        review_payload = {'status': 'Approved', 'note': 'Approved after reviewing the teacher rubric'}
        self.assertEqual(self.client.put(f'/api/adjustment-requests/{request_id}', headers=self.teacher_headers, json=review_payload).status_code, 403)
        self.assertEqual(self.client.put(f'/api/adjustment-requests/{request_id}', headers=self.admin_headers, json=review_payload).status_code, 403)
        self.assertEqual(self.client.delete(f'/api/adjustment-requests/{request_id}', headers=self.admin_headers).status_code, 403)
        listed = self.client.get('/api/admin/grade-adjustments', headers=self.admin_headers)
        self.assertEqual(listed.status_code, 200, listed.text)
        self.assertIn(request_id, {row['id'] for row in listed.json()})
        reviewed = self.client.put(f'/api/adjustment-requests/{request_id}', headers=self.records_headers, json=review_payload)
        self.assertEqual(reviewed.status_code, 200, reviewed.text)
        notifications = self.client.get('/api/teacher/grade-adjustment-notifications', headers=self.teacher_headers)
        self.assertEqual(notifications.status_code, 200, notifications.text)
        notification = next(item for item in notifications.json() if item['id'] == request_id)
        self.assertEqual(notification['status'], 'Approved')
        self.assertEqual(notification['grading_period_name'], period['name'])
        self.assertEqual(self.client.get('/api/teacher/grade-adjustment-notifications', headers=self.other_teacher_headers).json(), [])
        ignored = self.client.post('/api/teacher/grade-adjustment-notifications/read', headers=self.other_teacher_headers,
                                   json={'request_ids': [request_id]})
        self.assertEqual(ignored.status_code, 200, ignored.text)
        self.assertEqual(ignored.json()['read'], 0)
        self.assertIn(request_id, {item['id'] for item in self.client.get('/api/teacher/grade-adjustment-notifications', headers=self.teacher_headers).json()})
        self.assertEqual(self.client.post('/api/teacher/grade-adjustment-notifications/read', headers=self.teacher_headers,
                                          json={'request_ids': [request_id]}).status_code, 200)
        remaining = self.client.get('/api/teacher/grade-adjustment-notifications', headers=self.teacher_headers).json()
        self.assertEqual({item['id'] for item in remaining}, {item['id'] for item in notifications.json()} - {request_id})
        after = self.client.get(f'/api/gradebooks/{book_id}', headers=self.teacher_headers).json()
        self.assertEqual(after['students'][0]['scores'][str(item['id'])], proposal['new_score'])
        self.assertEqual(after['gradebook']['status'], 'Finalized')
        self.assertEqual(self.client.put(f'/api/gradebooks/{book_id}/scores', headers=self.teacher_headers,
                         json={'scores': scores, 'change_reason': 'Attempt direct edit after finalization'}).status_code, 422)
        locked = self.client.post(f'/api/gradebooks/{book_id}/lock', headers=self.records_headers,
                                 json={'reason': 'Approved term archived by records officer'})
        self.assertEqual(locked.status_code, 200, locked.text)
        locked_request = self.client.post(f'/api/gradebooks/{book_id}/adjustments', headers=self.teacher_headers, json={**proposal, 'new_score': item['max_score'] - 2})
        self.assertEqual(locked_request.status_code, 200, locked_request.text)
        reviewed = self.client.put(f"/api/adjustment-requests/{locked_request.json()['id']}", headers=self.records_headers, json=review_payload)
        self.assertEqual(reviewed.status_code, 200, reviewed.text)
        after = self.client.get(f'/api/gradebooks/{book_id}', headers=self.teacher_headers).json()
        self.assertEqual(after['gradebook']['status'], 'Locked')
        self.assertEqual(after['students'][0]['scores'][str(item['id'])], item['max_score'] - 2)
        rejected_request = self.client.post(f'/api/gradebooks/{book_id}/adjustments', headers=self.teacher_headers, json=proposal)
        self.assertEqual(rejected_request.status_code, 200, rejected_request.text)
        rejected = self.client.put(f"/api/adjustment-requests/{rejected_request.json()['id']}", headers=self.records_headers,
                                   json={'status': 'Rejected', 'note': 'Rubric does not support the requested change'})
        self.assertEqual(rejected.status_code, 200, rejected.text)
        rejection_notice = self.client.get('/api/teacher/grade-adjustment-notifications', headers=self.teacher_headers).json()
        self.assertIn('Rejected', {item['status'] for item in rejection_notice})
        final = self.client.get(f'/api/gradebooks/{book_id}', headers=self.teacher_headers).json()
        self.assertEqual(final['students'][0]['scores'][str(item['id'])], item['max_score'] - 2)
        score_before_criteria = final['students'][0]['scores'][str(item['id'])]
        criteria = [{'component_id': component['id'], 'weight': weight}
                    for component, weight in zip(final['components'], (25, 45, 30))]
        criteria_request = self.client.post(f'/api/gradebooks/{book_id}/criteria-adjustments', headers=self.teacher_headers,
                                            json={'components': criteria, 'reason': 'Teacher requests revised assessment criteria'})
        self.assertEqual(criteria_request.status_code, 200, criteria_request.text)
        criteria_id = criteria_request.json()['id']
        self.assertLess(criteria_id, 0)
        self.assertIn(criteria_id, {row['id'] for row in self.client.get('/api/admin/grade-adjustments', headers=self.admin_headers).json()})
        self.assertEqual(self.client.put(f'/api/adjustment-requests/{criteria_id}', headers=self.admin_headers,
                                         json={'status': 'Approved', 'note': 'Admin cannot approve criteria changes'}).status_code, 403)
        approved_criteria = self.client.put(f'/api/adjustment-requests/{criteria_id}', headers=self.records_headers,
                                            json={'status': 'Approved', 'note': 'Approved revised assessment criteria'})
        self.assertEqual(approved_criteria.status_code, 200, approved_criteria.text)
        after_criteria = self.client.get(f'/api/gradebooks/{book_id}', headers=self.teacher_headers).json()
        self.assertEqual([component['weight'] for component in after_criteria['components']], [25, 45, 30])
        self.assertEqual(after_criteria['students'][0]['scores'][str(item['id'])], score_before_criteria)
        self.assert_exports_match_gradebook(book_id)
        self.assertIn(criteria_id, {row['id'] for row in self.client.get('/api/teacher/grade-adjustment-notifications', headers=self.teacher_headers).json()})
        self.assertEqual(self.client.get(f'/api/gradebooks/{book_id}', headers=self.other_teacher_headers).status_code, 403)
        self.assertEqual(self.client.get('/api/gradebooks/lookup', headers=self.other_teacher_headers, params=context).status_code, 403)
        denied_creation = self.client.post('/api/gradebooks', headers=self.other_teacher_headers, json=context)
        self.assertEqual(denied_creation.status_code, 422, denied_creation.text)
        self.assertIn('not assigned', denied_creation.json()['detail'])
        self.assertEqual(self.client.get('/api/admin/grade-adjustments', headers=self.teacher_headers).status_code, 403)

    def test_06_admin_cannot_change_grading_policies(self):
        policy = self.client.get('/api/grading-policies', headers=self.admin_headers).json()[0]
        denied = self.client.post('/api/grading-policies', headers=self.admin_headers, json=policy)
        self.assertEqual(denied.status_code, 403, denied.text)
        denied = self.client.put(f"/api/grading-policies/{policy['id']}", headers=self.admin_headers, json=policy)
        self.assertEqual(denied.status_code, 403, denied.text)


if __name__ == "__main__":
    unittest.main()
