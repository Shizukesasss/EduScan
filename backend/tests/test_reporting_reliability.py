"""Export isolation and scheduler retry regressions using disposable school records."""
import hashlib
import json
import os
import sys
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import Mock, patch

TEST_DIR = tempfile.TemporaryDirectory(prefix="eduscan-reporting-tests-")
os.environ["DATABASE_URL"] = "sqlite://"
os.environ["EDUSCAN_DATA_DIR"] = TEST_DIR.name
os.environ["EDUSCAN_SECRET_KEY"] = "isolated-reporting-tests"
os.environ["ATTENDANCE_SCHEDULER_ENABLED"] = "false"
os.environ["SMS_GATEWAY_ENABLED"] = "false"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ["SF2_TEMPLATE_PATH"] = str(Path(__file__).resolve().parents[1] / "data/templates/School Form 2 (SF2) Daily Attendance Report of Learners.xlsx")

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from app.database import Base
from app.models import AttendanceEvent, Person, User
from app.services import gate_exports, reports, scheduler, sf2
from app.services.attendance import MANILA
from app.services.settings_store import set_json


class ReportingReliabilityTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(dir=TEST_DIR.name)
        self.root = Path(self.directory.name)
        self.engine = create_engine(f"sqlite:///{(self.root / 'test.db').as_posix()}", connect_args={"check_same_thread": False})
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine, expire_on_commit=False)
        self.teacher = User(username="teacher", full_name="Teacher", role="teacher", password_hash="unused")
        self.first = Person(external_id="FIRST", full_name="Learner One", sex="Male", role="Student", grade="10", section="One", enrollment_start_date=date(2026, 10, 1))
        self.second = Person(external_id="SECOND", full_name="Learner Two", sex="Female", role="Student", grade="10", section="Two", enrollment_start_date=date(2026, 10, 1))
        self.db.add_all([self.teacher, self.first, self.second]); self.db.commit()
        self.registry = self.root / "registered"; self.registry.mkdir()
        self.day = date(2026, 10, 12)
        self.now = datetime(2026, 10, 12, 10, 0, tzinfo=MANILA)
        self.patches = [patch.object(reports, "EXPORT_DIR", self.root), patch.object(reports, "REPORT_DIR", self.registry),
                        patch.object(sf2, "EXPORT_DIR", self.root),
                        patch.object(gate_exports, "SessionLocal", side_effect=lambda: Session(self.engine)),
                        patch.object(scheduler, "SessionLocal", side_effect=lambda: Session(self.engine)),
                        patch.object(scheduler, "dispatch_outbox")]
        for active in self.patches:
            active.start()
        with gate_exports._state_lock:
            gate_exports._pending_days.clear(); gate_exports._running = False

    def tearDown(self):
        for active in reversed(self.patches):
            active.stop()
        with gate_exports._state_lock:
            gate_exports._pending_days.clear(); gate_exports._running = False
        self.db.close(); self.engine.dispose(); self.directory.cleanup()

    def test_concurrent_sections_receive_their_own_registered_attendance_exports(self):
        barrier = threading.Barrier(2)
        def export(section):
            with Session(self.engine) as db:
                source = reports.generate_attendance_range_xlsx(db, self.day, self.day, grade="10", section=section)
            barrier.wait(timeout=10)
            with Session(self.engine) as db:
                record = reports.register_report(db, source, "Attendance Range", {"grade": "10", "section": section}, self.teacher)
                return source, Path(record.file_path), record.filename, record.file_sha256, json.loads(record.parameters_json)
        with ThreadPoolExecutor(max_workers=2) as executor:
            first, second = list(executor.map(export, ("One", "Two")))
        self.assertNotEqual(first[0], second[0])
        self.assertEqual(first[2], second[2])
        self.assertEqual(first[2], "Attendance-2026-10-12-to-2026-10-12.xlsx")
        for result, name, section in ((first, "Learner One", "One"), (second, "Learner Two", "Two")):
            with result[1].open("rb") as stream:
                workbook = load_workbook(stream)
                self.assertEqual(workbook["Attendance Detail"]["C5"].value, name)
                self.assertIsNone(workbook["Attendance Detail"]["C6"].value)
                workbook.close()
            self.assertEqual(result[3], hashlib.sha256(result[1].read_bytes()).hexdigest())
            self.assertEqual(result[4]["section"], section)

    def test_download_snapshot_is_not_overwritten_by_daily_refresh(self):
        stable = sf2.generate_temporary_log(self.db, self.day, "10", "One")
        snapshot = sf2.generate_temporary_log(self.db, self.day, "10", "One", snapshot=True)
        self.first.full_name = "Updated Test Name"; self.db.commit()
        refreshed = sf2.generate_temporary_log(self.db, self.day, "10", "One")
        second = sf2.generate_temporary_log(self.db, self.day, "10", "One", snapshot=True)
        self.assertEqual(stable, refreshed)
        self.assertEqual(snapshot.name, stable.name)
        self.assertNotEqual(snapshot, second)
        with snapshot.open("rb") as stream:
            workbook = load_workbook(stream)
            self.assertEqual(workbook.active["D2"].value, "Learner One")
            workbook.close()
        with second.open("rb") as stream:
            workbook = load_workbook(stream)
            self.assertEqual(workbook.active["D2"].value, "Updated Test Name")
            workbook.close()

    def test_concurrent_sf2_exports_keep_school_details_and_download_names(self):
        barrier = threading.Barrier(2)
        def export(school_name):
            with Session(self.engine) as db:
                path = sf2.generate_sf2(db, 2026, 10, "10", "One", "TEST", "2026-2027", school_name)
            barrier.wait(timeout=10)
            with path.open("rb") as stream:
                book = load_workbook(stream)
                value = book[sf2.SHEET_NAME]["C8"].value
                book.close()
            return path, value
        with ThreadPoolExecutor(max_workers=2) as executor:
            first, second = list(executor.map(export, ("School One", "School Two")))
        self.assertNotEqual(first[0], second[0])
        self.assertEqual(first[0].name, second[0].name)
        self.assertEqual((first[1], second[1]), ("School One", "School Two"))

    def test_failed_export_retries_on_next_tick_without_reclosing_attendance(self):
        set_json(self.db, "attendance", {"auto_close_enabled": True, "absence_cutoff": "09:00"})
        worker = scheduler.AttendanceScheduler()
        with patch.object(gate_exports, "generate_temporary_log", side_effect=[OSError("test disk failure"), self.root / "retried.xlsx"]) as output:
            with self.assertLogs(gate_exports.__name__, level="ERROR"):
                result = worker.run_once(self.now)
            self.assertEqual(result["created"], 2)
            self.assertIn(self.day, gate_exports._pending_days)
            self.assertEqual(output.call_count, 1)
            original_run = worker.run_once
            worker._stop = Mock()
            worker._stop.wait.side_effect = [False, True]
            with patch.object(worker, "run_once", side_effect=lambda: original_run(self.now)), patch.object(worker, "run_scheduled_backup"):
                worker._run()
            self.assertEqual(output.call_count, 2)
        self.assertEqual(len(self.db.scalars(select(AttendanceEvent)).all()), 2)
        self.assertEqual(gate_exports._pending_days, set())
        self.assertFalse(gate_exports._running)

    def test_one_failed_day_does_not_block_other_pending_days(self):
        later = self.day + timedelta(days=1)
        gate_exports._pending_days.update((self.day, later))
        def generate(_db, day):
            if day == self.day:
                raise OSError("test failure for one day")
        with patch.object(gate_exports, "generate_temporary_log", side_effect=generate) as output:
            with self.assertLogs(gate_exports.__name__, level="ERROR"):
                gate_exports.retry_gate_exports()
        self.assertEqual(output.call_count, 2)
        self.assertEqual(gate_exports._pending_days, {self.day})
        with patch.object(gate_exports, "generate_temporary_log") as output:
            gate_exports.retry_gate_exports()
        self.assertEqual(output.call_count, 1)
        self.assertFalse(gate_exports._pending_days)

    def test_partial_close_failure_still_exports_already_committed_rows(self):
        set_json(self.db, "attendance", {"auto_close_enabled": True})
        worker = scheduler.AttendanceScheduler()
        def close(db, day, _time):
            db.add(AttendanceEvent(id="partial-close", person_id=self.first.id, event_date=day,
                                   event_time=self.now.time(), direction="Daily Close", status="Absent", source="Test close"))
            db.commit()
            raise RuntimeError("failure after first committed row")
        observed = []
        def export(db, day):
            observed.append((day, len(db.scalars(select(AttendanceEvent)).all())))
        with patch.object(scheduler, "close_due_attendance", side_effect=close), patch.object(gate_exports, "generate_temporary_log", side_effect=export):
            with self.assertRaisesRegex(RuntimeError, "first committed row"):
                worker.run_once(self.now)
        self.assertEqual(observed, [(self.day, 1)])
        self.assertIsNotNone(self.db.get(AttendanceEvent, "partial-close"))

    def test_pending_exports_retry_even_when_automatic_closing_is_disabled(self):
        set_json(self.db, "attendance", {"auto_close_enabled": False})
        gate_exports._pending_days.add(self.day)
        worker = scheduler.AttendanceScheduler()
        worker._stop = Mock(); worker._stop.wait.side_effect = [False, True]
        original_run = worker.run_once
        with patch.object(gate_exports, "generate_temporary_log") as output, \
             patch.object(worker, "run_once", side_effect=lambda: original_run(self.now)), \
             patch.object(worker, "run_scheduled_backup"):
            worker._run()
        output.assert_called_once()
        self.assertFalse(gate_exports._pending_days)
        self.assertEqual(len(self.db.scalars(select(AttendanceEvent)).all()), 0)

    def test_scheduler_logs_failures_and_continues_other_jobs(self):
        worker = scheduler.AttendanceScheduler()
        worker._stop = Mock(); worker._stop.wait.side_effect = [False, True]
        with patch.object(scheduler, "retry_gate_exports", side_effect=RuntimeError("retry failure")) as retry, \
             patch.object(worker, "run_once", side_effect=RuntimeError("closing failure")) as close, \
             patch.object(scheduler, "dispatch_outbox", side_effect=RuntimeError("SMS failure")) as sms, \
             patch.object(worker, "run_scheduled_backup", side_effect=RuntimeError("backup failure")) as backup:
            with self.assertLogs(scheduler.__name__, level="ERROR") as logs:
                worker._run()
        self.assertEqual(len(logs.records), 4)
        for job in (retry, close, sms, backup):
            job.assert_called_once()


if __name__ == "__main__":
    unittest.main()
