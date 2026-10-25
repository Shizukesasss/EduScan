"""Run: .venv/Scripts/python -m unittest tests.test_jhs_export_parity.JHSExportParityTests -v

Only an in-memory database and temporary exports are used, never live records.
Checks exported inputs, independent arithmetic, saved results and template layout.
Desktop Excel's calculation engine is not exercised by these tests.
"""
import copy
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from tests import test_grading_regressions as fixtures
from tests.test_grading_regressions import grading, GradingPeriod, date, TEST_DIR
from app.services import reports
from app.services.class_record_export import subject_template, assessment_groups
from openpyxl import load_workbook

SUBJECTS = (
    ('ARALING PANLIPUNAN', False), ('BIOCHEMISTRY', True), ('Consumer Chemistry', True),
    ('ENGLISH', False), ('FILIPINO', False), ('SCIENCE', False),
    ('MAPEH', False), ('MATHEMATICS', False), ('RESEARCH 1', True),
    ('Research 2', True), ('Research 3', True), ('Research 4', True), ('ROBOTICS', True),
    ('TECHNOLOGY AND LIVELIHOOD EDUCATION', False), ('VALUES EDUCATION', False),
)
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


class JHSExportParityTests(unittest.TestCase):
    def test_core_export_transmutation_uses_unrounded_decimal_grade(self):
        for score, expected in ((71.18, 76), (71.175, 75)):
            fixture = fixtures.ClassRecordRegressionTest()
            fixture.setUp()
            path = None
            try:
                fixture.configure_subject_export()
                book = fixture.create(fixture.general)
                fixture.label_export_exams(book)
                for component in book.components:
                    for item in component.items:
                        item.max_score = 100
                fixture.db.commit()
                entries = [{'assessment_item_id': item.id, 'score': score, 'status': 'Scored'}
                           for component in book.components for item in component.items]
                for learner in (fixture.first, fixture.second):
                    grading.save_gradebook_scores(fixture.db, book.id, fixture.admin,
                        {str(learner.id): entries}, 'Synthetic decimal-boundary parity fixture')
                grading.finalize_gradebook(fixture.db, book.id, fixture.teacher, 'Isolated export boundary test')
                before = copy.deepcopy(grading.get_gradebook_full(fixture.db, book.id))
                with patch.object(reports, 'EXPORT_DIR', Path(TEST_DIR.name)):
                    path = reports.generate_gradebook_report_xlsx(fixture.db, book.id, 'V', 'V', '301874')
                values, formulas = load_workbook(path, data_only=True), load_workbook(path)
                try:
                    for row in (18, 69):
                        self.assertEqual(before['students'][0]['initial_grade'], 71.18)
                        self.assertEqual(values['TERM 1'][f'AB{row}'].value, 71.18)
                        self.assertEqual(values['TERM 1'][f'AC{row}'].value, expected)
                        formula = formulas['TERM 1'][f'AC{row}'].value
                        formula = formula.text if hasattr(formula, 'text') else formula
                        self.assertIn(f'ROUND(AB{row},12)', formula)
                    self.assertEqual(before, grading.get_gradebook_full(fixture.db, book.id))
                finally:
                    values.close()
                    formulas.close()
            finally:
                if path:
                    path.unlink(missing_ok=True)
                fixture.tearDown()

    def check_subject(self, name, elective, export_mode=3):
        fixture = fixtures.ClassRecordRegressionTest()
        fixture.setUp()
        try:
            kind, filename = subject_template(name, elective)
            fixture.configure_subject_export(kind)
            fixture.subject.name, fixture.subject.is_elective = name, elective
            fixture.section.track = 'STE' if elective else None
            fixture.db.commit()
            books, snapshots = [], []
            for term in (1, 2, 3):
                period = fixture.period
                if term != 1:
                    period = GradingPeriod(school_year_id=fixture.year.id, term_order=term, name=f'Term {term}',
                                           starts_on=date(2026, 6, 1), ends_on=date(2027, 3, 31))
                    fixture.db.add(period)
                    fixture.db.commit()
                book = fixture.create(fixture.general, grading_period_id=period.id)
                fixture.label_export_exams(book)
                for component in book.components:
                    for index, item in enumerate(component.items):
                        item.max_score = 30 + index * 10 if component == book.components[-1] else 20 + component.id % 3 * 10
                fixture.db.commit()
                for learner, offset in ((fixture.first, 0), (fixture.second, 0.05)):
                    entries = []
                    for index, component in enumerate(book.components):
                        for activity, item in enumerate(component.items):
                            ratio = (0.50 + activity * 0.10 if component == book.components[-1] else 0.60 + index * 0.03) + term * 0.05 + offset
                            if term == 3 and learner == fixture.second:
                                ratio = 1.0
                            entries.append({'assessment_item_id': item.id, 'score': round(item.max_score * ratio, 2), 'status': 'Scored'})
                    grading.save_gradebook_scores(fixture.db, book.id, fixture.admin, {str(learner.id): entries}, 'Synthetic varied scores and maxima')
                grading.finalize_gradebook(fixture.db, book.id, fixture.teacher, 'Isolated export parity test')
                books.append(book)
                snapshots.append(copy.deepcopy(grading.get_gradebook_full(fixture.db, book.id)))
            with patch.object(reports, 'EXPORT_DIR', Path(TEST_DIR.name)):
                selected = books[export_mode - 1] if export_mode in (1, 2, 3) else books[1]
                path = reports.generate_gradebook_report_xlsx(fixture.db, selected.id, 'V', 'V', '301874',
                        all_terms=export_mode == 3, through_term=2 if export_mode == 12 else None)
            values, formulas = load_workbook(path, data_only=True), load_workbook(path)
            try:
                self.assertEqual(values['INPUT DATA']['E24'].value, name)
                gmrc = kind == 'gmrc'
                hps, male_row, female_row, initial_col, term_col = (16, 19, 70, 'AV', 'AW') if gmrc else (15, 18, 69, 'AB', 'AC')
                columns = ({'ww_cognitive': 6, 'ww_affective': 14, 'pt_cognitive': 22,
                            'pt_affective': 28, 'pt_behavioral': 34, 'exam': 40} if gmrc else
                           {'ww': 6, 'pt': 14, 'exam': 20})
                for full in snapshots:
                    suffix = ' M and A' if kind == 'mapeh' else ''
                    sheet = values[f'TERM {full["gradebook"]["term_order"]}{suffix}']
                    included = full['gradebook']['term_order'] in ((1, 2, 3) if export_mode == 3 else (1, 2) if export_mode == 12 else (export_mode,))
                    if not included:
                        self.assertIsNone(sheet.cell(male_row, 6).value)
                        self.assertIsNone(sheet.cell(female_row, 6).value)
                        self.assertIsNone(sheet[f'{initial_col}{male_row}'].value)
                        self.assertIsNone(sheet[f'{term_col}{male_row}'].value)
                        continue
                    for student in full['students']:
                        row = male_row if student['sex'] == 'Male' else female_row
                        initial = 0
                        for group, component in assessment_groups(full, kind).items():
                            earned, maxima = [], []
                            for index, item in enumerate(item for item in component['items'] if item['active']):
                                col = columns[group] + index
                                score, maximum = sheet.cell(row, col).value, sheet.cell(hps, col).value
                                self.assertEqual(score, student['scores'][str(item['id'])])
                                self.assertEqual(maximum, item['max_score'])
                                earned.append(score)
                                maxima.append(maximum)
                            percentage = (sum(score / maximum * weight for score, maximum, weight in zip(earned, maxima, (30, 30, 40)))
                                          if group == 'exam' else sum(earned) / sum(maxima) * 100)
                            initial += percentage * component['weight'] / 100
                        self.assertAlmostEqual(initial, student['initial_grade'], places=2)
                        self.assertEqual(sheet[f'{initial_col}{row}'].value, student['initial_grade'])
                        self.assertEqual(sheet[f'{term_col}{row}'].value, student['reported_grade'])
                        converted = next(entry['value'] for entry in sorted(full['policy']['transmutation_table'], key=lambda entry: entry['min'], reverse=True) if initial >= entry['min'])
                        self.assertEqual(converted, student['reported_grade'])
                        self.assertEqual(formulas[sheet.title][f'{term_col}{row}'].data_type, 'f')
                template = reports.TEMPLATES_DIR / 'SUBJECT CLASS RECORDS' / filename
                with ZipFile(template) as source, ZipFile(path) as exported:
                    self.assertEqual(source.read('xl/styles.xml'), exported.read('xl/styles.xml'))
                    for part in source.namelist():
                        if part.startswith('xl/worksheets/') and part.endswith('.xml'):
                            old, new = ET.fromstring(source.read(part)), ET.fromstring(exported.read(part))
                            for tag in ('mergeCells', 'pageSetup', 'pageMargins'):
                                a, b = old.find(NS + tag), new.find(NS + tag)
                                self.assertEqual(ET.tostring(a) if a is not None else None, ET.tostring(b) if b is not None else None)
                            if 'HIGHEST POSSIBLE SCORE' in ''.join(new.itertext()) or part in ('xl/worksheets/sheet2.xml', 'xl/worksheets/sheet3.xml', 'xl/worksheets/sheet4.xml'):
                                for column in new.findall(NS + 'cols/' + NS + 'col'):
                                    width = float(column.get('width', 0))
                                    self.assertFalse(5 <= width < 8, (name, part, column.attrib))
                self.assertEqual(snapshots, [grading.get_gradebook_full(fixture.db, book.id) for book in books])
            finally:
                values.close()
                formulas.close()
                path.unlink(missing_ok=True)
        finally:
            fixture.tearDown()


def subject_test(name, elective):
    def run(self):
        self.check_subject(name, elective)
    return run


for index, (name, elective) in enumerate(SUBJECTS):
    setattr(JHSExportParityTests, f'test_{index + 1:02d}_{name.lower().replace(" ", "_")}', subject_test(name, elective))


def scoped_subject_test(name, mode):
    def run(self):
        self.check_subject(name, False, export_mode=mode)
    return run


for name in ('ENGLISH', 'MAPEH', 'VALUES EDUCATION'):
    for mode in (1, 2, 12):
        setattr(JHSExportParityTests, f'test_scope_{name.lower().replace(" ", "_")}_{mode}', scoped_subject_test(name, mode))
