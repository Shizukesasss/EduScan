import copy
import unittest
from tests import test_grading_regressions as fixtures
from app.services import grading

class LegacyJHSExportTests(unittest.TestCase):
    def test_ap_and_mapeh_preserve_saved_quarterly_policy(self):
        for name in ('ARALING PANLIPUNAN', 'MAPEH'):
            fixture = fixtures.ClassRecordRegressionTest()
            fixture.setUp()
            self.addCleanup(fixture.tearDown)
            with self.subTest(subject=name):
                fixture.grade.name='7'
                fixture.first.grade=fixture.second.grade='7'
                fixture.subject.name=name
                fixture.subject.category='JHS'
                fixture.db.commit()
                book=fixture.create(fixture.general)
                book.subject_name=name
                book.status='Draft'
                fixture.db.commit()
                fixture.score_student(book,80)
                fixture.score_second_student(book,70)
                book.status='Finalized'
                fixture.db.commit()
                before=copy.deepcopy(grading.get_gradebook_full(fixture.db,book.id))
                workbook=fixture.export(book)
                try:
                    sheet=workbook['TERM 1 M and A' if name=='MAPEH' else 'TERM 1']
                    self.assertEqual(sheet['AB12'].value,'Initial Grade')
                    self.assertEqual(sheet['F14'].value, 1)
                    self.assertEqual(sheet['N14'].value, 1)
                finally:workbook.close()
                from openpyxl import load_workbook
                values=load_workbook(fixture.export_path,data_only=True)
                try:
                    sheet=values[sheet.title]
                    initial=28
                    term=29
                    for row,student in zip((18,69),before['students']):
                        self.assertEqual(sheet.cell(row,initial).value,student['initial_grade'])
                        self.assertEqual(sheet.cell(row,term).value,student['reported_grade'])
                finally:values.close()
                self.assertEqual(before,grading.get_gradebook_full(fixture.db,book.id))
                if name in ('ARALING PANLIPUNAN', 'MAPEH'):
                    from zipfile import ZipFile
                    from xml.etree import ElementTree as ET
                    from app.services import reports
                    filename = 'MUSIC AND ARTS AND PE-HEALTH - Grade 7-10 Class Record.xlsx' if name == 'MAPEH' else 'ARALING PANLIPUNAN - Grade 7-10 Class Record.xlsx'
                    template = reports.TEMPLATES_DIR / 'SUBJECT CLASS RECORDS' / filename
                    with ZipFile(template) as source, ZipFile(fixture.export_path) as exported:
                        self.assertEqual(source.read('xl/styles.xml'), exported.read('xl/styles.xml'))
                        ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                        for part in source.namelist():
                            if part.startswith('xl/worksheets/') and part.endswith('.xml'):
                                old, new = ET.fromstring(source.read(part)), ET.fromstring(exported.read(part))
                                for tag in ('mergeCells', 'pageSetup', 'pageMargins'):
                                    a, b = old.find(ns + tag), new.find(ns + tag)
                                    self.assertEqual(ET.tostring(a) if a is not None else None, ET.tostring(b) if b is not None else None)
