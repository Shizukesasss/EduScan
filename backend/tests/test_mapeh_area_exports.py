import copy
import unittest
from datetime import date
from unittest.mock import patch
from pathlib import Path
from openpyxl import load_workbook
from fastapi import HTTPException
from tests import test_grading_regressions as f
from app.models import GradingPeriod
from app.services import grading,reports
from app.services.grading_engine import round_value

class MapehAreaExportTests(unittest.TestCase):
 def fixture(self,grade):
  x=f.ClassRecordRegressionTest();x.setUp();self.addCleanup(x.tearDown)
  x.grade.name=str(grade);x.first.grade=x.second.grade=str(grade);x.subject.name='MAPEH';x.subject.category='JHS';x.db.commit()
  return x
 def books(self,x,boundary_score=None):
  books={};snapshots={}
  for term in (1,2,3):
   period=x.period
   if term>1:
    period=GradingPeriod(school_year_id=x.year.id,quarter=term,name=f'Term {term}',starts_on=date(2026,6,1),ends_on=date(2027,3,31));x.db.add(period);x.db.commit()
   for area in ('music_arts','pe_health'):
    b=x.create(learning_area=area,grading_period_id=period.id)
    for c in b.components:
     for i,item in enumerate(c.items):item.max_score=100 if boundary_score is not None else 20+i*10
    x.db.commit()
    entries={}
    for n,student in enumerate((x.first,x.second)):
     entries[str(student.id)]=[{'assessment_item_id':item.id,'score':boundary_score if boundary_score is not None else round(item.max_score*(1 if n==0 else .4+term*.1+(i%2)*.03+(0.05 if area=='pe_health' else 0)),2),'status':'Scored'} for c in b.components for i,item in enumerate(c.items)]
    grading.save_gradebook_scores(x.db,b.id,x.teacher,entries,'Isolated area test')
    books[(term,area)]=b
   for area in ('music_arts','pe_health'):
    grading.finalize_gradebook(x.db,books[(term,area)].id,x.teacher,'Isolated finalization')
  snapshots={key:copy.deepcopy(grading.get_gradebook_full(x.db,b.id)) for key,b in books.items()}
  return books,snapshots
 def test_area_exports_use_unrounded_decimal_grade_at_transmutation_boundary(self):
  for boundary_score,expected in ((71.18,76),(71.175,75)):
   x=self.fixture(7);books,snapshots=self.books(x,boundary_score)
   with self.subTest(score=boundary_score),patch.object(reports,'EXPORT_DIR',Path(f.TEST_DIR.name)):
    path=reports.generate_gradebook_report_xlsx(x.db,books[(3,'music_arts')].id,'V','V','301874',allowed_teacher_id=x.teacher.id,all_terms=True)
    values=load_workbook(path,data_only=True);formulas=load_workbook(path)
    try:
     for term in (1,2,3):
      for area,suffix in (('music_arts','M and A'),('pe_health','PE and H')):
       full=snapshots[(term,area)];sheet=values[f'TERM {term} {suffix}'];formula_sheet=formulas[sheet.title]
       for student in full['students']:
        row=18 if student['sex']=='Male' else 69
        self.assertEqual(student['initial_grade'],71.18)
        self.assertEqual(sheet[f'AB{row}'].value,71.18)
        self.assertEqual(sheet[f'AC{row}'].value,expected)
        self.assertIn(f'ROUND(AB{row},12)',formula_sheet[f'AC{row}'].value)
     self.assertEqual(snapshots,{key:grading.get_gradebook_full(x.db,b.id) for key,b in books.items()})
    finally:values.close();formulas.close();path.unlink(missing_ok=True)
 def test_area_cannot_finalize_until_its_mapeh_partner_has_complete_grades(self):
  x=self.fixture(7)
  music=x.create(learning_area='music_arts')
  pe=x.create(learning_area='pe_health')
  x.db.add(f.SubjectTeachingAssignment(school_year_id=x.year.id,subject_id=x.subject.id,
                                       section_id=x.section.id,teacher_user_id=x.teacher.id))
  x.db.commit()
  entries={str(student.id):[{'assessment_item_id':item.id,'score':item.max_score,'status':'Scored'}
      for component in music.components for item in component.items]
      for student in (x.first,x.second)}
  grading.save_gradebook_scores(x.db,music.id,x.teacher,entries,'Isolated Music and Arts scores')

  validation=f.validate_gradebook_endpoint(music.id,x.db,x.teacher)
  self.assertFalse(validation['valid'])
  self.assertIn('MAPEH_PARTNER_INCOMPLETE',[issue['code'] for issue in validation['issues']])
  with self.assertRaisesRegex(ValueError,'complete Physical Education and Health grades'):
   grading.finalize_gradebook(x.db,music.id,x.teacher,'Attempt incomplete MAPEH finalization')
  self.assertEqual(music.status,'Draft')

  entries={str(student.id):[{'assessment_item_id':item.id,'score':item.max_score,'status':'Scored'}
      for component in pe.components for item in component.items]
      for student in (x.first,x.second)}
  grading.save_gradebook_scores(x.db,pe.id,x.teacher,entries,'Isolated PE and Health scores')
  self.assertTrue(f.validate_gradebook_endpoint(music.id,x.db,x.teacher)['valid'])
  grading.finalize_gradebook(x.db,music.id,x.teacher,'Finalize complete Music and Arts')
  grading.finalize_gradebook(x.db,pe.id,x.teacher,'Finalize complete PE and Health')
  self.assertEqual((music.status,pe.status),('Finalized','Finalized'))
 def test_partner_warnings_name_the_matching_area_and_term(self):
  from app.services.mapeh_learning_areas import partner_gradebook_issue
  x=self.fixture(7)
  periods={1:x.period}
  for term in (2,3):
   period=GradingPeriod(school_year_id=x.year.id,quarter=term,name=f'Term {term}',starts_on=date(2026,6,1),ends_on=date(2027,3,31))
   x.db.add(period);x.db.commit();periods[term]=period
  for term,period in periods.items():
   music=x.create(learning_area='music_arts',grading_period_id=period.id)
   pe=x.create(learning_area='pe_health',grading_period_id=period.id)
   with self.subTest(term=term,area='music_arts',partner='incomplete'):
    issue=partner_gradebook_issue(x.db,music)
    self.assertIn('Physical Education and Health',issue)
    self.assertIn(f'Term {term}',issue)
   with self.subTest(term=term,area='pe_health',partner='incomplete'):
    issue=partner_gradebook_issue(x.db,pe)
    self.assertIn('Music and Arts',issue)
    self.assertIn(f'Term {term}',issue)
   x.db.delete(pe);x.db.commit()
   with self.subTest(term=term,area='music_arts',partner='missing'):
    issue=partner_gradebook_issue(x.db,music)
    self.assertIn('Physical Education and Health',issue)
    self.assertIn(f'Term {term}',issue)
   x.db.delete(music);x.db.commit()
   with self.subTest(term=term,area='pe_health',partner='missing'):
    issue=partner_gradebook_issue(x.db,pe)
    self.assertIn('Music and Arts',issue)
    self.assertIn(f'Term {term}',issue)
 def test_grades_7_to_10_and_all_export_scopes(self):
  for grade in (7,8,9,10):
   x=self.fixture(grade);books,snapshots=self.books(x)
   scopes=[(3,(1,2,3),{'all_terms':True})]
   if grade==7:scopes += [(t,(t,),{}) for t in (1,2,3)]+[(2,(1,2),{'through_term':2})]
   for selected,terms,options in scopes:
    with self.subTest(grade=grade,terms=terms),patch.object(reports,'EXPORT_DIR',Path(f.TEST_DIR.name)):
     path=reports.generate_gradebook_report_xlsx(x.db,books[(selected,'music_arts')].id,'V','V','301874',allowed_teacher_id=x.teacher.id,**options)
     values=load_workbook(path,data_only=True);formulas=load_workbook(path)
     try:
      for term in terms:
       for area,suffix in [('music_arts','M and A'),('pe_health','PE and H')]:
        full=snapshots[(term,area)];sheet=values[f'TERM {term} {suffix}']
        for student,row in zip(full['students'],(18,69)):
         initial=0
         for c,start in zip(full['components'],(6,14,20)):
          items=[i for i in c['items'] if i['active']]
          scores=[sheet.cell(row,start+i).value for i in range(len(items))]
          maxima=[sheet.cell(15,start+i).value for i in range(len(items))]
          self.assertEqual(scores,[student['scores'][str(i['id'])] for i in items])
          percentage=sum(a/b*w for a,b,w in zip(scores,maxima,(30,30,40))) if c['component_type']=='Term Examinations' else sum(scores)/sum(maxima)*100
          initial+=percentage*c['weight']/100
         self.assertAlmostEqual(initial,student['initial_grade'],places=2)
         self.assertEqual(sheet[f'AB{row}'].value,student['initial_grade']);self.assertEqual(sheet[f'AC{row}'].value,student['reported_grade'])
         self.assertEqual(formulas[sheet.title][f'AC{row}'].data_type,'f')
      for student_index,row in ((0,15),(1,66)):
       combined=[]
       for term,col in ((1,'H'),(2,'K'),(3,'N')):
        if term in terms:
         result=round_value(sum(snapshots[(term,a)]['students'][student_index]['reported_grade'] for a in ('music_arts','pe_health'))/2,0,'half_up')
         self.assertEqual(values['FINAL GRADES'][f'{col}{row}'].value,result);combined.append(result)
         self.assertEqual(snapshots[(selected,'music_arts')]['mapeh_summary'][student_index]['terms'][term]['combined'],result)
        else:self.assertIsNone(values['FINAL GRADES'][f'{col}{row}'].value)
       expected=round_value(sum(combined)/3,0,'half_up') if len(combined)==3 else None
       self.assertEqual(values['FINAL GRADES'][f'O{row}'].value,expected)
      for t in set((1,2,3))-set(terms):
       self.assertIsNone(values[f'TERM {t} M and A']['F18'].value)
       self.assertIsNone(values[f'TERM {t} PE and H']['F18'].value)
     finally:values.close();formulas.close()
   self.assertEqual(snapshots,{key:grading.get_gradebook_full(x.db,b.id) for key,b in books.items()})
 def test_missing_or_incomplete_pair_blocks_export(self):
  x=self.fixture(7);books,_=self.books(x)
  books[(1,'pe_health')].status='Draft';x.db.commit()
  with self.assertRaises(HTTPException):reports.generate_gradebook_report_xlsx(x.db,books[(1,'music_arts')].id,'V','V','301874')

 def test_marked_legacy_export_excludes_other_parent_subjects(self):
  from app.models import Subject
  x=self.fixture(7);old=x.create(x.general);x.score_student(old,80);x.score_second_student(old,70)
  old.status='Finalized';old.is_legacy_readonly=True
  other=Subject(code='PE',name='Physical Education and Health',category='JHS');x.db.add(other);x.db.commit()
  pe=x.create(x.general,subject_id=other.id);x.score_student(pe,90);x.score_second_student(pe,80);pe.status='Finalized';x.db.commit()
  with patch.object(reports,'EXPORT_DIR',Path(f.TEST_DIR.name)):
   path=reports.generate_gradebook_report_xlsx(x.db,old.id,'V','V','301874')
   workbook=load_workbook(path,data_only=True)
   try:self.assertIsNone(workbook['TERM 1 PE and H']['F18'].value)
   finally:workbook.close()
