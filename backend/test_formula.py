import openpyxl
wb = openpyxl.load_workbook('templates/GRADE 7-10_MAPEH 1ST QUARTER.xlsx')
ws = wb['ARTS_Q1']
for row in range(12, 14):
    ws.cell(row=row, column=17).value = f'=IF(ISERROR(IF($P{row}="","",ROUND(($P{row}/SUMIF($F{row}:$O{row},"<>",$F$10:$O$10))*100,2))),"",IF($P{row}="","",ROUND(($P{row}/SUMIF($F{row}:$O{row},"<>",$F$10:$O$10))*100,2)))'
    ws.cell(row=row, column=30).value = f'=IF(ISERROR(IF($AC{row}="","",ROUND(($AC{row}/SUMIF($S{row}:$AB{row},"<>",$S$10:$AB$10))*100,2))),"",IF($AC{row}="","",ROUND(($AC{row}/SUMIF($S{row}:$AB{row},"<>",$S$10:$AB$10))*100,2)))'
wb.save('test_formula.xlsx')
print('Saved test_formula.xlsx')
