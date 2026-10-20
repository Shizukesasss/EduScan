"""Formula-backed exports for assessment setups that do not fit a fixed template."""
import math
from decimal import Decimal
from tempfile import TemporaryDirectory
from pathlib import Path

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select

from ..models import GradingPeriod, Subject
from ..models_grading import Gradebook
from .grading import get_gradebook_full
from .grading_engine import normalize_component_weights, calculate_student_grade, round_value


def _text(sheet, row, column, value):
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = 's'  # Names and activity labels must never become Excel formulas.
    return cell


def _round(expression, policy):
    function = {'floor': 'ROUNDDOWN', 'ceiling': 'ROUNDUP'}.get(policy['rounding_method'], 'ROUND')
    return f"{function}({expression},{policy['rounding_decimal_places']})"


def _term_sheet(workbook, full, cached, sheet_name=None):
    from .reports import _headers, _title
    gb, policy = full['gradebook'], full['policy']
    sheet = workbook.create_sheet(sheet_name or f"Term {gb['quarter']}")
    components = normalize_component_weights(full['components'])
    components = [{**component, 'items': [item for item in component['items'] if item.get('active')]}
                  for component in components]
    if not components or any(not component['items'] for component in components):
        raise HTTPException(409, 'Each grading component needs a used assessment activity before exporting.')
    students = sorted(full['students'], key=lambda student: (student['sex'] != 'Male', student['full_name'].casefold(), student['person_id']))
    final_column = 6 + sum(len(component['items']) + 3 for component in components)
    if final_column + 2 > 16384:
        raise HTTPException(409, 'This assessment setup exceeds the Excel column limit.')
    _title(sheet, 'CLASS RECORD', f"{gb['subject_name']} · {gb['school_year_name']} · Term {gb['quarter']}", final_column)
    # Metadata is literal text, including subjects whose names begin with '='.
    _text(sheet, 2, 1, f"{gb['subject_name']} · {gb['school_year_name']} · Term {gb['quarter']}")
    details = [('School', 'San Jose National High School'), ('Grade & section', f"Grade {gb['grade_name']} — {gb['section_name']}"),
               ('Teacher', gb['teacher_name']), ('School head', gb.get('school_head') or 'Not configured')]
    for row, (label, value) in enumerate(details, 3):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        _text(sheet, row, 1, label).font = Font(bold=True)
        sheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=min(final_column, 8))
        _text(sheet, row, 4, value)
    sheet.merge_cells('A7:C7')
    _text(sheet, 7, 1, 'Region / Division / School ID')
    sheet.merge_cells('D7:H7')
    sheet['D7'] = "='INPUT DATA'!E10&\" / \"&'INPUT DATA'!E11&\" / \"&'INPUT DATA'!E13"
    sheet.merge_cells(start_row=8, start_column=1, end_row=8, end_column=final_column)
    _text(sheet, 8, 1, 'Saved gradebook setup; spreadsheet edits do not update EduScan.')
    _headers(sheet, 10, ['No.', 'Learner', 'LRN / ID', 'Sex'])
    _text(sheet, 11, 2, 'HIGHEST POSSIBLE SCORE')
    caches = cached.setdefault(sheet.title, {})
    blocks = []
    column = 5
    for component in components:
        items = component['items']
        first, last = column, column + len(items) - 1
        total, ps, ws = last + 1, last + 2, last + 3
        sheet.merge_cells(start_row=9, start_column=first, end_row=9, end_column=ws)
        group = _text(sheet, 9, first, component['name'])
        group.font = Font(bold=True, color='FFFFFF')
        group.fill = PatternFill('solid', fgColor='1B3968')
        group.alignment = Alignment(horizontal='center', wrap_text=True)
        for offset, item in enumerate(items):
            col = get_column_letter(first + offset)
            _text(sheet, 10, first + offset, item['label'])
            sheet.cell(11, first + offset, item['max_score'])
            validation = DataValidation(type='decimal', operator='between', formula1='0', formula2=f'{col}$11')
            validation.errorTitle, validation.error = 'Invalid score', 'Enter a score between zero and the activity maximum.'
            validation.showErrorMessage, validation.errorStyle = True, 'stop'
            sheet.add_data_validation(validation)
            validation.add(f'{col}12:{col}{11 + len(students)}')
        for col, label in ((total, 'TOTAL'), (ps, 'PS'), (ws, 'WS')):
            _text(sheet, 10, col, label)
        start, end = get_column_letter(first), get_column_letter(last)
        sheet.cell(11, total, f'=SUM({start}11:{end}11)')
        caches[f'{get_column_letter(total)}11'] = sum(item['max_score'] for item in items)
        sheet.cell(11, ps, 100)
        sheet.cell(11, ws, component['weight'] / 100).number_format = '0%' if float(component['weight']).is_integer() else '0.00%'
        blocks.append((component, first, last, total, ps, ws))
        column = ws + 1
    initial_col, grade_col, raw_col = column, column + 1, column + 2
    for col, label in ((initial_col, 'INITIAL GRADE'), (grade_col, 'TERM GRADE'), (raw_col, 'Unrounded initial grade')):
        _text(sheet, 10, col, label)
    sheet.column_dimensions[get_column_letter(raw_col)].hidden = True
    for index in range(len(components)):
        sheet.column_dimensions[get_column_letter(raw_col + 1 + index)].hidden = True
    table = policy.get('transmutation_table') or []
    template_exams = any(component.get('component_type') == 'Term Examinations' for component in components)
    chunk_size = max(1, min(16, 64 // len(components))) if template_exams else 40
    if raw_col + len(components) + max(0, (len(table) - 1) // chunk_size) > 16384:
        raise HTTPException(409, 'This assessment setup exceeds the Excel column limit.')
    if template_exams:
        table = sorted(table, key=lambda entry: float(entry['min']), reverse=True)
    table_row = 15 + len(students)
    _text(sheet, table_row - 1, 1, f"Grade conversion · rounding: {policy['rounding_method']} ({policy['rounding_decimal_places']} decimals)")
    _headers(sheet, table_row, ['Minimum initial grade', 'Maximum initial grade', 'Term grade'])
    for index, entry in enumerate(table, table_row + 1):
        for col, key in enumerate(('min', 'max', 'value'), 1):
            sheet.cell(index, col, float(entry[key]))
    for row, student in enumerate(students, 12):
        for col, value in enumerate((row - 11, student['full_name'], student.get('lrn') or student['external_id'], student['sex']), 1):
            _text(sheet, row, col, value)
        expressions, guards, scores, percentages = [], [], [], []
        for block_index, (component, first, last, total, ps, ws) in enumerate(blocks):
            start, end = get_column_letter(first), get_column_letter(last)
            score_range = f'{start}{row}:{end}{row}'
            for offset, item in enumerate(component['items']):
                value = student['scores'].get(str(item['id']))
                status = student['score_statuses'].get(str(item['id']))
                if status != 'Scored' or value is None or not math.isfinite(value) or not 0 <= value <= item['max_score']:
                    raise HTTPException(409, 'A required exported assessment score is missing or invalid.')
                sheet.cell(row, first + offset, value)
                scores.append({'item_id': item['id'], 'score': value, 'status': status})
            scored_maxima = ','.join(f'IF(ISNUMBER({get_column_letter(col)}{row}),{get_column_letter(col)}$11,0)' for col in range(first, last + 1))
            percentage = f'SUM({score_range})/SUM({scored_maxima})*100'
            if component.get('component_type') == 'Term Examinations':
                portions = {'ST1': 30, 'ST2': 30, 'TE': 40}
                if {item['label'] for item in component['items']} != set(portions) or len(component['items']) != 3:
                    raise HTTPException(409, 'Term Examinations requires ST1, ST2, and TE to preserve the saved calculation.')
                terms, used = [], []
                for offset, item in enumerate(component['items']):
                    col = get_column_letter(first + offset)
                    terms.append(f'IF(ISNUMBER({col}{row}),{col}{row}/{col}$11*{portions[item["label"]]},0)')
                    used.append(f'IF(ISNUMBER({col}{row}),{portions[item["label"]]},0)')
                used_total, earned_total = f'SUM({",".join(used)})', f'SUM({",".join(terms)})'
                percentage = f'IF({used_total}=100,{earned_total},{earned_total}*100/{used_total})'
            raw_ps_column = raw_col + 1 + block_index
            raw_ps = f'{get_column_letter(raw_ps_column)}{row}'
            sheet.cell(row, raw_ps_column, f'=IF(COUNT({score_range})=0,"",{percentage})')
            weighted = f'{raw_ps}*({get_column_letter(ws)}$11*100)/100'
            percentages.append((raw_ps, f'{get_column_letter(ws)}$11'))
            sheet.cell(row, total, f'=IF(COUNT({score_range})=0,"",SUM({score_range}))')
            sheet.cell(row, ps, f'=IF({raw_ps}="","",{_round(raw_ps, policy)})')
            sheet.cell(row, ws, f'=IF(COUNT({score_range})=0,"",{_round(weighted, policy)})')
            expressions.append(weighted)
            guards.append(f'COUNT({score_range})={len(component["items"])}')
        raw_ref, initial_ref = f'{get_column_letter(raw_col)}{row}', f'{get_column_letter(initial_col)}{row}'
        weights = ','.join(f'{get_column_letter(block[5])}$11' for block in blocks)
        guards.append(f'ABS(SUM({weights})-1)<=0.0001')
        sheet.cell(row, raw_col, f'=IF(AND({",".join(guards)}),SUM({",".join(expressions)}),"")')
        sheet.cell(row, initial_col, f'=IF({raw_ref}="","",{_round(raw_ref, policy)})')
        grade_input = f'MAX(0,MIN(100,{raw_ref if template_exams else initial_ref}))'
        if table:
            # Keep first-matching-range semantics without array formulas or Excel's nesting limit.
            converted = '""' if template_exams else f'IF({grade_input}>=100,100,"")'
            for chunk in reversed(range(0, len(table), chunk_size)):
                for index in reversed(range(chunk, min(chunk + chunk_size, len(table)))):
                    conversion_row = table_row + 1 + index
                    condition = f'{grade_input}>=$A${conversion_row}'
                    if template_exams:
                        # Scale differences before summing so tiny cutoff differences survive spreadsheet precision tolerances.
                        terms = ','.join(f'({ref}-$A${conversion_row})*({weight}*100)*100000000' for ref, weight in percentages)
                        weight_sum = ','.join(f'{weight}*100' for _, weight in percentages)
                        difference = f'SUM({terms})+$A${conversion_row}*(SUM({weight_sum})-100)*100000000'
                        condition = f'IF(AND($A${conversion_row}>0,$A${conversion_row}<=100),{difference}>=0,{condition})'
                    else:
                        condition = f'AND({condition},{grade_input}<=$B${conversion_row})'
                    converted = f'IF({condition},$C${conversion_row},{converted})'
                if chunk:
                    helper_column = raw_col + len(components) + chunk // chunk_size
                    helper = get_column_letter(helper_column)
                    sheet.cell(row, helper_column, f'=IF({initial_ref}="","",{converted})')
                    sheet.column_dimensions[helper].hidden = True
                    converted = f'{helper}{row}'
        else:
            converted = f'IF({grade_input}>=100,100,IF({grade_input}>=60,INT(({grade_input}-60+0.000000001)/1.6)+75,INT(({grade_input}+0.000000001)/4)+60))'
        sheet.cell(row, grade_col, f'=IF({initial_ref}="","",{converted})')
        calculated = calculate_student_grade(scores, components, {component['id']: component['items'] for component in components},
                                             policy.get('transmutation_table'), gb['passing_grade'], policy['rounding_decimal_places'],
                                             policy['rounding_final_decimal_places'], policy['rounding_method'])
        if calculated['initial_grade'] != student['initial_grade'] or calculated['reported_grade'] != student['reported_grade']:
            raise HTTPException(409, 'Export calculations do not match the saved gradebook. No grades were changed.')
        for block_index, (block, detail, result) in enumerate(zip(blocks, calculated['component_details'], calculated['component_grades'])):
            caches[f'{get_column_letter(raw_col + 1 + block_index)}{row}'] = detail['percentage']
            for col, value in ((block[3], detail['total_earned']), (block[4], result['percentage']), (block[5], result['weighted_score'])):
                caches[f'{get_column_letter(col)}{row}'] = value
        caches[raw_ref] = float(sum(Decimal(str(detail['percentage'])) * Decimal(str(detail['weight'])) / 100 for detail in calculated['component_details']))
        caches[initial_ref] = calculated['initial_grade']
        caches[f'{get_column_letter(grade_col)}{row}'] = calculated['reported_grade']
    for cells in sheet.iter_rows(min_row=10, max_row=11, max_col=grade_col):
        for cell in cells:
            cell.font = Font(bold=True, color='1B3968')
            cell.fill = PatternFill('solid', fgColor='DCE7F5')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for cells in sheet.iter_rows(min_row=12, max_row=11 + len(students), min_col=5, max_col=grade_col):
        for cell in cells:
            cell.number_format = '0.' + '0' * policy['rounding_decimal_places'] if policy['rounding_decimal_places'] else '0'
    for col in range(1, grade_col + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 12
    sheet.column_dimensions['A'].width = 8
    sheet.column_dimensions['B'].width = 34
    sheet.column_dimensions['C'].width = 20
    sheet.row_dimensions[9].height, sheet.row_dimensions[10].height = 30, 42
    sheet.freeze_panes, sheet.print_title_rows = 'E12', '1:11'
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = 'landscape'
    return sheet


def _combined_arts_health(workbook, terms, sheets, students, cached):
    from .reports import _headers, _title
    sheet = workbook.create_sheet('FINAL GRADES')
    periods = sorted({1, 2, 3} | {key[0] for key in terms})
    headings = ['Learner', 'LRN / ID']
    for period in periods:
        headings.extend((f'Term {period} Music and Arts', f'Term {period} PE and Health', f'Term {period} Combined'))
    headings.append('FINAL GRADE')
    _title(sheet, 'MUSIC AND ARTS / PE AND HEALTH', 'Combined term grades require both learning areas.', len(headings))
    _headers(sheet, 4, headings)
    caches = cached.setdefault(sheet.title, {})
    for row, student in enumerate(sorted(students, key=lambda item: (item['sex'] != 'Male', item['full_name'].casefold(), item['person_id'])), 5):
        _text(sheet, row, 1, student['full_name'])
        _text(sheet, row, 2, student.get('lrn') or student['external_id'])
        combined_refs, combined_grades = [], []
        for index, period in enumerate(periods):
            grades = []
            first = 3 + index * 3
            for offset, area in enumerate(('M and A', 'PE and H')):
                term, source = terms.get((period, area)), sheets.get((period, area))
                if term is None:
                    continue
                roster = sorted(term['students'], key=lambda item: (item['sex'] != 'Male', item['full_name'].casefold(), item['person_id']))
                match = next(((i, item) for i, item in enumerate(roster, 12) if item['person_id'] == student['person_id']), None)
                if match is None:
                    continue
                grade_col = next(cell.column_letter for cell in source[10] if cell.value == 'TERM GRADE')
                reference = f"'{source.title}'!{grade_col}{match[0]}"
                sheet.cell(row, first + offset, f'=IF({reference}="","",{reference})')
                caches[f'{get_column_letter(first + offset)}{row}'] = match[1]['reported_grade']
                grades.append(match[1]['reported_grade'])
            pair = f'{get_column_letter(first)}{row}:{get_column_letter(first + 1)}{row}'
            ref = f'{get_column_letter(first + 2)}{row}'
            sheet[ref] = f'=IF(COUNT({pair})=2,ROUND(AVERAGE({pair}),0),"")'
            combined_refs.append(ref)
            if len(grades) == 2:
                caches[ref] = round_value(sum(grades) / 2, 0)
                combined_grades.append(caches[ref])
        reference = f'{get_column_letter(len(headings))}{row}'
        sheet[reference] = f'=IF(COUNT({",".join(combined_refs)})={len(periods)},ROUND(AVERAGE({",".join(combined_refs)}),0),"")'
        if len(combined_grades) == len(periods):
            caches[reference] = round_value(sum(combined_grades) / len(periods), 0)
    for col in range(1, len(headings) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 20 if col > 1 else 34
    sheet.row_dimensions[4].height = 42
    sheet.freeze_panes = 'C5'


def export_adaptive(db, book, full, region, division, school_id, allowed_teacher_id=None, all_terms=False):
    from .reports import _complete_finalized_record, _write_class_record_workbook, new_report_path, EXPORT_DIR, _safe_name
    if allowed_teacher_id is not None and book.teacher_user_id != allowed_teacher_id:
        raise HTTPException(403, 'This saved class record belongs to another teacher.')
    if not _complete_finalized_record(full):
        raise HTTPException(409, 'Finalize the complete class record before exporting.')
    if not region.strip() or not division.strip() or not school_id.strip().isdigit():
        raise HTTPException(422, 'Enter the region, division, and numeric school ID before exporting.')
    workbook = Workbook()
    inputs = workbook.active
    inputs.title = 'INPUT DATA'
    gb = full['gradebook']
    for row, label, value in ((10, 'REGION', region), (11, 'DIVISION', division), (13, 'SCHOOL ID', school_id),
                              (14, 'SCHOOL NAME', 'San Jose National High School'), (15, 'SCHOOL YEAR', gb['school_year_name']),
                              (16, 'SCHOOL HEAD', gb.get('school_head') or ''), (23, 'TEACHER', gb['teacher_name']),
                              (24, 'SUBJECT', gb['subject_name']), (25, 'GRADE', gb['grade_name']), (26, 'SECTION', gb['section_name'])):
        _text(inputs, row, 2, label).font = Font(bold=True)
        _text(inputs, row, 5, value)
    inputs.column_dimensions['B'].width, inputs.column_dimensions['E'].width = 24, 42
    from .class_record_export import subject_template
    def area(subject):
        try:
            kind, _ = subject_template(subject)
        except HTTPException:
            return None
        return ('PE and H' if 'HEALTH' in subject.upper() else 'M and A') if kind == 'mapeh' else None
    combined = not all_terms and area(gb['subject_name']) is not None
    query = select(Gradebook).where(Gradebook.school_year_id == book.school_year_id, Gradebook.grade_level_id == book.grade_level_id,
                                   Gradebook.section_id == book.section_id,
                                   Gradebook.grading_period.has(GradingPeriod.deleted_at.is_(None)), Gradebook.subject.has(Subject.deleted_at.is_(None)))
    if not combined:
        query = query.where(Gradebook.subject_id == book.subject_id)
    if allowed_teacher_id is not None:
        query = query.where(Gradebook.teacher_user_id == allowed_teacher_id)
    if all_terms:
        query = query.where(Gradebook.teacher_user_id == book.teacher_user_id, Gradebook.quarter.in_((1, 2, 3)))
    terms = {}
    for sibling in db.scalars(query.order_by(Gradebook.quarter)).all():
        learning_area = area(sibling.subject_name) if combined else ''
        if combined and learning_area is None:
            continue
        term = full if sibling.id == book.id else get_gradebook_full(db, sibling.id)
        if not _complete_finalized_record(term):
            continue
        key = (sibling.quarter, learning_area)
        if key in terms:
            raise HTTPException(409, 'Multiple finalized records match this subject and term. Resolve the assignment before exporting.')
        terms[key] = term
    if not any(term['gradebook']['id'] == book.id for term in terms.values()):
        raise HTTPException(409, 'The selected class record no longer belongs to a current subject and grading period.')
    if all_terms and set(terms) != {(1, ''), (2, ''), (3, '')}:
        raise HTTPException(409, 'All three terms must be complete and finalized or locked for this teacher, subject, section, and school year before exporting together.')
    cached, sheets = {}, {}
    for key, term in terms.items():
        sheet_name = f'Term {key[0]} {key[1]}' if combined else None
        sheets[key] = _term_sheet(workbook, term, cached, sheet_name)
    if combined:
        _combined_arts_health(workbook, terms, sheets, full['students'], cached)
    for sheet in workbook:
        for row in sheet:
            if any(cell.data_type == 'f' and len(cell.value) > 8192 for cell in row):
                raise HTTPException(409, 'This assessment setup exceeds the Excel formula length limit.')
    period_name = 'AllTerms' if all_terms else f"Term{gb['quarter']}"
    target = new_report_path(EXPORT_DIR, f"ClassRecord-{_safe_name(gb['school_year_name'])}-{period_name}-Grade{_safe_name(gb['grade_name'])}-{_safe_name(gb['section_name'])}-{_safe_name(gb['subject_name'])}.xlsx")
    with TemporaryDirectory(dir=target.parent) as directory:
        source = Path(directory) / 'gradebook.xlsx'
        workbook.save(source)
        _write_class_record_workbook(source, target, {name: {} for name in cached}, formula_values=cached)
    return target
