from __future__ import annotations

import hashlib
import html
import io
import re
import uuid
from datetime import date, datetime, timedelta
from types import SimpleNamespace

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from sqlalchemy import select

from ..models import (AnnexGTermComment, GradeLevel, GradingPeriod, LearnerEnrollment,
                      Person, RosterImportAudit, SchoolSection, SchoolYear,
                      Subject, SubjectTeachingAssignment, User)
from ..models_grading import Gradebook
from .academic_context import normalize_grade_name, subject_matches_grade
from .attendance import is_instructional_day, local_now, rows_for_people
from .grading import get_gradebook_full
from .grading_engine import round_value
from .settings_store import get_json


def _header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().casefold().replace("_", " "))


def _value_date(value: object, field: str, row_number: int) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as error:
        raise ValueError(f"row {row_number}: {field} must be an Excel date or YYYY-MM-DD") from error


def _records(data: bytes, required: dict[str, set[str]], optional: dict[str, set[str]] | None = None) -> list[tuple[int, dict[str, object]]]:
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Workbook exceeds 15 MB")
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
    except Exception as error:
        raise HTTPException(status_code=422, detail=f"Workbook could not be read: {error}") from error
    indexes: dict[str, int] = {}
    fields = {**required, **(optional or {})}
    for index, value in enumerate(headers):
        name = _header(value)
        for field, names in fields.items():
            if name in names:
                indexes[field] = index
                break
    missing = set(required) - set(indexes)
    if missing:
        book.close()
        raise HTTPException(status_code=422, detail=f"Workbook is missing columns: {', '.join(sorted(missing))}")
    result = []
    for row_number, values in enumerate(rows, start=2):
        item = {field: values[index] if index < len(values) else None for field, index in indexes.items()}
        if any(value not in (None, "") for value in item.values()):
            result.append((row_number, item))
    book.close()
    return result


def _approval(approved_reference: str) -> str:
    reference = approved_reference.strip()
    if len(reference) < 5:
        raise HTTPException(status_code=422, detail="Enter the school approval/document reference for this file")
    return reference


def _person_by_lrn(db: Session, value: object, row_number: int) -> Person:
    lrn = str(value or "").strip()
    if not lrn:
        raise ValueError(f"row {row_number}: exact LRN is required; name matching is not accepted")
    person = db.scalar(select(Person).where(Person.lrn == lrn))
    if not person or person.role != "Student":
        raise ValueError(f"row {row_number}: no existing learner has exact LRN {lrn}")
    return person


def import_annex_g_enrollments(db: Session, data: bytes, filename: str, approved_reference: str,
                               actor: User, dry_run: bool = True) -> dict:
    reference = _approval(approved_reference)
    rows = _records(data, {
        "lrn": {"lrn", "learner reference number"},
        "school_year": {"school year", "schoolyear"},
        "grade": {"grade", "grade level"},
        "section": {"section"},
        "starts_on": {"enrollment start", "enrollment start date", "start date"},
    }, optional={"ends_on": {"enrollment end", "enrollment end date", "end date"}})
    errors: list[dict] = []
    entries = []
    seen: set[tuple[int, int, date]] = set()
    for row_number, raw in rows:
        try:
            person = _person_by_lrn(db, raw["lrn"], row_number)
            school_year_name = str(raw["school_year"] or "").strip()
            school_year = db.scalar(select(SchoolYear).where(SchoolYear.name == school_year_name, SchoolYear.deleted_at.is_(None)))
            if not school_year:
                raise ValueError(f"row {row_number}: school year {school_year_name or '(blank)'} is not configured")
            grade_name = normalize_grade_name(str(raw["grade"] or ""))
            grade = db.scalar(select(GradeLevel).where(GradeLevel.name == grade_name, GradeLevel.active.is_(True), GradeLevel.deleted_at.is_(None)))
            if not grade or grade_name not in {"7", "8", "9", "10"}:
                raise ValueError(f"row {row_number}: Grade 7–10 is required")
            section_name = str(raw["section"] or "").strip()
            section = db.scalar(select(SchoolSection).where(
                SchoolSection.grade_level_id == grade.id,
                SchoolSection.active.is_(True), SchoolSection.deleted_at.is_(None),
            ).order_by(SchoolSection.id))
            if not section:
                raise ValueError(f"row {row_number}: no active section exists for Grade {grade_name}")
            if section.name.casefold() != section_name.casefold():
                section = db.scalar(select(SchoolSection).where(
                    SchoolSection.grade_level_id == grade.id,
                    SchoolSection.active.is_(True), SchoolSection.deleted_at.is_(None),
                    SchoolSection.name == section_name,
                ))
            if not section or section.name.casefold() != section_name.casefold():
                raise ValueError(f"row {row_number}: section {section_name or '(blank)'} is not active in Grade {grade_name}")
            starts_on = _value_date(raw["starts_on"], "enrollment start", row_number)
            ends_on = _value_date(raw.get("ends_on"), "enrollment end", row_number)
            if not starts_on:
                raise ValueError(f"row {row_number}: enrollment start date is required")
            if ends_on and ends_on < starts_on:
                raise ValueError(f"row {row_number}: enrollment end date is before the start date")
            if starts_on > school_year.ends_on or (ends_on and ends_on < school_year.starts_on):
                raise ValueError(f"row {row_number}: enrollment interval does not overlap school year {school_year_name}")
            key = (person.id, school_year.id, starts_on)
            if key in seen:
                raise ValueError(f"row {row_number}: duplicate learner/school-year/start date")
            seen.add(key)
            previous = db.scalars(select(LearnerEnrollment).where(
                LearnerEnrollment.person_id == person.id,
                LearnerEnrollment.school_year_id == school_year.id,
            )).all()
            for item in previous:
                old_end = item.ends_on or school_year.ends_on
                new_end = ends_on or school_year.ends_on
                if item.starts_on != starts_on and starts_on <= old_end and item.starts_on <= new_end:
                    raise ValueError(f"row {row_number}: enrollment dates overlap another Grade {item.grade_level.name} section placement")
            existing = db.scalar(select(LearnerEnrollment).where(
                LearnerEnrollment.person_id == person.id,
                LearnerEnrollment.school_year_id == school_year.id,
                LearnerEnrollment.starts_on == starts_on,
            ))
            entries.append({"row": row_number, "person": person, "year": school_year, "grade": grade,
                            "section": section, "starts_on": starts_on, "ends_on": ends_on, "existing": existing})
        except (ValueError, TypeError) as error:
            errors.append({"row": row_number, "errors": [str(error)]})
    if dry_run:
        return {"dry_run": True, "valid_rows": len(entries), "invalid_rows": len(errors), "errors": errors[:100]}
    if errors:
        raise HTTPException(status_code=422, detail={"message": "Correct invalid enrollment rows before import; no rows were changed", "errors": errors[:100]})
    inserted = updated = 0
    digest = hashlib.sha256(data).hexdigest()
    for entry in entries:
        record = entry["existing"]
        if record:
            updated += 1
        else:
            record = LearnerEnrollment(person_id=entry["person"].id, school_year_id=entry["year"].id,
                starts_on=entry["starts_on"], imported_by=actor.id, imported_by_name=actor.full_name,
                source_reference=reference, source_sha256=digest)
            db.add(record)
            inserted += 1
        record.grade_level_id = entry["grade"].id
        record.section_id = entry["section"].id
        record.ends_on = entry["ends_on"]
        record.source_reference = reference
        record.source_sha256 = digest
        record.imported_by = actor.id
        record.imported_by_name = actor.full_name
    db.add(RosterImportAudit(id=str(uuid.uuid4()), filename=filename[:255], file_sha256=digest,
        approved_reference=reference, inserted_count=inserted, updated_count=updated, skipped_count=0,
        import_mode="annex_g_enrollment", actor_user_id=actor.id, actor_name=actor.full_name))
    db.commit()
    return {"dry_run": False, "inserted": inserted, "updated": updated, "skipped": 0, "invalid_rows": 0}


def backfill_annex_g_dobs(db: Session, data: bytes, filename: str, approved_reference: str,
                          actor: User, dry_run: bool = True) -> dict:
    reference = _approval(approved_reference)
    rows = _records(data, {"lrn": {"lrn", "learner reference number"},
                           "dob": {"dob", "date of birth", "dateofbirth"}})
    errors: list[dict] = []
    entries = []
    seen: set[int] = set()
    skipped = 0
    for row_number, raw in rows:
        try:
            person = _person_by_lrn(db, raw["lrn"], row_number)
            dob = _value_date(raw["dob"], "date of birth", row_number)
            if not dob:
                raise ValueError(f"row {row_number}: date of birth is required")
            if dob > date.today():
                raise ValueError(f"row {row_number}: date of birth cannot be in the future")
            if person.id in seen:
                raise ValueError(f"row {row_number}: duplicate LRN in workbook")
            seen.add(person.id)
            if person.date_of_birth:
                skipped += 1
                continue
            entries.append({"row": row_number, "person": person, "dob": dob})
        except (ValueError, TypeError) as error:
            errors.append({"row": row_number, "errors": [str(error)]})
    if dry_run:
        return {"dry_run": True, "valid_rows": len(entries), "invalid_rows": len(errors), "would_fill": len(entries),
                "would_skip_existing": skipped, "errors": errors[:100]}
    if errors:
        raise HTTPException(status_code=422, detail={"message": "Correct invalid DOB rows before import; no dates were changed", "errors": errors[:100]})
    for entry in entries:
        entry["person"].date_of_birth = entry["dob"]
    digest = hashlib.sha256(data).hexdigest()
    db.add(RosterImportAudit(id=str(uuid.uuid4()), filename=filename[:255], file_sha256=digest,
        approved_reference=reference, inserted_count=0, updated_count=len(entries), skipped_count=skipped,
        import_mode="annex_g_dob", actor_user_id=actor.id, actor_name=actor.full_name))
    db.commit()
    return {"dry_run": False, "inserted": 0, "updated": len(entries), "skipped": skipped, "invalid_rows": 0}


def first_friday_of_june(school_year: SchoolYear) -> date:
    first = date(school_year.starts_on.year, 6, 1)
    return first + timedelta(days=(4 - first.weekday()) % 7)


def age_on(birth_date: date, reference_date: date) -> int:
    return reference_date.year - birth_date.year - ((reference_date.month, reference_date.day) < (birth_date.month, birth_date.day))


def _descriptor(grade: int | None) -> str:
    if grade is None:
        return ""
    return "Advancing" if grade >= 90 else "Benchmarking" if grade >= 80 else "Connecting" if grade >= 75 else "Developing" if grade >= 65 else "Emerging"


def _term_average(values: list[int | None]) -> int | None:
    if not values or any(value is None for value in values):
        return None
    return round_value(sum(values) / len(values), 0, "half_up")


def _attendance_by_learner(db: Session, enrollments: list[LearnerEnrollment], term_end: date) -> dict[int, list[dict]]:
    today = local_now().date()
    through = min(term_end, today)
    if not enrollments:
        return {}
    first = min(item.school_year.starts_on for item in enrollments)
    proxies = {}
    totals: dict[int, dict[str, dict[str, int | str]]] = {item.id: {} for item in enrollments}
    for enrollment in enrollments:
        person = enrollment.person
        proxy = SimpleNamespace(id=person.id, external_id=person.external_id, lrn=person.lrn,
            full_name=person.full_name, sex=person.sex, role=person.role, assignment=person.assignment,
            enrollment_note=person.enrollment_note, transfer_school=person.transfer_school,
            enrollment_start_date=enrollment.starts_on, enrollment_end_date=enrollment.ends_on,
            enrollment_status="Regular", active=True, grade=enrollment.grade_level.name,
            section=enrollment.section.name)
        proxies[enrollment.id] = proxy
    day = first
    while day <= through:
        if is_instructional_day(db, day):
            eligible_ids = [item.id for item in enrollments
                            if item.starts_on <= day and (not item.ends_on or item.ends_on >= day)]
            eligible_people = [proxies[item_id] for item_id in eligible_ids]
            for item_id in eligible_ids:
                key = day.strftime("%Y-%m")
                totals[item_id].setdefault(key, {"month": key, "days_of_class": 0, "days_present": 0,
                                                  "days_absent": 0, "unresolved_days": 0})["days_of_class"] += 1
            for row in rows_for_people(db, eligible_people, day):
                enrollment_id = next((item_id for item_id in eligible_ids
                                      if proxies[item_id].id == row["person_id"]), None)
                if enrollment_id is None:
                    continue
                key = day.strftime("%Y-%m")
                if row["status"] in {"Present", "Late", "Time Out"}:
                    totals[enrollment_id][key]["days_present"] += 1
                elif row["status"] in {"Absent", "Excused"}:
                    totals[enrollment_id][key]["days_absent"] += 1
                elif row["status"] == "No scan":
                    totals[enrollment_id][key]["unresolved_days"] += 1
        day = day.fromordinal(day.toordinal() + 1)
    result = {}
    for enrollment_id, months in totals.items():
        result[enrollment_id] = [months[key] for key in sorted(months)]
    return result


def build_annex_g_section(db: Session, school_year_id: int, grade: str, section: str,
                          term_order: int) -> dict:
    if term_order not in {1, 2, 3}:
        raise HTTPException(status_code=422, detail="Annex G supports Terms 1, 2, and 3")
    year = db.scalar(select(SchoolYear).where(SchoolYear.id == school_year_id, SchoolYear.deleted_at.is_(None)))
    if not year:
        raise HTTPException(status_code=404, detail="School year was not found")
    grade_name = normalize_grade_name(grade)
    if grade_name not in {"7", "8", "9", "10"}:
        raise HTTPException(status_code=422, detail="Annex G is available for Grades 7–10 only")
    grade_record = db.scalar(select(GradeLevel).where(
        GradeLevel.name == grade_name, GradeLevel.deleted_at.is_(None)))
    section_record = db.scalar(select(SchoolSection).where(
        SchoolSection.grade_level_id == grade_record.id if grade_record else False,
        SchoolSection.name == section, SchoolSection.deleted_at.is_(None))) if grade_record else None
    if not section_record:
        raise HTTPException(status_code=404, detail="Grade and section were not found")
    terms = {item.term_order: item for item in db.scalars(select(GradingPeriod).where(
        GradingPeriod.school_year_id == year.id, GradingPeriod.deleted_at.is_(None),
        GradingPeriod.term_order <= term_order)).all()}
    if term_order not in terms:
        raise HTTPException(status_code=409, detail=f"Term {term_order} is not configured for {year.name}")
    report_end = terms[term_order].ends_on
    reference_date = first_friday_of_june(year)
    enrollments = db.scalars(select(LearnerEnrollment).where(
        LearnerEnrollment.school_year_id == year.id,
        LearnerEnrollment.grade_level_id == grade_record.id,
        LearnerEnrollment.section_id == section_record.id,
        LearnerEnrollment.starts_on <= report_end,
        (LearnerEnrollment.ends_on.is_(None) | (LearnerEnrollment.ends_on >= year.starts_on)),
    ).order_by(LearnerEnrollment.person_id, LearnerEnrollment.starts_on)).all()
    enrollments = [item for item in enrollments if item.person.role == "Student"]
    person_ids = list(dict.fromkeys(item.person_id for item in enrollments))
    assignments = db.scalars(select(SubjectTeachingAssignment).join(
        Subject, Subject.id == SubjectTeachingAssignment.subject_id,
    ).where(
        SubjectTeachingAssignment.school_year_id == year.id,
        SubjectTeachingAssignment.section_id == section_record.id,
        SubjectTeachingAssignment.active.is_(True), Subject.active.is_(True), Subject.deleted_at.is_(None),
    ).order_by(Subject.name)).all()
    assignments = [item for item in assignments if subject_matches_grade(item.subject, grade_name)]
    if not assignments:
        assignment_issue = "No active learning-area assignments were found for this section and school year."
    else:
        assignment_issue = None

    books = db.scalars(select(Gradebook).where(
        Gradebook.school_year_id == year.id,
        Gradebook.section_id == section_record.id,
        Gradebook.subject_id.in_([item.subject_id for item in assignments] or [-1]),
        Gradebook.term_order <= term_order,
        Gradebook.is_legacy_readonly.is_(False),
    )).all()
    book_map = {(book.subject_id, book.term_order, book.learning_area or ""): book for book in books}
    computed: dict[int, dict[int, dict[int, int | None]]] = {}
    for book in books:
        full = get_gradebook_full(db, book.id, include_mapeh_summary=False, roster_person_ids=person_ids)
        computed[book.id] = {row["person_id"]: row for row in full["students"]} if full else {}

    attendance = _attendance_by_learner(db, enrollments, report_end)
    comments = db.scalars(select(AnnexGTermComment).join(
        LearnerEnrollment, LearnerEnrollment.id == AnnexGTermComment.enrollment_id,
    ).where(AnnexGTermComment.enrollment_id.in_([item.id for item in enrollments] or [-1]),
            AnnexGTermComment.term_order <= term_order)).all()
    comments_by_enrollment: dict[int, dict[str, str]] = {}
    for comment in comments:
        comments_by_enrollment.setdefault(comment.enrollment_id, {})[str(comment.term_order)] = comment.comment

    learning_area_specs = []
    for assignment in assignments:
        subject = assignment.subject
        is_mapeh = subject.code.casefold() == "mapeh" or "mapeh" in subject.name.casefold()
        learning_area_specs.append({"subject": subject, "mapeh": is_mapeh})

    learner_rows = []
    for enrollment in enrollments:
        person = enrollment.person
        blockers = []
        age = age_on(person.date_of_birth, reference_date) if person.date_of_birth and person.date_of_birth <= reference_date else None
        if person.date_of_birth is None:
            blockers.append("Verified date of birth is missing; an administrator must enter it in People & Records from an official learner record.")
        elif person.date_of_birth > reference_date:
            blockers.append("Date of birth is after the school-year age reference date; verify the source record.")
        if assignment_issue:
            blockers.append(assignment_issue)
        rows = []
        final_grades = []
        for spec in learning_area_specs:
            subject = spec["subject"]
            if spec["mapeh"]:
                area_rows = []
                for area_key, area_name in (("music_arts", "Music and Arts"), ("pe_health", "Physical Education and Health")):
                    area_terms: dict[str, int | None] = {}
                    for number in range(1, term_order + 1):
                        book = book_map.get((subject.id, number, area_key))
                        grade_value = None
                        if not book:
                            blockers.append(f"{subject.name} · {area_name} · Term {number}: no gradebook exists.")
                        elif book.status not in {"Finalized", "Locked"}:
                            blockers.append(f"{subject.name} · {area_name} · Term {number}: gradebook is {book.status}.")
                        else:
                            student = computed.get(book.id, {}).get(person.id)
                            if not student:
                                blockers.append(f"{subject.name} · {area_name} · Term {number}: learner is missing from the gradebook roster.")
                            elif student["reported_grade"] is None:
                                blockers.append(f"{subject.name} · {area_name} · Term {number}: score calculation is incomplete.")
                            else:
                                grade_value = student["reported_grade"]
                        area_terms[str(number)] = grade_value
                    area_final = _term_average([area_terms.get(str(number)) for number in (1, 2, 3)]) if term_order == 3 else None
                    area_rows.append({"name": area_name, "term_grades": area_terms, "final_grade": area_final,
                                      "descriptor": _descriptor(area_final)})
                terms_combined = {}
                for number in range(1, term_order + 1):
                    values = [area_rows[index]["term_grades"].get(str(number)) for index in range(2)]
                    terms_combined[str(number)] = _term_average(values)
                mapeh_final = _term_average([terms_combined.get(str(number)) for number in (1, 2, 3)]) if term_order == 3 else None
                rows.append({"name": "MAPEH", "term_grades": terms_combined, "final_grade": mapeh_final,
                             "descriptor": _descriptor(mapeh_final),
                             "remarks": "Passed" if mapeh_final is not None and mapeh_final >= 75 else "Failed" if mapeh_final is not None else "",
                             "subareas": area_rows})
                final_grades.append(mapeh_final)
                continue
            terms_grades: dict[str, int | None] = {}
            for number in range(1, term_order + 1):
                book = book_map.get((subject.id, number, ""))
                grade_value = None
                if not book:
                    blockers.append(f"{subject.name} · Term {number}: no gradebook exists.")
                elif book.status not in {"Finalized", "Locked"}:
                    blockers.append(f"{subject.name} · Term {number}: gradebook is {book.status}.")
                else:
                    student = computed.get(book.id, {}).get(person.id)
                    if not student:
                        blockers.append(f"{subject.name} · Term {number}: learner is missing from the gradebook roster.")
                    elif student["reported_grade"] is None:
                        blockers.append(f"{subject.name} · Term {number}: score calculation is incomplete.")
                    else:
                        grade_value = student["reported_grade"]
                terms_grades[str(number)] = grade_value
            final_grade = _term_average([terms_grades.get(str(number)) for number in (1, 2, 3)]) if term_order == 3 else None
            rows.append({"name": subject.name, "term_grades": terms_grades, "final_grade": final_grade,
                         "descriptor": _descriptor(final_grade),
                         "remarks": "Passed" if final_grade is not None and final_grade >= 75 else "Failed" if final_grade is not None else ""})
            final_grades.append(final_grade)
        general_average = _term_average(final_grades) if term_order == 3 else None
        if term_order == 3 and final_grades and general_average is None:
            blockers.append("General Average is unavailable until every assigned learning area has all three finalized term grades.")
        monthly = attendance.get(enrollment.id, [])
        unresolved = sum(item["unresolved_days"] for item in monthly)
        if unresolved:
            blockers.append(f"Attendance has {unresolved} unresolved instructional day(s); reconcile attendance through the selected term.")
        learner_rows.append({
            "enrollment_id": enrollment.id, "person_id": person.id, "full_name": person.full_name,
            "lrn": person.lrn, "sex": person.sex, "grade": grade_name, "section": section_record.name,
            "age": age, "learning_areas": rows, "general_average": general_average,
            "attendance": monthly, "comments": comments_by_enrollment.get(enrollment.id, {}),
            "blockers": list(dict.fromkeys(blockers)),
        })
    profile = get_json(db, "annex_g.school", {})
    school_head = get_json(db, "school.heads", {}).get(str(year.id), {}).get("full_name", "")
    return {
        "school_year_id": year.id, "school_year": year.name, "grade": grade_name,
        "section": section_record.name, "term_order": term_order, "term_name": terms[term_order].name,
        "age_reference_date": reference_date.isoformat(),
        "school": {"name": profile.get("name", ""), "region": profile.get("region", ""),
                   "division": profile.get("division", ""), "school_id": profile.get("school_id", ""),
                   "school_head": school_head},
        "learners": learner_rows,
        "blockers": ["No verified learner enrollment placements exist for this section and school year."] if not enrollments else [],
    }


def render_annex_g_html(report: dict, enrollment_id: int) -> str:
    learner = next((row for row in report["learners"] if row["enrollment_id"] == enrollment_id), None)
    if not learner:
        raise HTTPException(status_code=404, detail="Learner is not part of this report")
    school = report["school"]
    table_rows = []
    for area in learner["learning_areas"]:
        table_rows.append((area["name"], area["term_grades"], area["final_grade"],
                           area.get("descriptor", ""), area["remarks"], False))
        for subarea in area.get("subareas", []):
            table_rows.append((subarea["name"], subarea["term_grades"], subarea["final_grade"],
                               subarea.get("descriptor", ""), "", True))
    if not table_rows:
        table_rows.append(("", {}, None, "", "", False))
    grade_rows = []
    for index, (name, grades, final_grade, descriptor, remarks, subarea) in enumerate(table_rows):
        identity = ""
        if index == 0:
            span = len(table_rows)
            identity = (f"<td rowspan='{span}'>{html.escape(learner['full_name'])}</td>"
                f"<td rowspan='{span}'>{html.escape(learner['lrn'] or '')}</td>"
                f"<td rowspan='{span}'>{learner['age'] if learner['age'] is not None else ''}</td>"
                f"<td rowspan='{span}'>{html.escape(learner['sex'])}</td>")
        area_cell = f"<td>{html.escape(name)}</td>" if subarea else f"<th>{html.escape(name)}</th>"
        term_cells = "".join(f"<td>{html.escape(str(grades.get(str(number)))) if grades.get(str(number)) is not None and number <= report['term_order'] else ''}</td>"
                              for number in (1, 2, 3))
        grade_rows.append(f"<tr class='subarea'>{identity}{area_cell}{term_cells}"
            f"<td>{final_grade if final_grade is not None else ''}</td><td>{html.escape(descriptor)}</td>"
            f"<td>{html.escape(remarks)}</td></tr>")
    attendance_rows = "".join(f"<tr><td>{html.escape(item['month'])}</td><td>{item['days_of_class']}</td>"
        f"<td>{item['days_present']}</td><td>{item['days_absent']}</td></tr>" for item in learner["attendance"])
    comments = "".join(f"<p><strong>Term {number} teacher comment:</strong> {html.escape(learner['comments'].get(str(number), ''))}</p>"
        f"<p class='signature'>Teacher signature: ____________________　Parent/Guardian signature: ____________________</p>"
        for number in range(1, report["term_order"] + 1))
    blockers = "" if not learner["blockers"] else "<aside><strong>Needs attention — not ready for release</strong><ul>" + "".join(
        f"<li>{html.escape(item)}</li>" for item in learner["blockers"]) + "</ul></aside>"
    school_name = html.escape(school.get("name") or "")
    return f"""<!doctype html><html><head><meta charset='utf-8'><title>Annex G — {html.escape(learner['full_name'])}</title>
<style>@page{{size:A4 landscape;margin:12mm}}body{{font:12px Arial,sans-serif;color:#17253b}}h1,h2{{text-align:center;margin:4px}}.school{{text-align:center;margin-bottom:14px}}table{{width:100%;border-collapse:collapse;margin:8px 0 14px}}th,td{{border:1px solid #53647c;padding:5px;text-align:center}}th:first-child,td:first-child{{text-align:left}}thead th{{background:#e8eef6}}.subarea{{background:#f5f7fa}}aside{{border:1px solid #bd5a2a;background:#fff4ec;padding:8px;margin:10px 0}}.signature{{border-bottom:1px solid #9aa6b5;padding:8px 0}}.columns{{display:grid;grid-template-columns:1fr 1fr;gap:18px}}.descriptor{{font-size:11px}}</style></head>
<body><div class='school'><h1>PERFORMANCE REPORT TEMPLATE — ANNEX G</h1><h2>{school_name}</h2><p>Region: {html.escape(school.get('region') or '')}　Division: {html.escape(school.get('division') or '')}　School ID: {html.escape(school.get('school_id') or '')}</p>
<p>School Head: {html.escape(school.get('school_head') or '')}</p><p>School Year {html.escape(report['school_year'])} · Grade {html.escape(learner['grade'])} · Section {html.escape(learner['section'])} · Through {html.escape(report['term_name'])}</p></div>
{blockers}<table><thead><tr><th>Learner</th><th>LRN</th><th>Age (as of {html.escape(report['age_reference_date'])})</th><th>Sex</th><th>Learning Area</th><th>Term 1</th><th>Term 2</th><th>Term 3</th><th>Final Grade</th><th>Descriptor</th><th>Remarks</th></tr></thead>
<tbody>{''.join(grade_rows)}</tbody></table><p class='descriptor'><strong>Descriptor scale:</strong> Advancing 90–100 · Benchmarking 80–89 · Connecting 75–79 · Developing 65–74 · Emerging 0–64</p>
<p><strong>General Average:</strong> {learner['general_average'] if learner['general_average'] is not None else ''}</p>
<div class='columns'><section><h2>Monthly Attendance</h2><table><thead><tr><th>Month</th><th>Days of Class</th><th>Days Present</th><th>Days Absent</th></tr></thead><tbody>{attendance_rows}</tbody></table></section><section><h2>Teacher Comments and Acknowledgment</h2>{comments}</section></div>
<h2>Certificate of Transfer</h2><p>Admitted to: ______________________________　Date: ______________　School: ______________________________</p>
<h2>Cancellation of Transfer</h2><p>Cancel transfer to: ________________________　Date: ______________　Parent/Guardian signature: __________________</p></body></html>"""
