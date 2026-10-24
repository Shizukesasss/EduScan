from __future__ import annotations

import hashlib
import html
import json
import math
import posixpath
import re
import shutil
import uuid
from collections import Counter, defaultdict
from datetime import date, timedelta
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import EXPORT_DIR
from ..models import GeneratedReport, Person, User
from .attendance import active_people, rows_for_people
from .grading import gradebook_summary
from .academic_context import is_senior_high
from .grading_engine import round_value
from .export_files import new_report_path
from .settings_store import get_json


REPORT_DIR = EXPORT_DIR / "registered"
REPORT_DIR.mkdir(parents=True, exist_ok=True)
TEMPLATES_DIR = Path(__file__).resolve().parents[2] / "templates"


def _safe_name(value: object) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "-", str(value)).strip("-._") or "report"


def register_report(db: Session, source: Path, report_type: str, parameters: dict, actor: User) -> GeneratedReport:
    if not source.exists():
        raise HTTPException(status_code=500, detail="Generated report file was not found")
    report_id = str(uuid.uuid4())
    target = REPORT_DIR / f"{report_id}-{source.name}"
    shutil.copy2(source, target)
    digest = hashlib.sha256(target.read_bytes()).hexdigest()
    record = GeneratedReport(
        id=report_id, report_type=report_type, filename=source.name, file_path=str(target), file_sha256=digest,
        parameters_json=json.dumps(parameters, default=str, sort_keys=True), status="Generated",
        generated_by=actor.id, generated_by_name=actor.full_name, generated_by_role=actor.role,
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return record


def official_sf2_preview(path: Path) -> dict:
    from .sf2 import SHEET_NAME, MALE_ROWS, FEMALE_ROWS, DAY_COLUMNS, NS_XDR
    tardy = set()
    with ZipFile(path) as archive:
        if "xl/drawings/drawing1.xml" in archive.namelist():
            drawing = ET.fromstring(archive.read("xl/drawings/drawing1.xml"))
            for node in drawing.iter(f"{{{NS_XDR}}}cNvPr"):
                match = re.fullmatch(r"EduScan Tardy (\d+)-(\d+)", node.get("name", ""))
                if match:
                    tardy.add(tuple(map(int, match.groups())))
    with path.open('rb') as source:
        book = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = book[SHEET_NAME]
            cells = list(sheet.iter_rows(min_row=1, max_row=62, max_col=31, values_only=True))
            days = [(index, cells[10][column - 1]) for index, column in enumerate(DAY_COLUMNS) if cells[10][column - 1] is not None]
            learners = []
            for sex, rows in (("Male", MALE_ROWS), ("Female", FEMALE_ROWS)):
                for row in rows:
                    values = cells[row - 1]
                    if not values[1]:
                        continue
                    learners.append({"name": str(values[1]), "sex": sex,
                        "marks": ["L" if (row, DAY_COLUMNS[index]) in tardy else str(values[DAY_COLUMNS[index] - 1] or "") for index, _ in days],
                        "absences": values[28] or 0, "late": values[29] or 0, "remarks": str(values[30] or "")})
            return {"days": [day for _, day in days], "learners": learners,
                    "school_name": str(cells[7][2] or ""), "school_id": str(cells[5][2] or ""),
                    "school_year": str(cells[5][10] or "")}
        finally:
            book.close()


def _title(sheet, title: str, subtitle: str, last_column: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet.cell(1, 1, title).font = Font(size=16, bold=True)
    sheet.cell(1, 1).alignment = Alignment(horizontal="center")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet.cell(2, 1, subtitle).alignment = Alignment(horizontal="center")


def _headers(sheet, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, 1):
        cell = sheet.cell(row, column, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def generate_grade_report_xlsx(db: Session, class_key: str) -> Path:
    summary = gradebook_summary(db, class_key)
    metadata, rows, statistics = summary["metadata"], summary["rows"], summary["statistics"]
    if not metadata["grade"] or not metadata["section"]:
        raise HTTPException(status_code=409, detail="Save the gradebook before generating its report")
    book = Workbook()
    sheet = book.active
    sheet.title = "Grade Summary"
    labels = ["LRN / ID", "Learner", "Initial Grade", "Transmuted Grade", "Result", "Completion"]
    _title(sheet, "EDUSCAN GRADING SUMMARY", f"SY {metadata['school_year']} · Term {metadata['quarter']} · {metadata['subject']} · Grade {metadata['grade']} {metadata['section']}", len(labels))
    _headers(sheet, 4, labels)
    for row_number, item in enumerate(rows, 5):
        values = [item["lrn"] or item["external_id"], item["full_name"], item["initial_grade"], item["transmuted_grade"], item["status"], "Complete" if item["complete"] else "Incomplete"]
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)
    stats_row = len(rows) + 7
    sheet.cell(stats_row, 1, "CLASS STATISTICS").font = Font(bold=True)
    stats = [("Learners", statistics["learners"]), ("Complete", statistics["complete"]), ("Incomplete", statistics["incomplete"]),
             ("Passed", statistics["passed"]), ("Below rule", statistics["below_rule"]), ("Average", statistics["average"]),
             ("Highest", statistics["highest"]), ("Lowest", statistics["lowest"])]
    for index, (label, value) in enumerate(stats, stats_row + 1):
        sheet.cell(index, 1, label).font = Font(bold=True)
        sheet.cell(index, 2, value)
    sheet.freeze_panes = "A5"
    for column, width in {"A": 20, "B": 34, "C": 16, "D": 18, "E": 16, "F": 14}.items():
        sheet.column_dimensions[column].width = width
    target = new_report_path(EXPORT_DIR, f"Grade-Summary-{_safe_name(metadata['school_year'])}-Term{metadata['quarter']}-{_safe_name(metadata['grade'])}-{_safe_name(metadata['section'])}-{_safe_name(metadata['subject'])}.xlsx")
    book.save(target)
    return target


def grade_report_html(db: Session, class_key: str) -> str:
    summary = gradebook_summary(db, class_key)
    metadata, rows, statistics = summary["metadata"], summary["rows"], summary["statistics"]
    body = "".join(
        f"<tr><td>{html.escape(item['lrn'] or item['external_id'])}</td><td>{html.escape(item['full_name'])}</td>"
        f"<td>{item['initial_grade'] if item['initial_grade'] is not None else '—'}</td><td>{item['transmuted_grade'] if item['transmuted_grade'] is not None else '—'}</td>"
        f"<td>{html.escape(item['status'])}</td></tr>" for item in rows
    )
    return _printable_html(
        "EduScan Grading Summary",
        f"SY {metadata['school_year']} · Term {metadata['quarter']} · {metadata['subject']} · Grade {metadata['grade']} {metadata['section']}",
        "<table><thead><tr><th>LRN / ID</th><th>Learner</th><th>Initial</th><th>Transmuted</th><th>Result</th></tr></thead><tbody>" + body + "</tbody></table>"
        f"<h2>Class statistics</h2><p>Learners: {statistics['learners']} · Complete: {statistics['complete']} · Incomplete: {statistics['incomplete']} · "
        f"Passed: {statistics['passed']} · Below rule: {statistics['below_rule']} · Average: {statistics['average'] if statistics['average'] is not None else '—'} · "
        f"Highest: {statistics['highest'] if statistics['highest'] is not None else '—'} · Lowest: {statistics['lowest'] if statistics['lowest'] is not None else '—'}</p>",
    )


_SHEET_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_SHEET_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _xlsx_sheet_paths(archive: ZipFile) -> dict[str, str]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {item.get("Id"): item.get("Target") for item in relationships.findall(f"{{{_PACKAGE_REL_NS}}}Relationship")}
    sheets = workbook.find(f"{{{_SHEET_MAIN_NS}}}sheets")
    result = {}
    if sheets is None:
        return result
    for sheet in sheets.findall(f"{{{_SHEET_MAIN_NS}}}sheet"):
        target = targets.get(sheet.get(f"{{{_SHEET_REL_NS}}}id"))
        if target:
            result[sheet.get("name", "")] = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
    return result


def _xlsx_cell_xml(attributes: bytes, value: object) -> bytes:
    attributes = re.sub(rb'\s+t="[^"]*"', b"", attributes)
    if value is None:
        content = b"<is><t></t></is>"
        attributes += b' t="inlineStr">'
    elif isinstance(value, str):
        text = html.escape(value, quote=False).encode("utf-8")
        preserve = b' xml:space="preserve"' if value[:1].isspace() or value[-1:].isspace() else b""
        content = b"<is><t" + preserve + b">" + text + b"</t></is>"
        attributes += b' t="inlineStr">'
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        if not math.isfinite(number):
            raise ValueError("Class-record template value is not finite")
        content = b"<v>" + format(number, ".15g").encode("ascii") + b"</v>"
        attributes += b">"
    else:
        raise ValueError("Unsupported class-record template value")
    return attributes + content + b"</c>"


def _set_xlsx_cell(sheet_xml: bytes, coordinate: str, value: object) -> bytes:
    pattern = re.compile(rb'<c\b(?=[^>]*\br="' + re.escape(coordinate.encode("ascii")) + rb'")[^>]*?(?:/>|>.*?</c>)', re.DOTALL)
    match = pattern.search(sheet_xml)
    if match:
        original = match.group(0)
        opening = original[:original.find(b">") + 1]
        attributes = opening[:-1].rstrip()
        if attributes.endswith(b"/"):
            attributes = attributes[:-1].rstrip()
        replacement = _xlsx_cell_xml(attributes, value)
        return sheet_xml[:match.start()] + replacement + sheet_xml[match.end():]

    row_number = re.search(r"\d+$", coordinate).group()
    column_letters = re.match(r"[A-Z]+", coordinate).group()
    column_number = 0
    for letter in column_letters:
        column_number = column_number * 26 + ord(letter) - ord("A") + 1
    row_pattern = re.compile(rb'<row\b(?=[^>]*\br="' + row_number.encode("ascii") + rb'")[^>]*(?:/>|>.*?</row>)', re.DOTALL)
    row_match = row_pattern.search(sheet_xml)
    if not row_match:
        raise ValueError(f"Template row {row_number} for cell {coordinate} is missing")
    original_row = row_match.group(0)
    opening_end = original_row.find(b">") + 1
    opening = original_row[:opening_end]
    body = b"" if opening.endswith(b"/>") else original_row[opening_end:-6]
    cells = []
    cell_pattern = re.compile(rb'<c\b[^>]*(?:/>|>.*?</c>)', re.DOTALL)
    for cell in cell_pattern.finditer(body):
        cell_opening = cell.group(0)[:cell.group(0).find(b">") + 1]
        reference = re.search(rb'\br="([A-Z]+)(\d+)"', cell_opening)
        if not reference:
            continue
        number = 0
        for letter in reference.group(1).decode("ascii"):
            number = number * 26 + ord(letter) - ord("A") + 1
        cells.append((cell.start(), cell.end(), number, cell_opening))
    nearest_style = None
    insertion = len(body)
    for start, _, number, cell_opening in cells:
        if number > column_number:
            insertion = start
            break
    if cells:
        nearest = min(cells, key=lambda cell: abs(cell[2] - column_number))
        style = re.search(rb'\bs="([^"]+)"', nearest[3])
        if style:
            nearest_style = b' s="' + style.group(1) + b'"'
    attributes = b'<c r="' + coordinate.encode("ascii") + b'"' + (nearest_style or b"")
    new_cell = _xlsx_cell_xml(attributes, value)
    if opening.endswith(b"/>"):
        opening = opening[:-2] + b">"
        updated_row = opening + new_cell + b"</row>"
    else:
        updated_row = original_row[:opening_end] + body[:insertion] + new_cell + body[insertion:] + b"</row>"
    return sheet_xml[:row_match.start()] + updated_row + sheet_xml[row_match.end():]


def _set_calc_attribute(opening: bytes, name: bytes, value: bytes) -> bytes:
    pattern = re.compile(rb"(\s" + name + rb"=)[\"'][^\"']*[\"']")
    if pattern.search(opening):
        return pattern.sub(lambda match: match.group(1) + b'"' + value + b'"', opening)
    return opening[:-2] + b" " + name + b'="' + value + b'"/>'


def _write_class_record_workbook(template: Path, target: Path, changes: dict[str, dict[str, object]], formula_values: dict | None = None, formula_guards: dict | None = None, formula_overrides: dict | None = None) -> None:
    from tempfile import NamedTemporaryFile

    if not template.is_file():
        raise HTTPException(status_code=409, detail=f"The required updated class-record workbook '{template.name}' is missing.")
    temporary_path = None
    def widen_numeric_column(match):
        tag = match.group()
        width = re.search(rb'\bwidth="([^"]+)"', tag)
        if width and 5 <= float(width.group(1)) < 8:
            return tag[:width.start(1)] + b'8' + tag[width.end(1):]
        return tag
    try:
        with ZipFile(template, "r") as source:
            paths = _xlsx_sheet_paths(source)
            updated_sheets = {}
            for sheet_name, values in changes.items():
                if sheet_name not in paths:
                    raise ValueError(f"Template sheet '{sheet_name}' is missing")
                content = source.read(paths[sheet_name])
                for coordinate, value in values.items():
                    content = _set_xlsx_cell(content, coordinate, value)
                for coordinate, formula in (formula_overrides or {}).get(sheet_name, {}).items():
                    if len(formula.lstrip('=')) > 8192:
                        raise HTTPException(409, 'The saved policy exceeds the Excel formula length limit.')
                    content = _set_xlsx_cell(content, coordinate, None)
                    pattern = rb'<c\b(?=[^>]*\br="' + coordinate.encode('ascii') + rb'")[^>]*>.*?</c>'
                    match = re.search(pattern, content, re.DOTALL)
                    opening = re.sub(rb'\s+t="[^"]*"', b'', match.group()[:match.group().find(b'>') + 1])
                    cell = opening + b'<f>' + html.escape(formula.lstrip('='), quote=False).encode('utf-8') + b'</f></c>'
                    content = content[:match.start()] + cell + content[match.end():]
                for coordinate, condition in (formula_guards or {}).get(sheet_name, {}).items():
                    pattern = rb'<c\b(?=[^>]*\br="' + coordinate.encode('ascii') + rb'")[^>]*>.*?</c>'
                    match = re.search(pattern, content, re.DOTALL)
                    if not match or not re.search(rb'<f\b', match.group()):
                        raise ValueError(f'Expected combined grade formula in {sheet_name}!{coordinate}')
                    cell = re.sub(rb'(<f\b[^>]*>)(.*?)(</f>)', lambda value: value.group(1) + b'IF(' + condition.encode('ascii') + b',' + value.group(2) + b',"")' + value.group(3), match.group(), flags=re.DOTALL)
                    cell = re.sub(rb'<v\b[^>]*(?:/>|>.*?</v>)', b'', cell, flags=re.DOTALL)
                    content = content[:match.start()] + cell + content[match.end():]
                for coordinate, value in (formula_values or {}).get(sheet_name, {}).items():
                    pattern = rb'<c\b(?=[^>]*\br="' + coordinate.encode('ascii') + rb'")[^>]*>.*?</c>'
                    match = re.search(pattern, content, re.DOTALL)
                    if not match or not re.search(rb'<f\b', match.group()):
                        raise ValueError(f'Expected calculation formula in {sheet_name}!{coordinate}')
                    cell = re.sub(rb'<v\b[^>]*(?:/>|>.*?</v>)', b'', match.group(), flags=re.DOTALL)
                    opening_end = cell.find(b'>') + 1
                    cell = re.sub(rb'\s+t="[^"]*"', b'', cell[:opening_end]) + cell[opening_end:]
                    cell = cell.replace(b'</c>', b'<v>' + format(float(value), '.15g').encode('ascii') + b'</v></c>')
                    content = content[:match.start()] + cell + content[match.end():]
                updated_sheets[paths[sheet_name]] = content
            # Include unused term sheets in combined templates; 100.00 needs more space.
            for sheet_name, worksheet_path in paths.items():
                if sheet_name.upper().startswith('TERM '):
                    content = updated_sheets.get(worksheet_path, source.read(worksheet_path))
                    updated_sheets[worksheet_path] = re.sub(rb'<col\b[^>]*/>', widen_numeric_column, content)
            with NamedTemporaryFile(dir=target.parent, suffix=".xlsx", delete=False) as temporary:
                temporary_path = Path(temporary.name)
            with ZipFile(temporary_path, "w", ZIP_DEFLATED) as output:
                for item in source.infolist():
                    if item.filename == "xl/calcChain.xml":
                        continue
                    if item.filename in updated_sheets:
                        content = updated_sheets[item.filename]
                    elif item.filename == "xl/workbook.xml":
                        content = source.read(item.filename)
                        calc = re.search(rb"<calcPr\b[^>]*/>", content)
                        if calc:
                            opening = calc.group(0)
                            for key, value in ((b"calcMode", b"auto"), (b"fullCalcOnLoad", b"1"), (b"forceFullCalc", b"1")):
                                opening = _set_calc_attribute(opening, key, value)
                            content = content[:calc.start()] + opening + content[calc.end():]
                        else:
                            content = content.replace(b"</workbook>", b'<calcPr calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/></workbook>')
                    elif item.filename == "xl/_rels/workbook.xml.rels":
                        content = re.sub(rb"<Relationship\b(?=[^>]*calcChain)[^>]*/>", b"", source.read(item.filename))
                    elif item.filename == "[Content_Types].xml":
                        content = re.sub(rb"<Override\b(?=[^>]*PartName=\"/xl/calcChain\.xml\")[^>]*/>", b"", source.read(item.filename))
                    else:
                        content = source.read(item.filename)
                    output.writestr(item, content)
        temporary_path.replace(target)
    except (OSError, KeyError, ValueError, ET.ParseError) as exc:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise HTTPException(status_code=500, detail=f"Could not fill the updated class-record workbook: {exc}") from exc


def _student_groups(students: list[dict]) -> tuple[list[dict], list[dict]]:
    males, females, unknown = [], [], []
    for student in students:
        sex = (student.get("sex") or "").strip().upper()
        if sex in {"M", "MALE"}:
            males.append(student)
        elif sex in {"F", "FEMALE"}:
            females.append(student)
        else:
            unknown.append(student)
    if unknown:
        raise HTTPException(status_code=409, detail="Set each student's sex to Male or Female before exporting this class record.")
    if len(males) > 50 or len(females) > 50:
        raise HTTPException(status_code=409, detail="The updated DepEd workbook has 50 student rows per sex; split this section before exporting.")
    return males, females


def _assessment_category(component: dict) -> str | None:
    value = f"{component.get('name', '')} {component.get('component_type', '')}".upper()
    if "WRITTEN" in value or re.search(r'\bORAL\b', value):
        return "ww"
    if "PERFORMANCE" in value or "TASK" in value:
        return "pt"
    if "EXAM" in value or "ASSESSMENT" in value or "QUIZ" in value:
        return "exam"
    return None


def _column_name(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _complete_finalized_record(full: dict | None) -> bool:
    return bool(
        full
        and full["gradebook"]["status"] in {"Finalized", "Locked"}
        and full["students"]
        and full["statistics"]["incomplete"] == 0
    )


def generate_gradebook_report_xlsx(db: Session, gradebook_id: int, region: str = "", division: str = "", school_id: str = "", allowed_teacher_id: int | None = None, all_terms: bool = False, through_term: int | None = None) -> Path:
    from .grading import get_gradebook_full
    from ..models_grading import Gradebook
    from .adaptive_class_record import export_adaptive

    full = get_gradebook_full(db, gradebook_id)
    book = db.get(Gradebook, gradebook_id)
    if not full or not book:
        raise HTTPException(status_code=404, detail="Gradebook not found")
    # Fixed JHS templates require ST1/ST2/TE; older policies keep their own formulas.
    uses_term_exams = any(component.get('component_type') == 'Term Examinations' for component in full['components'])
    fixed_rounding = full['policy']['rounding_method'] == 'half_up' and full['policy']['rounding_decimal_places'] == 2
    if not is_senior_high(full['gradebook']['grade_name']) and ((uses_term_exams and fixed_rounding) or full['gradebook']['subject_name'].strip().upper() in {'ARALING PANLIPUNAN', 'MAPEH'}):
        from .class_record_export import export_jhs, subject_template
        try:
            kind, filename = subject_template(full['gradebook']['subject_name'] or '',
                                              book.subject.is_elective is True)
        except HTTPException as exc:
            if exc.status_code != 409:
                raise
        else:
            return export_jhs(db, book, full, kind, TEMPLATES_DIR / 'SUBJECT CLASS RECORDS' / filename,
                              region, division, school_id, allowed_teacher_id, all_terms=all_terms, through_term=through_term)
    return export_adaptive(db, book, full, region, division, school_id, allowed_teacher_id, all_terms=all_terms, through_term=through_term)


def gradebook_report_html(db: Session, gradebook_id: int) -> str:

    from .grading import get_gradebook_full
    full = get_gradebook_full(db, gradebook_id)
    if not full:
        raise HTTPException(status_code=404, detail="Gradebook not found")
        
    gb = full["gradebook"]
    rows = full["students"]
    statistics = full["statistics"]
    
    body = "".join(
        f"<tr><td>{html.escape(item['lrn'] or item['external_id'])}</td><td>{html.escape(item['full_name'])}</td>"
        f"<td>{item['initial_grade'] if item['initial_grade'] is not None else '—'}</td><td>{item['reported_grade'] if item['reported_grade'] is not None else '—'}</td>"
        f"<td>{html.escape(item['status'])}</td></tr>" for item in rows
    )
    return _printable_html(
        "EduScan Grading Summary",
        f"SY {gb['school_year_name']} · Term {gb['quarter']} · {gb['subject_name']} · Grade {gb['grade_name']} {gb['section_name']}",
        "<table><thead><tr><th>LRN / ID</th><th>Learner</th><th>Initial</th><th>Transmuted</th><th>Result</th></tr></thead><tbody>" + body + "</tbody></table>"
        f"<h2>Class statistics</h2><p>Learners: {statistics['learners']} · Complete: {statistics['complete']} · Incomplete: {statistics['incomplete']} · "
        f"Passed: {statistics['passed']} · Below rule: {statistics['below_passing']} · Average: {statistics['average'] if statistics['average'] is not None else '—'} · "
        f"Highest: {statistics['highest'] if statistics['highest'] is not None else '—'} · Lowest: {statistics['lowest'] if statistics['lowest'] is not None else '—'}</p>",
    )



def attendance_range_rows(db: Session, starts_on: date, ends_on: date, role: str | None = None,
                          grade: str | None = None, section: str | None = None,
                          person_id: int | None = None) -> list[dict]:
    if ends_on < starts_on or (ends_on - starts_on).days > 366:
        raise HTTPException(status_code=422, detail="Attendance report range must be between 1 and 367 days")
    people = active_people(db, grade, section, students_only=bool(grade or section), include_departed=True)
    if role:
        people = [item for item in people if item.role == role]
    if person_id:
        people = [item for item in people if item.id == person_id]
    result = []
    current = starts_on
    while current <= ends_on:
        for row in rows_for_people(db, people, current):
            row["attendance_date"] = current
            result.append(row)
        current += timedelta(days=1)
    return result


def generate_attendance_range_xlsx(db: Session, starts_on: date, ends_on: date, role: str | None = None,
                                   grade: str | None = None, section: str | None = None,
                                   person_id: int | None = None) -> Path:
    rows = attendance_range_rows(db, starts_on, ends_on, role, grade, section, person_id)
    book = Workbook()
    detail = book.active
    detail.title = "Attendance Detail"
    labels = ["Date", "ID / LRN", "Name", "Role", "Grade", "Section / assignment", "Time in", "Time out", "Attendance status", "Enrollment / employment", "Start date", "Last enrolled / duty date", "Previous / receiving school", "Movement note", "Source"]
    _title(detail, "EDUSCAN ATTENDANCE REPORT", f"{starts_on.isoformat()} through {ends_on.isoformat()}", len(labels))
    _headers(detail, 4, labels)
    for row_number, item in enumerate(rows, 5):
        values = [item["attendance_date"], item["lrn"] or item["external_id"], item["full_name"], item["role"], item["grade"] or "",
                  item["section"] or item["assignment"] or "", item["time_in"].strftime("%H:%M:%S") if item["time_in"] else "",
                  item["time_out"].strftime("%H:%M:%S") if item["time_out"] else "", item["status"], item["enrollment_status"], item["enrollment_start_date"], item["enrollment_end_date"], item["transfer_school"], item["enrollment_note"], item["source"]]
        for column, value in enumerate(values, 1):
            detail.cell(row_number, column, value)
    detail.freeze_panes = "A5"
    for column, width in {"A": 13, "B": 20, "C": 34, "D": 24, "E": 10, "F": 24, "G": 12, "H": 12, "I": 16, "J": 30}.items():
        detail.column_dimensions[column].width = width

    summary_sheet = book.create_sheet("Summary")
    _title(summary_sheet, "ATTENDANCE SUMMARY", f"{starts_on.isoformat()} through {ends_on.isoformat()}", 9)
    summary_labels = ["ID / LRN", "Name", "Role", "Present", "Late", "Absent", "Excused", "No scan", "Time-out days"]
    _headers(summary_sheet, 4, summary_labels)
    grouped: dict[int, list[dict]] = defaultdict(list)
    for item in rows:
        grouped[item["person_id"]].append(item)
    for row_number, person_rows in enumerate(grouped.values(), 5):
        first = person_rows[0]
        counts = Counter(item["status"] for item in person_rows)
        values = [first["lrn"] or first["external_id"], first["full_name"], first["role"], counts["Present"], counts["Late"], counts["Absent"], counts["Excused"], counts["No scan"], counts["Time Out"]]
        for column, value in enumerate(values, 1):
            summary_sheet.cell(row_number, column, value)
    target = new_report_path(EXPORT_DIR, f"Attendance-{starts_on.isoformat()}-to-{ends_on.isoformat()}.xlsx")
    book.save(target)
    return target


def attendance_report_html(db: Session, starts_on: date, ends_on: date, role: str | None = None,
                           grade: str | None = None, section: str | None = None,
                           person_id: int | None = None) -> str:
    rows = attendance_range_rows(db, starts_on, ends_on, role, grade, section, person_id)
    if starts_on == ends_on:
        headings = ["LRN / ID", "Name", "Role", "Grade / section", "Time in", "Time out", "Attendance", "Enrollment / employment", "Movement note"]
        body = ""
        for item in rows:
            values = [item["lrn"] or item["external_id"], item["full_name"], item["role"],
                      f"{item['grade'] or ''} {item['section'] or item['assignment'] or ''}", item["time_in"] or "", item["time_out"] or "",
                      item["status"], item["enrollment_status"], item["enrollment_note"] or ""]
            body += "<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in values) + "</tr>"
        return _printable_html("Daily Attendance and Movement", starts_on.isoformat(),
            "<table><thead><tr>" + "".join(f"<th>{label}</th>" for label in headings) + "</tr></thead><tbody>" + body + "</tbody></table>")
    grouped: dict[int, list[dict]] = defaultdict(list)
    for item in rows:
        grouped[item["person_id"]].append(item)
    body = ""
    for person_rows in grouped.values():
        first = person_rows[0]
        counts = Counter(item["status"] for item in person_rows)
        body += f"<tr><td>{html.escape(first['lrn'] or first['external_id'])}</td><td>{html.escape(first['full_name'])}</td><td>{html.escape(first['role'])}</td><td>{counts['Present']}</td><td>{counts['Late']}</td><td>{counts['Absent']}</td><td>{counts['Excused']}</td><td>{counts['No scan']}</td></tr>"
    return _printable_html("EduScan Attendance Summary", f"{starts_on.isoformat()} through {ends_on.isoformat()}",
                           "<table><thead><tr><th>ID / LRN</th><th>Name</th><th>Role</th><th>Present</th><th>Late</th><th>Absent</th><th>Excused</th><th>No scan</th></tr></thead><tbody>" + body + "</tbody></table>")


def _printable_html(title: str, subtitle: str, content: str) -> str:
    return f"""<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title><style>
    body{{font-family:Arial,sans-serif;margin:32px;color:#111}}h1{{text-align:center;margin-bottom:4px}}.subtitle{{text-align:center;margin-bottom:24px}}
    table{{border-collapse:collapse;width:100%;font-size:12px}}th,td{{border:1px solid #777;padding:6px;text-align:left}}th{{background:#dbeafe}}
    @media print{{button{{display:none}}body{{margin:12mm}}}}button{{padding:8px 14px;margin-bottom:18px}}
    </style></head><body><button onclick='window.print()'>Print report</button><h1>{html.escape(title)}</h1><p class='subtitle'>{html.escape(subtitle)}</p>{content}</body></html>"""
